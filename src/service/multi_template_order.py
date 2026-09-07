"""Read one mixed-template workbook and build stable template groups."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from types import MappingProxyType
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from .multi_template_excel_reader import read_sheet_rows
from .v2_template_store_utils import safe_segment


TEMPLATE_COLUMN = "模板"
PARENT_JOB_STATUSES = frozenset({"preflighting", "preflight_failed", "ready", "canary_running", "running", "completed", "completed_with_errors", "failed", "interrupted"})
TEMPLATE_GROUP_STATUSES = frozenset({"pending", "canary_running", "canary_failed", "ready", "running", "succeeded", "failed", "interrupted"})
FAILURE_SCOPES = frozenset({"template", "system"})
ORDER_NUMBER_COLUMNS = ("内部订单号", "订单号", "订单编号", "Order", "Order ID", "OrderId", "order_id")


@dataclass(frozen=True)
class MultiTemplateIssue:
    code: str
    message: str
    suggestion: str
    template_id: str = ""
    excel_row: int = 0
    order_no: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {"code": self.code, "message": self.message, "suggestion": self.suggestion, "template_id": self.template_id, "excel_row": self.excel_row, "order_no": self.order_no}

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "MultiTemplateIssue":
        return cls(str(payload.get("code") or ""), str(payload.get("message") or ""), str(payload.get("suggestion") or ""), str(payload.get("template_id") or ""), _positive_int(payload.get("excel_row")), str(payload.get("order_no") or ""))


@dataclass(frozen=True)
class MultiTemplateOrderRow:
    sheet_name: str
    excel_row: int
    order_no: str
    template_id: str
    values: Mapping[str, Any]
    raw_values: tuple[Any, ...] = ()

    def __post_init__(self) -> None:
        object.__setattr__(self, "values", MappingProxyType(dict(self.values)))
        object.__setattr__(self, "raw_values", tuple(self.raw_values))

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "excel_row": self.excel_row,
            "order_no": self.order_no,
            "template_id": self.template_id,
            "values": _json_value(dict(self.values)),
            "raw_values": _json_value(self.raw_values),
        }

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "MultiTemplateOrderRow":
        values = _from_json_value(payload.get("values"))
        raw_values = _from_json_value(payload.get("raw_values"))
        return cls(
            str(payload.get("sheet_name") or ""),
            _positive_int(payload.get("excel_row")),
            str(payload.get("order_no") or ""),
            str(payload.get("template_id") or ""),
            dict(values) if isinstance(values, Mapping) else {},
            tuple(raw_values) if isinstance(raw_values, list) else (),
        )


@dataclass(frozen=True)
class TemplateOrderGroup:
    template_id: str
    rows: tuple[MultiTemplateOrderRow, ...]

    def __post_init__(self) -> None:
        rows = tuple(self.rows)
        if not rows or any(row.template_id != self.template_id for row in rows):
            raise ValueError("模板分组必须包含同一模板的订单行")
        object.__setattr__(self, "rows", rows)

    def to_dict(self) -> dict[str, Any]:
        return {"template_id": self.template_id, "rows": [row.to_dict() for row in self.rows]}

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "TemplateOrderGroup":
        rows = payload.get("rows") or []
        return cls(str(payload.get("template_id") or ""), tuple(MultiTemplateOrderRow.from_dict(row) for row in rows if isinstance(row, Mapping)))


@dataclass(frozen=True)
class MultiTemplateOrderBatch:
    sheet_name: str
    headers: tuple[str, ...]
    rows: tuple[MultiTemplateOrderRow, ...]
    groups: tuple[TemplateOrderGroup, ...]
    issues: tuple[MultiTemplateIssue, ...]

    def __post_init__(self) -> None:
        object.__setattr__(self, "headers", tuple(self.headers))
        object.__setattr__(self, "rows", tuple(self.rows))
        object.__setattr__(self, "groups", tuple(self.groups))
        object.__setattr__(self, "issues", tuple(self.issues))

    def to_dict(self) -> dict[str, Any]:
        return {"sheet_name": self.sheet_name, "headers": list(self.headers), "rows": [row.to_dict() for row in self.rows], "groups": [group.to_dict() for group in self.groups], "issues": [issue.to_dict() for issue in self.issues]}

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "MultiTemplateOrderBatch":
        return cls(
            str(payload.get("sheet_name") or ""),
            tuple(_header_text(value) for value in payload.get("headers") or []),
            tuple(MultiTemplateOrderRow.from_dict(row) for row in payload.get("rows") or [] if isinstance(row, Mapping)),
            tuple(TemplateOrderGroup.from_dict(group) for group in payload.get("groups") or [] if isinstance(group, Mapping)),
            tuple(MultiTemplateIssue.from_dict(issue) for issue in payload.get("issues") or [] if isinstance(issue, Mapping)),
        )


class MultiTemplateOrderParser:
    """Parse the selected sheet once without invoking template-specific parsers."""

    def parse(self, order_file: Path | str, *, sheet_name: str = "") -> MultiTemplateOrderBatch:
        source = Path(order_file)
        if not source.is_file():
            return _failed_batch("order_file_missing", "订单表格不存在，请重新选择后再试。", "请重新上传订单表格。")
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except Exception:
            return _failed_batch("order_file_unreadable", "订单表格无法读取，请确认文件格式后重试。", "请上传可读取的 Excel 文件。")
        try:
            if not workbook.sheetnames:
                return _failed_batch("orders_empty", "订单表格为空，请至少提供表头。", "请上传包含表头和订单行的表格。")
            if sheet_name and sheet_name not in workbook.sheetnames:
                return _failed_batch("sheet_not_found", "指定的订单工作表不存在。", "请选择订单表中存在的工作表。")
            if not sheet_name and len(workbook.sheetnames) > 1:
                return _failed_batch("sheet_name_required", "订单表包含多个工作表，请先选择订单所在工作表。", "请选择包含“模板”列的订单工作表。")
            resolved_sheet_name = sheet_name or workbook.sheetnames[0]
            return self._parse_sheet(resolved_sheet_name, read_sheet_rows(workbook, workbook[resolved_sheet_name]))
        finally:
            workbook.close()

    def _parse_sheet(
        self,
        sheet_name: str,
        parsed_rows: Iterable[tuple[int, tuple[Any, ...], tuple[bool, ...]]],
    ) -> MultiTemplateOrderBatch:
        rows = tuple(parsed_rows)
        headers = tuple(_header_text(value) for excel_row, values, _ in rows if excel_row == 1 for value in values)
        if not headers:
            return _failed_batch("orders_empty", "订单表格为空，请至少提供表头。", "请上传包含表头和订单行的表格。", sheet_name)
        template_indexes = [index for index, header in enumerate(headers) if header == TEMPLATE_COLUMN]
        if not template_indexes:
            return _failed_batch("template_column_missing", "订单表缺少“模板”列。", "请新增名为“模板”的列。", sheet_name)
        if len(template_indexes) != 1:
            return _failed_batch("template_column_duplicate", "订单表存在多个“模板”列。", "请只保留一个名为“模板”的列。", sheet_name)
        return self._parse_rows(sheet_name, headers, template_indexes[0], rows)

    def _parse_rows(
        self,
        sheet_name: str,
        headers: tuple[str, ...],
        template_index: int,
        parsed_rows: Iterable[tuple[int, tuple[Any, ...], tuple[bool, ...]]],
    ) -> MultiTemplateOrderBatch:
        rows: list[MultiTemplateOrderRow] = []
        issues: list[MultiTemplateIssue] = []
        grouped: dict[str, list[MultiTemplateOrderRow]] = {}
        for excel_row, raw_values, formula_cells in parsed_rows:
            if excel_row <= 1 or not any(_text(value) for value in raw_values):
                continue
            mapped = {header: raw_values[index] for index, header in enumerate(headers)}
            order_no = _order_no(mapped)
            missing_formula_indexes = [
                index
                for index, has_formula in enumerate(formula_cells)
                if has_formula and raw_values[index] is None
            ]
            if missing_formula_indexes:
                issues.append(
                    MultiTemplateIssue(
                        "formula_value_missing",
                        "订单表中存在未计算的公式，无法提供给现有渲染流程。",
                        "请使用 Excel 重新计算并保存订单表后再上传。",
                        _text(raw_values[template_index]),
                        excel_row,
                        order_no,
                    )
                )
                continue
            template_id = _text(raw_values[template_index])
            issue = _template_issue(template_id, excel_row, order_no)
            if issue:
                issues.append(issue)
                continue
            row = MultiTemplateOrderRow(sheet_name, excel_row, order_no, template_id, mapped, raw_values)
            rows.append(row)
            grouped.setdefault(template_id, []).append(row)
        groups = tuple(TemplateOrderGroup(template_id, tuple(group)) for template_id, group in grouped.items())
        return MultiTemplateOrderBatch(sheet_name, headers, tuple(rows), groups, tuple(issues))


def parse_multi_template_orders(order_file: Path | str, *, sheet_name: str = "") -> MultiTemplateOrderBatch:
    return MultiTemplateOrderParser().parse(order_file, sheet_name=sheet_name)


def _failed_batch(code: str, message: str, suggestion: str, sheet_name: str = "") -> MultiTemplateOrderBatch:
    return MultiTemplateOrderBatch(sheet_name, (), (), (), (MultiTemplateIssue(code, message, suggestion, excel_row=1),))


def _template_issue(template_id: str, excel_row: int, order_no: str) -> MultiTemplateIssue | None:
    if not template_id:
        return MultiTemplateIssue("template_id_missing", "模板 ID 不能为空。", "请在“模板”列填写已启用模板 ID。", excel_row=excel_row, order_no=order_no)
    if safe_segment(template_id) != template_id:
        return MultiTemplateIssue("template_id_invalid", "模板 ID 格式不合法。", "请填写已有模板的完整 ID。", template_id, excel_row, order_no)
    return None


def _order_no(values: Mapping[str, Any]) -> str:
    return next(
        (_text(value) for header, value in values.items() if _header_text(header).strip() in ORDER_NUMBER_COLUMNS and _text(value)),
        "",
    )


def _positive_int(value: Any) -> int:
    try:
        return max(int(value), 0)
    except (TypeError, ValueError):
        return 0


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _header_text(value: Any) -> str:
    return str(value) if value is not None else ""


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return {"__multi_template_type__": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"__multi_template_type__": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"__multi_template_type__": "time", "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {"__multi_template_type__": "timedelta", "value": value.total_seconds()}
    if isinstance(value, Decimal):
        return {"__multi_template_type__": "decimal", "value": str(value)}
    if isinstance(value, Mapping):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, tuple):
        return [_json_value(item) for item in value]
    if isinstance(value, list):
        return [_json_value(item) for item in value]
    return value


def _from_json_value(value: Any) -> Any:
    if isinstance(value, list):
        return [_from_json_value(item) for item in value]
    if not isinstance(value, Mapping):
        return value
    kind = value.get("__multi_template_type__")
    raw = value.get("value")
    if kind == "datetime":
        return datetime.fromisoformat(str(raw))
    if kind == "date":
        return date.fromisoformat(str(raw))
    if kind == "time":
        return time.fromisoformat(str(raw))
    if kind == "timedelta":
        return timedelta(seconds=float(raw))
    if kind == "decimal":
        return Decimal(str(raw))
    return {str(key): _from_json_value(item) for key, item in value.items()}


__all__ = ["FAILURE_SCOPES", "PARENT_JOB_STATUSES", "TEMPLATE_COLUMN", "TEMPLATE_GROUP_STATUSES", "MultiTemplateIssue", "MultiTemplateOrderBatch", "MultiTemplateOrderParser", "MultiTemplateOrderRow", "TemplateOrderGroup", "parse_multi_template_orders"]
