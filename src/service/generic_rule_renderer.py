"""Build generic render tasks directly from confirmed template rule packs."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, Mapping

from openpyxl import load_workbook

from .font_style_rules import font_style_by_option
from .name_color_cycle import normalize_name_color_cycle
from .template_registry import TemplateDefinition
from .template_rule_execution import resolve_mapped_text
from .template_rule_ast import RULE_AST_SCHEMA, runtime_actions


class GenericRuleRenderError(ValueError):
    pass


_MISSING = object()

IMPLICIT_ORDER_BINDINGS = {
    "order_no": (
        "order_no",
        "order",
        "order no",
        "order number",
        "\u8ba2\u5355\u53f7",
        "\u5185\u90e8\u8ba2\u5355\u53f7",
    ),
    "font": (
        "font",
        "font option",
        "font options",
        "\u5b57\u4f53",
        "\u5b57\u4f53\u9009\u9879",
    ),
    "design": ("design", "design option", "\u8bbe\u8ba1", "\u8bbe\u8ba1\u9009\u9879"),
    "style": ("style", "style option", "\u5c3a\u5bf8", "\u6b3e\u5f0f"),
    "color": ("color", "color option", "\u989c\u8272", "\u5b57\u4f53\u989c\u8272"),
    "department": ("department", "production department", "\u751f\u4ea7\u90e8\u95e8", "\u90e8\u95e8"),
    "text": (
        "text",
        "name",
        "names",
        "personalization",
        "\u540d\u5b57",
        "\u5b9a\u5236\u4fe1\u606f",
        "\u5b9a\u5236\u5185\u5bb9",
    ),
}

TEMPLATE_COLUMN_ALIASES = ("template", "\u6a21\u677f")


def build_generic_render_task(
    template: TemplateDefinition,
    rules: Mapping[str, Any],
    order_file: Path,
    output_ai: Path,
    *,
    sheet_name: str = "",
    columns: int = 4,
) -> Dict[str, Any]:
    if not template.template_ai or not template.template_ai.exists():
        raise GenericRuleRenderError("Template AI file does not exist.")
    bindings = rules.get("order_bindings", {})
    if not isinstance(bindings, Mapping) or not bindings:
        raise GenericRuleRenderError("Confirmed rules do not define order_bindings.")
    rows = _read_rows(order_file, sheet_name=sheet_name)
    _require_columns(rows, bindings.values())
    rows = _filter_rows_for_template(rows, template.template_id)
    assets = _asset_catalog(template.assets)
    orders = [
        _build_order(index, row, rules, bindings, assets)
        for index, row in enumerate(rows, start=1)
        if any(value not in (None, "") for value in row.values())
    ]
    if not orders:
        raise GenericRuleRenderError("Order sheet contains no data rows.")
    render_layout = _mapping(rules.get("render_layout"))
    if render_layout:
        orders = _build_layout_orders(orders, render_layout)
    single_output = str(render_layout.get("output_mode") or "").strip() == "single_file"
    for index, order in enumerate(orders, start=1):
        order["output_ai"] = str(
            (
                output_ai
                if single_output or len(orders) == 1
                else output_ai.with_name(f"{output_ai.stem}-{index:03d}{output_ai.suffix}")
            ).resolve()
        )
    return {
        "type": "generic_template_rules",
        "template_id": template.template_id,
        "template_ai": str(template.template_ai.resolve()),
        "output_ai": str(output_ai.resolve()),
        "output_ai_files": [str(output_ai.resolve())] if single_output else [order["output_ai"] for order in orders],
        "orders": orders,
        "option_groups": _list_of_mappings(rules.get("option_groups")),
        "dimensions": _mapping(rules.get("dimensions")),
        "text_policies": _mapping(rules.get("text_policies")),
        "render_layout": render_layout,
        "font_option_styles": _mapping(rules.get("font_option_styles")),
        "allow_unnamed_name_fallback": bool(rules.get("allow_unnamed_name_fallback", False)),
        "layout": {"columns": max(1, int(columns)), "gap_mm": 8.0, "margin_mm": 8.0},
        "output": _mapping(rules.get("output")),
        "transforms": _mapping(rules.get("transforms")),
    }


def _build_layout_orders(orders: list[Dict[str, Any]], layout: Mapping[str, Any]) -> list[Dict[str, Any]]:
    """Group source rows into declaratively configured output sheets."""

    if str(layout.get("type") or "").strip() != "name_columns":
        raise GenericRuleRenderError(f"Unsupported render layout: {layout.get('type')}")
    groups: Dict[tuple[str, ...], Dict[str, Any]] = {}
    for order in orders:
        mode = _layout_mode(order, layout)
        group_by = mode.get("group_by", ["row"])
        if not isinstance(group_by, list) or not group_by:
            group_by = ["row"]
        key = tuple(_layout_group_value(order, field) for field in group_by)
        bucket = groups.setdefault(key, {"mode": mode, "members": []})
        bucket["members"].append(order)

    result: list[Dict[str, Any]] = []
    for bucket in groups.values():
        members = bucket["members"]
        leader = dict(members[0])
        leader["layout_members"] = members
        leader["layout_mode"] = bucket["mode"]
        result.append(leader)
    return result


def _layout_mode(order: Mapping[str, Any], layout: Mapping[str, Any]) -> Dict[str, Any]:
    base = _mapping(layout.get("default"))
    department = str(_mapping(order.get("values")).get("department") or "").strip().casefold()
    overrides = _mapping(layout.get("department_overrides"))
    for name, override in overrides.items():
        if str(name).strip().casefold() == department and isinstance(override, Mapping):
            return {**base, **dict(override)}
    return base


def _layout_group_value(order: Mapping[str, Any], field: Any) -> str:
    if str(field) == "row":
        return str(order.get("row_index") or "")
    values = _mapping(order.get("values"))
    selections = _mapping(order.get("selections"))
    return str(values.get(str(field), selections.get(str(field), "")) or "").strip().casefold()


def _build_order(
    index: int,
    row: Mapping[str, Any],
    rules: Mapping[str, Any],
    bindings: Mapping[str, Any],
    assets: Mapping[str, str],
) -> Dict[str, Any]:
    values = _bound_values(row, bindings)
    selections = {
        key: ("" if values.get(key) is None else str(values.get(key))).strip()
        for key in ("font", "design", "style", "color")
        if ("" if values.get(key) is None else str(values.get(key))).strip()
    }
    rule_ast = rules.get("rule_ast")
    use_rule_ast = isinstance(rule_ast, Mapping) and rule_ast.get("$schema") == RULE_AST_SCHEMA
    font_styles = {} if use_rule_ast else font_style_by_option(
        rules.get("font_style_rules"), legacy_option_overrides=rules.get("option_overrides")
    )
    variables = _build_variables(
        values,
        rules,
        font_style=font_styles.get(selections.get("font", "")),
        rule_ast=rule_ast if use_rule_ast else None,
        selections={**selections, "text": values.get("text")},
    )
    asset_tasks = []
    selected_options = set(selections.values())
    for mapping in _list_of_mappings(rules.get("asset_mappings")):
        selected_option = str(mapping.get("option") or "")
        if selected_option not in selected_options:
            continue
        asset_name = str(mapping.get("asset") or "")
        if asset_name not in assets:
            raise GenericRuleRenderError(f"Mapped asset is not registered: {asset_name}")
        asset_tasks.append(
            {
                "option": selected_option,
                "asset": assets[asset_name],
                "target": str(mapping.get("target") or selected_option),
            }
        )
    return {
        "row_index": index,
        "order_no": str(values.get("order_no") or f"ROW-{index}"),
        "values": values,
        "selections": selections,
        "variables": variables,
        "assets": asset_tasks,
        "transforms": _effective_transforms(rules.get("transforms"), selections),
    }


def _build_variables(
    values: Mapping[str, Any],
    rules: Mapping[str, Any],
    *,
    font_style: Mapping[str, Any] | None = None,
    rule_ast: Any = None,
    selections: Mapping[str, Any] | None = None,
) -> list[Dict[str, Any]]:
    mappings = _list_of_mappings(rules.get("slot_mappings"))
    if not mappings:
        targets = _list_of_mappings(rules.get("text_targets"))
        if len(targets) == 1 and "text" in values:
            mappings = [{"field": "text", "slot": targets[0].get("name", "")}]
    name_color_cycle = normalize_name_color_cycle(rules.get("name_color_cycle")) if rule_ast is None else {}
    variables = []
    for mapping in mappings:
        field = str(mapping.get("field") or mapping.get("source") or "")
        target = str(mapping.get("slot") or mapping.get("name") or "")
        if not field or not target or field not in values:
            continue
        text_policies = rules.get("text_policies", {})
        text_policies = text_policies if isinstance(text_policies, Mapping) else {}
        resolved, value = resolve_mapped_text(values.get(field), mapping, text_policies)
        if not resolved:
            raise GenericRuleRenderError(
                f"Order value cannot satisfy split policy for target: {target}"
            )
        variable = {"target": target, "field": field, "value": value}
        actions = runtime_actions(rule_ast, target=target, selections=selections or {}) if rule_ast is not None else []
        if actions:
            variable["actions"] = actions
        if target == "Name" and name_color_cycle:
            variable["name_color_cycle"] = name_color_cycle
        if font_style:
            variable["font_style"] = dict(font_style)
        variables.append(variable)
    if not variables:
        raise GenericRuleRenderError("Order row does not produce any template variables.")
    return variables


def _read_rows(path: Path, *, sheet_name: str) -> list[Dict[str, Any]]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise GenericRuleRenderError("Generic rule renderer currently requires an .xlsx order file.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise GenericRuleRenderError(f"Order sheet does not exist: {sheet_name}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]]


def _require_columns(rows: list[Mapping[str, Any]], columns: Iterable[Any]) -> None:
    if not rows:
        return
    available = _column_lookup(rows[0])
    missing = [str(column) for column in columns if _normalize_column(column) not in available]
    if missing:
        raise GenericRuleRenderError(f"Order sheet is missing bound columns: {', '.join(missing)}")


def _filter_rows_for_template(rows: list[Dict[str, Any]], template_id: str) -> list[Dict[str, Any]]:
    if not rows:
        return rows
    lookup = _column_lookup(rows[0])
    template_column = next(
        (lookup[_normalize_column(alias)] for alias in TEMPLATE_COLUMN_ALIASES if _normalize_column(alias) in lookup),
        "",
    )
    if not template_column:
        return rows
    target = str(template_id or "").strip().casefold()
    return [
        row
        for row in rows
        if not str(row.get(template_column) or "").strip()
        or str(row.get(template_column) or "").strip().casefold() == target
    ]


def _bound_values(row: Mapping[str, Any], bindings: Mapping[str, Any]) -> Dict[str, Any]:
    values: Dict[str, Any] = {}
    for field, column in bindings.items():
        values[str(field)] = _column_value(row, str(column))
    for field, aliases in IMPLICIT_ORDER_BINDINGS.items():
        if values.get(field) not in (None, ""):
            continue
        value = _first_column_value(row, aliases)
        if value is not _MISSING:
            values[field] = value
    return values


def _first_column_value(row: Mapping[str, Any], aliases: Iterable[Any]) -> Any:
    for alias in aliases:
        value = _column_value(row, str(alias), default=_MISSING)
        if value is not _MISSING:
            return value
    return _MISSING


def _column_value(row: Mapping[str, Any], column: str, *, default: Any = None) -> Any:
    if column in row:
        return row.get(column)
    actual = _column_lookup(row).get(_normalize_column(column))
    return row.get(actual) if actual is not None else default


def _column_lookup(row: Mapping[str, Any]) -> Dict[str, str]:
    return {_normalize_column(key): str(key) for key in row}


def _normalize_column(value: Any) -> str:
    return str(value or "").strip().casefold()


def _asset_catalog(items: Any) -> Dict[str, str]:
    catalog: Dict[str, str] = {}
    for item in _list_of_mappings(items):
        path = str(item.get("stored_path") or "")
        if not path:
            continue
        resolved = Path(path)
        if not resolved.is_absolute():
            resolved = Path(__file__).resolve().parents[2] / resolved
        if not resolved.exists():
            raise GenericRuleRenderError(f"Template asset file does not exist: {resolved}")
        catalog[str(item.get("file_name") or resolved.name)] = str(resolved.resolve())
        catalog[path] = str(resolved.resolve())
    return catalog


def _effective_transforms(value: Any, selections: Mapping[str, str]) -> Dict[str, Any]:
    transforms = _mapping(value)
    effective = {
        key: item
        for key, item in transforms.items()
        if not isinstance(item, Mapping)
    }
    targets = []
    for selected in selections.values():
        settings = transforms.get(selected)
        if isinstance(settings, Mapping):
            targets.append({"target": selected, **dict(settings)})
    if targets:
        effective["targets"] = targets
    return effective


def _mapping(value: Any) -> Dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _list_of_mappings(value: Any) -> list[Dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]
