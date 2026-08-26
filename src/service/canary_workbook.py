"""Create and verify the isolated one-order workbook used by a canary."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .multi_template_order import TEMPLATE_COLUMN
from .v2_template_store_utils import safe_segment


class CanaryWorkbookIntegrityError(RuntimeError):
    """Raised when the isolated workbook changes around the adapter boundary."""


def write_canary_workbook(
    *,
    work_dir: Path | str,
    sheet_name: str,
    headers: tuple[str, ...],
    template_id: str,
    row_values: tuple[Any, ...],
) -> Path:
    if len(row_values) != len(headers):
        raise ValueError("canary row values are incomplete")
    safe_id = safe_segment(template_id)
    if safe_id != template_id:
        raise ValueError("canary template id is unsafe")
    output = Path(work_dir).resolve() / "canary" / safe_id / "orders.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(list(headers))
        sheet.append(list(row_values))
        workbook.save(output)
    finally:
        workbook.close()
    _verify_canary_workbook(output, sheet_name, headers, template_id, row_values)
    return output


def require_canary_workbook_snapshot(
    path: Path,
    expected_sha256: str,
    sheet_name: str,
    headers: tuple[str, ...],
    template_id: str,
    row_values: tuple[Any, ...],
) -> None:
    try:
        _verify_canary_workbook(path, sheet_name, headers, template_id, row_values)
    except Exception as exc:
        raise CanaryWorkbookIntegrityError("canary workbook content changed") from exc
    if sha256_file(path) != expected_sha256:
        raise CanaryWorkbookIntegrityError("canary workbook digest changed")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _verify_canary_workbook(
    path: Path,
    sheet_name: str,
    headers: tuple[str, ...],
    template_id: str,
    row_values: tuple[Any, ...],
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()
    template_index = headers.index(TEMPLATE_COLUMN)
    if rows != [headers, row_values] or str(rows[1][template_index] or "").strip() != template_id:
        raise RuntimeError("canary workbook verification failed")


__all__ = [
    "CanaryWorkbookIntegrityError",
    "require_canary_workbook_snapshot",
    "sha256_file",
    "write_canary_workbook",
]
