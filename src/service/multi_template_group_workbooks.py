"""Write value-only workbooks for each validated template group."""

from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .multi_template_order import TEMPLATE_COLUMN, MultiTemplateOrderBatch, TemplateOrderGroup
from .v2_template_store_utils import safe_segment


class GroupWorkbookWriter:
    def write(self, batch: MultiTemplateOrderBatch, parent_job_dir: Path | str) -> dict[str, Path]:
        if batch.issues:
            raise ValueError("订单存在解析问题，不能生成分组工作簿")
        if TEMPLATE_COLUMN not in batch.headers:
            raise ValueError("订单分组缺少模板列")
        root = Path(parent_job_dir).resolve() / "groups"
        paths: dict[str, Path] = {}
        for group in batch.groups:
            output = self._output_path(root, group.template_id)
            output.parent.mkdir(parents=True, exist_ok=True)
            self._write(output, batch, group)
            self._verify(output, batch, group)
            paths[group.template_id] = output
        return paths

    def _output_path(self, root: Path, template_id: str) -> Path:
        safe_id = safe_segment(template_id)
        if safe_id != template_id:
            raise ValueError("模板 ID 不能用于生成安全分组目录")
        digest = hashlib.sha256(template_id.encode("utf-8")).hexdigest()[:12]
        return root / f"{safe_id}-{digest}" / "orders.xlsx"

    def _write(self, output: Path, batch: MultiTemplateOrderBatch, group: TemplateOrderGroup) -> None:
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = batch.sheet_name
            sheet.append(list(batch.headers))
            for row in group.rows:
                sheet.append(list(row.raw_values))
            workbook.save(output)
        finally:
            workbook.close()

    def _verify(self, path: Path, batch: MultiTemplateOrderBatch, group: TemplateOrderGroup) -> None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook[batch.sheet_name].iter_rows(values_only=True))
        finally:
            workbook.close()
        template_index = batch.headers.index(TEMPLATE_COLUMN)
        actual_ids = {str(row[template_index] or "").strip() for row in rows[1:]}
        if len(rows) - 1 != len(group.rows) or actual_ids != {group.template_id}:
            raise RuntimeError("分组工作簿回读校验失败")


__all__ = ["GroupWorkbookWriter"]
