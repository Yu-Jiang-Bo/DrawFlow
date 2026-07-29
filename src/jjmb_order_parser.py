"""Parser for JJMB202603281027102517 template order sheets."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List

import openpyxl


FIELD_ALIASES = {
    "order_no": ["内部订单号", "订单号", "订单编号"],
    "detail_id": ["订单明细id", "订单明细ID", "明细id"],
    "quantity": ["购买数量", "数量"],
    "template": ["模板"],
    "custom_info": ["定制信息"],
    "department": ["生产部门", "部门", "department", "production department"],
    "manufacturer": ["外协厂家代码", "厂家代码", "厂家", "厂商", "生产厂家", "供应商", "manufacturer", "factory", "supplier"],
    "product_name": ["产品中文名称", "产品名称", "product name", "product_name"],
    "color_option": ["字体颜色", "颜色", "color", "font color", "color_option"],
}

SPLIT_CUSTOM_FIELD_ALIASES = {
    "style_option": ["Style Option", "Style", "style_option", "款式", "设计款式"],
    "font_option": ["Font Option", "Font Options", "Font", "font_option", "字体", "字体选项"],
    "personalization": ["Personalization", "Name", "Names", "定制内容", "定制信息", "名字"],
}


@dataclass(frozen=True)
class JJMBOrderItem:
    order_no: str
    detail_id: str
    quantity: int
    template: str
    style_option: str
    font_option: str
    personalization_values: List[str]
    department: str = ""
    manufacturer: str = ""
    product_name: str = ""
    color_option: str = ""


def read_xlsx_rows(path: Path | str, sheet_name: str | None = None) -> List[Dict[str, str]]:
    workbook = openpyxl.load_workbook(Path(path), data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    result: List[Dict[str, str]] = []
    for row in rows[1:]:
        values = ["" if value is None else str(value).strip() for value in row]
        mapped = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
        if any(mapped.values()):
            result.append(mapped)
    return result


def normalize_row(row: Dict[str, str]) -> Dict[str, str]:
    lower_map = {key.strip().lower(): value for key, value in row.items()}
    normalized: Dict[str, str] = {}
    for target, aliases in FIELD_ALIASES.items():
        normalized[target] = ""
        for alias in aliases:
            value = lower_map.get(alias.strip().lower(), "")
            if value:
                normalized[target] = value
                break
    return normalized


def parse_custom_info(value: str) -> Dict[str, str]:
    result = {"style_option": "", "font_option": "", "personalization": ""}
    current_key: str | None = None
    current_lines: List[str] = []

    def flush() -> None:
        nonlocal current_key, current_lines
        if current_key:
            result[current_key] = "\n".join(current_lines).strip()
        current_key = None
        current_lines = []

    for raw_line in (value or "").replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^(Style Option|Font Option|Personalization)\s*:\s*(.*)$", line, re.I)
        if match:
            flush()
            key = match.group(1).lower()
            if key == "style option":
                current_key = "style_option"
            elif key == "font option":
                current_key = "font_option"
            else:
                current_key = "personalization"
            current_lines = [match.group(2).strip()] if match.group(2).strip() else []
        else:
            if current_key:
                current_lines.append(line)
    flush()
    return result


def parse_split_custom_columns(row: Dict[str, str], base: Dict[str, str] | None = None) -> Dict[str, str]:
    result = dict(base or {"style_option": "", "font_option": "", "personalization": ""})
    for target, aliases in SPLIT_CUSTOM_FIELD_ALIASES.items():
        if result.get(target):
            continue
        result[target] = first_row_value(row, aliases)
    return result


def first_row_value(row: Dict[str, str], aliases: List[str]) -> str:
    lower_map = {key.strip().lower(): value for key, value in row.items()}
    for alias in aliases:
        value = lower_map.get(alias.strip().lower(), "")
        if value:
            return value
    return ""


def normalize_option(value: str, prefix: str) -> str:
    compact = re.sub(r"\s+", "", value or "")
    match = re.search(re.escape(prefix) + r"(\d+)", compact, re.I)
    if not match:
        return ""
    return f"{prefix}{int(match.group(1))}"


def parse_quantity(value: str) -> int:
    try:
        quantity = int(float((value or "").strip()))
    except ValueError:
        return 1
    return max(quantity, 1)


def split_personalization(value: str) -> List[str]:
    text = (value or "").strip()
    if not text:
        return []
    lines = [line.strip() for line in text.replace("\r\n", "\n").split("\n") if line.strip()]
    if len(lines) <= 1 and "|" in text:
        lines = [part.strip() for part in text.split("|") if part.strip()]
    cleaned = [strip_list_marker(line) for line in lines]
    return [line for line in cleaned if line]


def strip_list_marker(value: str) -> str:
    return re.sub(r"^\s*\d{1,3}\s*(?:\.\s+|[\)\]\u3001:-]\s*)", "", value or "").strip()


def parse_order_items(rows: Iterable[Dict[str, str]], template_id: str | None = None) -> List[JJMBOrderItem]:
    items: List[JJMBOrderItem] = []
    for row in rows:
        normalized = normalize_row(row)
        if template_id and normalized["template"] != template_id:
            continue
        custom = parse_split_custom_columns(row, parse_custom_info(normalized["custom_info"]))
        style = normalize_option(custom["style_option"], "Style")
        font = normalize_option(custom["font_option"], "F")
        values = split_personalization(custom["personalization"])
        if not style or not font or not values:
            continue
        items.append(
            JJMBOrderItem(
                order_no=normalized["order_no"],
                detail_id=normalized["detail_id"],
                quantity=parse_quantity(normalized["quantity"]),
                template=normalized["template"],
                style_option=style,
                font_option=font,
                personalization_values=values,
                department=normalized["department"],
                manufacturer=normalized["manufacturer"],
                product_name=normalized["product_name"],
                color_option=normalized["color_option"],
            )
        )
    return items
