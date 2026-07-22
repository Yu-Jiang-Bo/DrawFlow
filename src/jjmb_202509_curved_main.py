"""Render JJMB202509231236046265 curved-title order sheets."""

from __future__ import annotations

import argparse
import html
import json
import re
from collections import OrderedDict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Mapping

from openpyxl import load_workbook

from .renderer.illustrator_bridge import IllustratorBridge, IllustratorBridgeError
from .service.render_integrity import RenderIntegrityError, compare_png_previews


TEMPLATE_ID = "JJMB202509231236046265"
DEFAULT_FONT_OPTION = "F1"
DEFAULT_DEPARTMENT = "ZW"
DEFAULT_TITLE = "Merry Christmas"
QUANTITY_FIELDS = ("购买数量", "数量", "Quantity", "Qty")


@dataclass(frozen=True)
class CurvedOrderItem:
    order_no: str
    detail_id: str
    department: str
    font_option: str
    text: str
    text_type: str
    quantity_index: int
    copy_index: int = 1

    def to_json_dict(self) -> Dict[str, object]:
        return {
            "order_no": self.order_no,
            "detail_id": self.detail_id,
            "department": self.department,
            "font_option": self.font_option,
            "text": self.text,
            "text_type": self.text_type,
            "quantity_index": self.quantity_index,
            "copy_index": self.copy_index,
        }


@dataclass(frozen=True)
class CurvedOrderGroup:
    order_no: str
    items: List[CurvedOrderItem]
    group_key: str = ""

    def to_json_dict(self) -> Dict[str, object]:
        return {
            "order_no": self.order_no,
            "items": [item.to_json_dict() for item in self.items],
        }


def read_xlsx_rows(path: Path, sheet_name: str | None = None) -> List[Dict[str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    result: List[Dict[str, str]] = []
    for row in rows[1:]:
        mapped = {
            headers[index]: "" if value is None else str(value).strip()
            for index, value in enumerate(row[: len(headers)])
        }
        if any(mapped.values()):
            result.append(mapped)
    return result


def get_field(row: Dict[str, str], *names: str) -> str:
    for name in names:
        value = row.get(name, "")
        if value:
            return value
    return ""


def parse_custom_info(value: str) -> Dict[str, str]:
    result = {"font": "", "title": "", "names": "", "style": ""}
    current_key = ""
    current_lines: List[str] = []

    def flush() -> None:
        nonlocal current_key, current_lines
        if current_key:
            result[current_key] = "\n".join(current_lines).strip()
        current_key = ""
        current_lines = []

    for raw_line in html.unescape(value or "").replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^(Style Option|Font Options?|Title|Names?)\s*[:：]\s*(.*)$", line, re.I)
        if match:
            flush()
            key = match.group(1).lower()
            if key == "style option":
                current_key = "style"
            elif key.startswith("font option"):
                current_key = "font"
            elif key == "title":
                current_key = "title"
            else:
                current_key = "names"
            current_lines = [match.group(2).strip()] if match.group(2).strip() else []
        elif current_key:
            current_lines.append(line)
    flush()
    return result


def normalize_font(value: str) -> str:
    match = re.search(r"F\s*(\d+)", value or "", re.I)
    if not match:
        return DEFAULT_FONT_OPTION
    index = int(match.group(1))
    if index < 1 or index > 14:
        return DEFAULT_FONT_OPTION
    return f"F{index}"


def split_names(value: str) -> List[str]:
    raw = html.unescape(value or "").strip()
    if not raw:
        return []
    lines = [line.strip() for line in raw.replace("\r\n", "\n").split("\n") if line.strip()]
    if len(lines) > 1:
        result: List[str] = []
        for line in lines:
            result.extend(split_names_inline(line))
        return result
    return split_names_inline(raw)


def split_names_inline(value: str) -> List[str]:
    text = re.sub(r"\s+", " ", html.unescape(value or "")).strip()
    if not text:
        return []
    if re.search(r"(?:^|\s)\d{1,3}\s*[\.\)\]\u3001:]?\s*", text):
        parts = re.split(r"(?:^|\s)\d{1,3}\s*[\.\)\]\u3001:]?\s*", text)
    elif "|" in text:
        parts = text.split("|")
    elif "," in text:
        parts = text.split(",")
    else:
        parts = [text]
    return [clean_text(part) for part in parts if clean_text(part)]


def clean_text(value: str) -> str:
    text = html.unescape(value or "").strip()
    text = re.sub(r"^\s*\d{1,3}\s*[\.\)\]\u3001:]?\s*", "", text)
    return re.sub(r"\s+", " ", text).strip(" ,;")


def parse_items(
    rows: Iterable[Dict[str, str]],
    *,
    multi_name_customization: bool = False,
) -> List[CurvedOrderItem]:
    items: List[CurvedOrderItem] = []
    for row in rows:
        if get_field(row, "模板") != TEMPLATE_ID:
            continue
        custom = parse_custom_info(get_field(row, "定制信息"))
        font_option = normalize_font(custom["font"])
        order_no = get_field(row, "内部订单号", "订单号")
        detail_id = get_field(row, "订单明细id", "订单明细ID")
        department = get_field(row, "生产部门") or DEFAULT_DEPARTMENT
        copy_count = _row_quantity(row) if multi_name_customization else 1
        names = split_names(custom["names"])
        title = clean_text(custom["title"]) or DEFAULT_TITLE
        for copy_index in range(1, copy_count + 1):
            index = 1
            for name in names:
                items.append(
                    CurvedOrderItem(
                        order_no=order_no,
                        detail_id=detail_id,
                        department=department,
                        font_option=font_option,
                        text=name,
                        text_type="name",
                        quantity_index=index,
                        copy_index=copy_index,
                    )
                )
                index += 1
            if title:
                items.append(
                    CurvedOrderItem(
                        order_no=order_no,
                        detail_id=detail_id,
                        department=department,
                        font_option=font_option,
                        text=title,
                        text_type="title",
                        quantity_index=index,
                        copy_index=copy_index,
                    )
                )
    return items


def group_items(items: Iterable[CurvedOrderItem]) -> List[CurvedOrderGroup]:
    grouped: "OrderedDict[str, List[CurvedOrderItem]]" = OrderedDict()
    for item in items:
        key = item.order_no + "\u001f" + item.detail_id + "\u001f" + str(item.copy_index)
        grouped.setdefault(key, []).append(item)
    return [
        CurvedOrderGroup(order_no=items[0].order_no, items=items, group_key=group_key)
        for group_key, items in grouped.items()
        if items
    ]


def load_font_map(mark_report: Path) -> Dict[str, Dict[str, str]]:
    report = json.loads(mark_report.read_text(encoding="utf-8-sig"))
    result: Dict[str, Dict[str, str]] = {}
    for entry in report.get("entries", []):
        if entry.get("status") != "ok":
            continue
        font_option = str(entry.get("font_option", ""))
        result[font_option] = {
            "font_name": str(entry.get("font_name", "")),
            "path_name": str(entry.get("path_name", "")),
            "bounds_name": str(entry.get("bounds_name", "")),
            "baseline_ratio": entry.get("baseline_ratio", {}),
            "bounds_shape_ratio": entry.get("bounds_shape_ratio", {}),
        }
    return result


def build_task(
    font_report: Path,
    output_ai: Path,
    groups: List[CurvedOrderGroup],
    columns: int,
    keep_title_frames: bool = False,
    keep_name_frames: bool = False,
    layout_overrides: Mapping[str, object] | None = None,
    color_mode: str = "CMYK",
    preview_png: Path | None = None,
    preview_dpi: int = 300,
) -> Dict[str, object]:
    if not groups:
        raise ValueError("No renderable orders")
    layout = {
        "columns": max(columns, 1),
        "margin_mm": 8.0,
        "gap_mm": 18.0,
        "item_gap_mm": 7.0,
        "order_label_height_mm": 7.0,
        "order_label_font_size_pt": 13.0,
        "name_width_mm": 16.0,
        "name_height_mm": 5.0,
        "title_width_mm": 40.0,
        "title_height_mm": 7.0,
        "keep_title_frames": keep_title_frames,
        "keep_name_frames": keep_name_frames,
    }
    for key in ("name_width_mm", "name_height_mm", "title_width_mm", "title_height_mm"):
        value = _positive_number((layout_overrides or {}).get(key))
        if value:
            layout[key] = value
    return {
        "type": "jjmb_202509_curved",
        "font_map": load_font_map(font_report),
        "output_ai": str(output_ai),
        "groups": [group.to_json_dict() for group in groups],
        "layout": layout,
        "fit": {
            "min_font_size_pt": 4.0,
            "max_font_size_pt": 80.0,
            "padding_mm": 0.2,
        },
        "output": {
            "format": "ai",
            "compatibility": "Illustrator 8",
            "color_mode": _normalize_color_mode(color_mode),
            "outline_text": True,
            "pathfinder_merge": True,
            "preview_png_path": str(preview_png) if preview_png else "",
            "preview_dpi": max(int(preview_dpi), 1) if preview_png else 0,
        },
        "debug": {
            "report_path": str(output_ai.with_suffix(".debug.json")),
        },
    }


def write_json(path: Path, payload: Dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


class CurvedRenderIntegrityError(RuntimeError):
    def __init__(self, message: str, *, code: str):
        super().__init__(message)
        self.code = code


def render_with_integrity_gate(
    *,
    output_ai: Path,
    task_options: Mapping[str, object],
    task_dir: Path,
    quality_dir: Path,
    quality_report: Path,
    visible: bool,
    bridge_factory: Callable[..., object] = IllustratorBridge,
) -> Path:
    """Render candidates until two saved-AI previews have identical pixels."""

    script = Path(__file__).resolve().parents[1] / "scripts" / "illustrator" / "render_202509_curved.jsx"
    candidates: List[Dict[str, object]] = []
    attempts: List[Dict[str, object]] = []
    accepted: Dict[str, object] | None = None

    for attempt in range(1, 4):
        candidate_ai = quality_dir / f"candidate-{attempt}.ai"
        candidate_preview = quality_dir / f"candidate-{attempt}.png"
        candidate_task = task_dir / f"render-task-quality-{attempt}.json"
        task = build_task(
            output_ai=candidate_ai,
            preview_png=candidate_preview,
            **task_options,
        )
        write_json(candidate_task, task)
        candidate: Dict[str, object] = {
            "attempt": attempt,
            "ai": candidate_ai,
            "preview": candidate_preview,
            "task": candidate_task,
        }
        candidates.append(candidate)
        try:
            bridge = bridge_factory(visible=visible)
            bridge.render(script, candidate_task)
        except Exception as exc:
            attempts.append(
                {
                    "attempt": attempt,
                    "error": str(exc),
                    "stage": "candidate_render",
                }
            )
            write_json(
                quality_report,
                {
                    "status": "failed",
                    "attempts": attempts,
                    "candidates": _integrity_candidate_report(candidates),
                },
            )
            raise
        for previous in candidates[:-1]:
            try:
                comparison = compare_png_previews(Path(previous["preview"]), candidate_preview)
            except RenderIntegrityError as exc:
                attempts.append(
                    {
                        "left_attempt": int(previous["attempt"]),
                        "right_attempt": attempt,
                        "error": str(exc),
                    }
                )
                write_json(
                    quality_report,
                    {
                        "status": "failed",
                        "attempts": attempts,
                        "candidates": _integrity_candidate_report(candidates),
                    },
                )
                raise CurvedRenderIntegrityError(str(exc), code="render_integrity_preview_invalid") from exc
            attempts.append(
                {
                    "left_attempt": int(previous["attempt"]),
                    "right_attempt": attempt,
                    **comparison.to_json_dict(),
                }
            )
            if comparison.equal:
                accepted = candidate
                break
        if accepted:
            break

    write_json(
        quality_report,
        {
            "status": "passed" if accepted else "failed",
            "attempts": attempts,
            "candidates": _integrity_candidate_report(candidates),
            "accepted_output": str(output_ai) if accepted else "",
        },
    )
    if not accepted:
        raise CurvedRenderIntegrityError(
            "渲染完整性校验失败：连续三次生成的 AI 预览不一致，已保留候选文件供人工复核。",
            code="render_integrity_mismatch",
        )
    Path(accepted["ai"]).replace(output_ai)
    return Path(accepted["task"])


def _integrity_candidate_report(candidates: Iterable[Mapping[str, object]]) -> List[Dict[str, object]]:
    return [
        {
            "attempt": int(candidate["attempt"]),
            "ai": str(candidate["ai"]),
            "preview": str(candidate["preview"]),
            "task": str(candidate["task"]),
        }
        for candidate in candidates
    ]


def _row_quantity(row: Mapping[str, str]) -> int:
    value = get_field(dict(row), *QUANTITY_FIELDS)
    if not value:
        raise ValueError("支持多姓名定制的订单数量不能为空，且必须是正整数。")
    try:
        number = Decimal(value)
    except (InvalidOperation, ValueError):
        raise ValueError("支持多姓名定制的订单数量必须是正整数。") from None
    if not number.is_finite() or number != number.to_integral_value() or number < 1:
        raise ValueError("支持多姓名定制的订单数量必须是正整数。")
    return int(number)


def _positive_number(value: object) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return number if number > 0 else 0.0


def _normalize_color_mode(value: object) -> str:
    text = str(value or "").strip().upper()
    return text if text in {"CMYK", "RGB"} else "CMYK"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render JJMB202509231236046265 curved title orders")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--font-report", required=True)
    parser.add_argument("--output-ai", required=True)
    parser.add_argument("--columns", type=int, default=5)
    parser.add_argument("--keep-title-frames", action="store_true")
    parser.add_argument("--keep-name-frames", action="store_true")
    parser.add_argument("--color-mode", choices=["CMYK", "RGB"], default="CMYK")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--visible", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_ai = Path(args.output_ai).resolve()
    output_ai.parent.mkdir(parents=True, exist_ok=True)
    try:
        rows = read_xlsx_rows(Path(args.xlsx).resolve())
        items = parse_items(rows)
        groups = group_items(items)
        task_file = output_ai.parent / "render-tasks" / "jjmb-202509-curved-render-task.json"
        task_options = {
            "font_report": Path(args.font_report).resolve(),
            "groups": groups,
            "columns": args.columns,
            "keep_title_frames": args.keep_title_frames,
            "keep_name_frames": args.keep_name_frames,
            "color_mode": args.color_mode,
        }
        print(f"render task: groups={len(groups)}, items={len(items)}")
        if args.dry_run:
            write_json(task_file, build_task(output_ai=output_ai, **task_options))
            print(task_file)
            return 0
        task_file = render_with_integrity_gate(
            output_ai=output_ai,
            task_options=task_options,
            task_dir=task_file.parent,
            quality_dir=output_ai.parent / "render-integrity",
            quality_report=output_ai.with_suffix(".render-integrity.json"),
            visible=args.visible,
        )
        print(task_file)
    except (CurvedRenderIntegrityError, IllustratorBridgeError, ValueError) as exc:
        print(f"render failed: {exc}")
        return 2
    print(f"rendered: {output_ai}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
