import hashlib
import json
import shutil
import subprocess
import struct
import zipfile
import zlib
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import Workbook

from src.renderer.illustrator_bridge import IllustratorBridgeError
from src.service import v2_order_render
from src.service import v2_order_io
from src.service import v2_order_render_support
from src.service.http_server import RenderRequestHandler
from src.service.local_client import LocalClientError, LocalDrawFlowClient
from src.service.template_registry import TemplateRegistry
from src.service.v2_template_api import V2TemplateApi
from src.service.v2_template_boundary import V2_RENDER_PIPELINE, V2_TEMPLATE_TYPE
from src.service.v2_template_contract import V2_CONTRACT_SCHEMA, V2_CONTRACT_VERSION
from src.service.v2_template_store import V2TemplateStore
from src.service.v2_order_plan import build_v2_order_units
from src.service.v2_order_render_support import stats


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _v2_payload(
    template_id: str,
    template_digest: str,
    *,
    with_styles: bool = False,
    with_design_dimensions: bool = False,
) -> dict:
    config = {
        "$schema": V2_CONTRACT_SCHEMA,
        "schema_version": V2_CONTRACT_VERSION,
        "template": {"template_id": template_id, "name": "V2 Order Demo"},
        "outputs": [
            {
                "key": "Output_main",
                "display_name": "主效果图",
                "font": {
                    "field": "font",
                    "options": [
                        {
                            "key": "F1",
                            "content_preset": "direct_text",
                            "font_dependencies": [],
                            "slots": [
                                {
                                    "key": "slot_name",
                                    "source_field": "name",
                                    "required": True,
                                    "preset": "direct_text",
                                    "font_dependencies": [],
                                }
                            ],
                        }
                    ],
                },
            }
        ],
        "field_bindings": {"font": "字体", "name": "定制信息"},
        "option_mappings": [
            {
                "field": "font",
                "source_value": "F1",
                "target": "F1",
                "output": "Output_main",
                "group": "font",
            }
        ],
        "audit": {
            "scan_version": "scan-1",
            "template_sha256": template_digest,
            "config_version": 1,
        },
    }
    if with_styles:
        output = config["outputs"][0]
        output["style"] = {
            "field": "style",
            "options": [
                {"key": "style1", "dimensions": {"mode": "style", "width_mm": 50, "height_mm": 30}},
                {"key": "style2", "dimensions": {"mode": "style", "width_mm": 200, "height_mm": 50}},
            ],
        }
        config["field_bindings"]["style"] = "尺寸"
        config["option_mappings"].extend(
            [
                {"field": "style", "source_value": "S", "target": "style1", "output": "Output_main", "group": "style"},
                {"field": "style", "source_value": "M", "target": "style2", "output": "Output_main", "group": "style"},
            ]
        )
    scan = {
        "$schema": "custom-renderer/v2-template-scan",
        "scan_protocol_version": 1,
        "evidence": {
            "template_sha256": template_digest,
            "scan_protocol_version": 1,
            "illustrator_version": "29.0",
            "scanned_at": "2026-08-12T00:00:00Z",
            "object_path_digest": "c" * 64,
        },
        "template": {"path": "Template"},
        "outputs": [
            {
                "key": "Output_main",
                "path": "Template/Output_main",
                "styles": (
                    [
                        {"key": "style1", "path": "Template/Output_main/Style/style1"},
                        {"key": "style2", "path": "Template/Output_main/Style/style2"},
                    ]
                    if with_styles
                    else []
                ),
                "designs": [],
                "fonts": [
                    {
                        "key": "F1",
                        "path": "Template/Output_main/Font/F1",
                        "slots": [
                            {
                                "key": "slot_name",
                                "path": "Template/Output_main/Font/F1/slot_name",
                            }
                        ],
                        "anchors": [],
                        "tails": [],
                    }
                ],
            }
        ],
        "dependencies": {"fonts": []},
        "issues": [],
        "blocked": False,
    }
    if with_design_dimensions:
        output = config["outputs"][0]
        output["design"] = {"field": "design", "options": [{"key": "Design03", "slots": []}]}
        config["field_bindings"]["design"] = "design"
        config["option_mappings"].append(
            {"field": "design", "source_value": "D3", "target": "Design03", "output": "Output_main", "group": "design"}
        )
        scan["outputs"][0]["designs"] = [
            {
                "key": "Design03",
                "path": "Template/Output_main/Design/Design03",
                "dimensions": {"width_mm": 155.035, "height_mm": 56.652},
                "slots": [],
                "anchors": [],
                "tails": [],
                "assets": [],
            }
        ]
    return {"config": config, "scan": scan}


def _bundle(
    path: Path,
    template_id: str,
    template_bytes: bytes,
    *,
    with_styles: bool = False,
    with_design_dimensions: bool = False,
    config_updates: dict | None = None,
    manifest_version: str = "v0001",
) -> None:
    digest = _sha256(template_bytes)
    payload = _v2_payload(
        template_id,
        digest,
        with_styles=with_styles,
        with_design_dimensions=with_design_dimensions,
    )
    if config_updates:
        payload["config"].update(config_updates)
    manifest = {
        "version": manifest_version,
        "config_sha256": "a" * 64,
        "scan_sha256": "b" * 64,
        "assets": [
            {
                "file_name": "template.ai",
                "role": "template",
                "path": "assets/template.ai",
                "extension": ".ai",
                "sha256": digest,
            }
        ],
    }
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(manifest))
        archive.writestr("metadata.json", json.dumps({"template_id": template_id, "name": "V2 Order Demo"}))
        archive.writestr("config.json", json.dumps(payload["config"], ensure_ascii=False))
        archive.writestr("scan.json", json.dumps(payload["scan"], ensure_ascii=False))
        archive.writestr("assets/template.ai", template_bytes)


def _write_order(path: Path, *, department: str = "", manufacturer: str = "") -> None:
    workbook = Workbook()
    sheet = workbook.active
    headers = ["订单号", "订单明细号", "产品名称", "字体颜色", "字体", "定制信息"]
    row = ["ORDER-1", "LINE-1", "Bracelet", "Gold", "F1", "Tom&Jerry"]
    if department:
        headers.append("生产部门")
        row.append(department)
    if manufacturer:
        headers.append("厂家")
        row.append(manufacturer)
    sheet.append(headers)
    sheet.append(row)
    workbook.save(path)


class V2PublishedCentral:
    def __init__(self, bundle_path: Path) -> None:
        self.bundle_path = bundle_path

    def get_manifest(self, template_id):
        raise LocalClientError("legacy missing", code="central_http_404")

    def get_v2_versions(self, template_id):
        return {
            "template_id": template_id,
            "publication": {"status": "active", "current_version": "v0001"},
            "versions": [{"version": "v0001"}],
        }

    def download_v2_version_bundle_to_file(self, template_id, version, target_path):
        assert (template_id, version) == ("V2ORDER001", "v0001")
        Path(target_path).write_bytes(self.bundle_path.read_bytes())
        return {"ok": True}


class CapturingRenderer:
    def __init__(self) -> None:
        self.calls = []
        self.compose_calls = []
        self.color_frame_calls = []
        self.png_master_calls = []

    def render(self, render_task, **kwargs):
        self.calls.append({"render_task": deepcopy(render_task), **kwargs})
        Path(kwargs["output_ai"]).parent.mkdir(parents=True, exist_ok=True)
        Path(kwargs["output_ai"]).write_bytes(b"rendered-ai")
        if kwargs.get("preview_png"):
            Path(kwargs["preview_png"]).write_bytes(_png_bytes())
        Path(kwargs["layout_warning_file"]).write_text('{"warnings": []}', encoding="utf-8")
        return str(kwargs["output_ai"])

    def compose_order_column(self, **kwargs):
        self.compose_calls.append(kwargs)
        Path(kwargs["output_ai"]).parent.mkdir(parents=True, exist_ok=True)
        Path(kwargs["output_ai"]).write_bytes(b"composed-ai")
        return str(kwargs["output_ai"])

    def compose_color_frames(self, **kwargs):
        self.color_frame_calls.append(kwargs)
        Path(kwargs["output_ai"]).parent.mkdir(parents=True, exist_ok=True)
        Path(kwargs["output_ai"]).write_bytes(b"color-master-ai")
        return str(kwargs["output_ai"])

    def compose_png_master_pages(self, **kwargs):
        self.png_master_calls.append(kwargs)
        output_ai = Path(kwargs["output_ai"])
        paths = _numbered_paths(output_ai, int(kwargs.get("page_count") or 1))
        for path in paths:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"png-master-ai")
        return str(output_ai)


class FailingRenderer(CapturingRenderer):
    def render(self, render_task, **kwargs):
        raise IllustratorBridgeError("Illustrator 自动化服务暂时不可用（HRESULT -2146959355）")


def test_v2_public_payload_cannot_override_version_snapshot_fields(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path)
    renderer = CapturingRenderer()
    service = v2_order_render.V2OrderRenderService(V2PublishedCentral(bundle_path), tmp_path / "local", renderer, [])

    record = service.render({
        "template_id": "V2ORDER001",
        "order_file": str(order_path),
        "dry_run": True,
        "_fixed_template_version": "v9999",
        "_fixed_template_sha256": "0" * 64,
    })

    assert record["status"] == "completed"
    assert record["request"]["template_version"] == "v0001"
    assert set(record["request"]) == {"template_id", "template_version", "order_file", "sheet_name", "dry_run"}
    assert renderer.calls == []


def test_v2_fixed_snapshot_rejects_a_bundle_with_another_version(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai", manifest_version="v0002")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path)
    service = v2_order_render.V2OrderRenderService(V2PublishedCentral(bundle_path), tmp_path / "local", CapturingRenderer(), [])

    record = service.render_fixed_snapshot(
        {"template_id": "V2ORDER001", "order_file": str(order_path), "dry_run": True},
        version="v0001",
        template_sha256=_sha256(b"template-ai"),
        config_sha256="a" * 64,
        scan_sha256="b" * 64,
    )

    assert (record["status"], record["error_code"]) == ("failed", "v2_template_version_unavailable")


@pytest.mark.parametrize("field", ("version", "template_sha256", "config_sha256", "scan_sha256"))
def test_v2_fixed_snapshot_rejects_missing_fixed_snapshot_fields(tmp_path, field):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path)
    snapshot = {
        "version": "v0001",
        "template_sha256": _sha256(b"template-ai"),
        "config_sha256": "a" * 64,
        "scan_sha256": "b" * 64,
    }
    snapshot[field] = ""
    service = v2_order_render.V2OrderRenderService(V2PublishedCentral(bundle_path), tmp_path / "local", CapturingRenderer(), [])

    with pytest.raises(v2_order_render.V2OrderRenderError) as caught:
        service.render_fixed_snapshot(
            {"template_id": "V2ORDER001", "order_file": str(order_path), "dry_run": True},
            **snapshot,
        )

    assert caught.value.code == "v2_template_snapshot_invalid"


def test_v2_fixed_snapshot_rejects_changed_config_or_scan(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path)
    with zipfile.ZipFile(bundle_path) as archive:
        config_sha = _sha256(archive.read("config.json"))
        scan_sha = _sha256(archive.read("scan.json"))
    service = v2_order_render.V2OrderRenderService(V2PublishedCentral(bundle_path), tmp_path / "local", CapturingRenderer(), [])
    payload = {"template_id": "V2ORDER001", "order_file": str(order_path), "dry_run": True}

    changed_config = service.render_fixed_snapshot(
        payload,
        version="v0001",
        template_sha256=_sha256(b"template-ai"),
        config_sha256="0" * 64,
        scan_sha256=scan_sha,
    )
    changed_scan = service.render_fixed_snapshot(
        payload,
        version="v0001",
        template_sha256=_sha256(b"template-ai"),
        config_sha256=config_sha,
        scan_sha256="0" * 64,
    )

    assert (changed_config["status"], changed_config["error_code"]) == ("failed", "v2_template_config_invalid")
    assert (changed_scan["status"], changed_scan["error_code"]) == ("failed", "v2_template_config_invalid")
    assert {"_fixed_template_sha256", "_fixed_config_sha256", "_fixed_scan_sha256"} <= set(changed_config["request"])


def _png_bytes() -> bytes:
    raw = b"\x00\x00\x00\x00\x00"
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0))
        + _png_chunk(b"IDAT", zlib.compress(raw))
        + _png_chunk(b"IEND", b"")
    )


def _png_chunk(kind: bytes, payload: bytes) -> bytes:
    return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)


def _numbered_paths(path: Path, page_count: int) -> list[Path]:
    if page_count <= 1:
        return [path]
    return [path.with_name(f"{path.stem}-{index:02d}{path.suffix}") for index in range(1, page_count + 1)]


def test_v2_formal_output_rejects_missing_department_before_template_renderer(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    with pytest.raises(LocalClientError) as exc_info:
        client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert exc_info.value.code == "v2_public_output_metadata_missing"
    assert "生产部门" in str(exc_info.value)
    assert renderer.calls == []
    job = client.jobs.list_recent(1)[0]
    assert job["status"] == "failed"
    assert job["error_code"] == "v2_public_output_metadata_missing"


def test_v2_order_preflight_failure_is_business_safe(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["字体", "定制信息"])
    sheet.append(["F2", "Tom&Jerry"])
    workbook.save(order_path)
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=CapturingRenderer(),
        font_dirs=[],
    )

    with pytest.raises(LocalClientError) as exc_info:
        client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert exc_info.value.code == "v2_order_preflight_failed"
    message = str(exc_info.value)
    assert "F2" in message
    assert "Traceback" not in message
    assert "C:" not in message
    assert "$." not in message


def test_v2_order_render_logs_technical_message_without_exposing_it_to_jobs(tmp_path, monkeypatch):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path, department="K")
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=FailingRenderer(),
        font_dirs=[],
    )
    logged_messages = []
    monkeypatch.setattr(
        v2_order_render.LOGGER,
        "error",
        lambda message, *args: logged_messages.append(message % args),
    )

    with pytest.raises(LocalClientError) as exc_info:
        client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert exc_info.value.code == "v2_order_render_failed"
    assert not exc_info.value.technical_message
    job = client.jobs.list_recent(1)[0]
    assert job["status"] == "failed"
    assert job["error_code"] == "v2_order_render_failed"
    assert job["failure_scope"] == "system"
    assert "technical_message" not in job
    assert "_technical_failure" not in job
    assert len(logged_messages) == 1
    assert "HRESULT -2146959355" in logged_messages[0]


def test_v2_download_disk_failure_is_saved_as_a_system_scope(tmp_path):
    class DiskFullCentral(V2PublishedCentral):
        def download_v2_version_bundle_to_file(self, template_id, version, target_path):
            raise OSError("disk full")

    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai")
    order_path = tmp_path / "order.xlsx"
    _write_order(order_path, department="K")
    client = LocalDrawFlowClient(DiskFullCentral(bundle_path), tmp_path / "local", v2_renderer=CapturingRenderer(), font_dirs=[])

    with pytest.raises(LocalClientError) as exc_info:
        client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert exc_info.value.code == "v2_template_version_unavailable"
    job = client.jobs.list_recent(1)[0]
    assert (job["status"], job["failure_scope"]) == ("failed", "system")


def test_v2_order_file_disk_failure_is_a_system_scope(monkeypatch):
    monkeypatch.setattr(v2_order_io, "read_xlsx_rows", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("disk full")))

    with pytest.raises(v2_order_render_support.V2OrderRenderError) as exc_info:
        v2_order_render_support.read_order_rows({"order_file": "orders.xlsx"})

    assert (exc_info.value.code, exc_info.value.failure_scope) == ("v2_order_file_unreadable", "system")


def test_v2_order_stats_counts_orders_times_outputs_once():
    result = stats(
        [{"Name": "A"}, {"Name": "B"}],
        {"outputs": [{"key": "front"}, {"key": "back"}]},
        dry_run=False,
    )

    assert result["orders"] == 2
    assert result["outputs"] == 2
    assert result["items"] == 4


def test_v2_multi_name_quantity_expands_full_text_per_copy():
    config = _v2_payload("V2ORDER001", "a" * 64)["config"]
    config["field_bindings"] = {"font": "字体", "name": "定制信息", "quantity": "数量"}
    config["multi_name_customization"] = {"enabled": True}
    preflight = {"preflight_rows": [{"row": 1, "outputs": [{"output": "Output_main", "font": "F1"}]}]}

    units = build_v2_order_units(
        config,
        {"outputs": [{"key": "Output_main"}]},
        [{"字体": "F1", "定制信息": "Alice|Bob", "数量": 3}],
        preflight,
    )

    assert [unit.values["name"] for unit in units] == ["Alice|Bob", "Alice|Bob", "Alice|Bob"]
    assert [unit.quantity_index for unit in units] == [1, 2, 3]


def test_v2_single_content_quantity_does_not_expand_without_switch():
    config = _v2_payload("V2ORDER001", "a" * 64)["config"]
    config["field_bindings"] = {"font": "字体", "name": "定制信息", "quantity": "数量"}
    preflight = {"preflight_rows": [{"row": 1, "outputs": [{"output": "Output_main", "font": "F1"}]}]}

    units = build_v2_order_units(
        config,
        {"outputs": [{"key": "Output_main"}]},
        [{"字体": "F1", "定制信息": "Alice|Bob", "数量": 3}],
        preflight,
    )

    assert len(units) == 1
    assert units[0].values["name"] == "Alice|Bob"


def test_v2_single_name_template_splits_newline_names_into_one_order_column(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(
        bundle_path,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={
            "render_mode": "single_customization",
            "field_bindings": {
                "order_no": "订单号",
                "detail_id": "订单明细号",
                "department": "生产部门",
                "product_name": "产品名称",
                "font": "字体",
                "style": "尺寸",
                "name": "定制信息",
                "quantity": "数量",
                "color": "字体颜色",
            }
        },
    )
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["订单号", "订单明细号", "生产部门", "产品名称", "字体", "尺寸", "定制信息", "数量", "字体颜色"])
    sheet.append(["ORDER-K", "LINE-K", "K", "Bracelet", "F1", "M", "Alice\nBob\nCarol", 3, "White"])
    workbook.save(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    record = client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert record["status"] == "completed"
    assert record["stats"]["items"] == 3
    assert record["outputs"]["single_order_files"][0]["name"] == "ORDER-K.ai"
    assert record["outputs"]["single_order_files"][0]["item_count"] == 3
    assert [call["values"]["name"] for call in renderer.calls[:3]] == ["Alice", "Bob", "Carol"]
    assert [call["selections"]["Output_main"]["style"] for call in renderer.calls[:3]] == [
        "style2",
        "style2",
        "style2",
    ]
    assert len(renderer.compose_calls) == 2
    assert len(renderer.color_frame_calls) == 1
    single_order_compose = renderer.compose_calls[0]
    assert single_order_compose["input_order_nos"] == ["ORDER-K", "ORDER-K", "ORDER-K"]
    assert single_order_compose["label_lines"] == ["ORDER-K", "\u767d\u8272"]
    color_component_compose = renderer.compose_calls[1]
    assert color_component_compose["input_order_nos"] == ["ORDER-K", "ORDER-K", "ORDER-K"]
    assert color_component_compose["label_lines"] == []
    assert renderer.color_frame_calls[0]["inputs"][0]["order_nos"] == ["ORDER-K"]


def test_v2_single_tail_name_slot_splits_newlines_without_splitting_spaces():
    config = {"field_bindings": {"name": "Name"}}
    render_task = {
        "outputs": [
            {
                "key": "Output_main",
                "actions": [
                    {
                        "type": "replace_slot_text",
                        "group": "font",
                        "option_key": "F2",
                        "source_field": "name",
                        "source_part_index": 0,
                        "tail_paths": ["Template/Output_main/Font/F2/tail_name_first_a"],
                        "tails": [
                            {"position": "first", "glyph_mode": "opentype_alternate"},
                            {"position": "last", "glyph_mode": "opentype_alternate"},
                        ],
                    }
                ],
            }
        ]
    }
    preflight = {
        "preflight_rows": [
            {
                "row": 1,
                "outputs": [{"output": "Output_main", "font": "F2"}],
            }
        ]
    }

    units = build_v2_order_units(
        config,
        render_task,
        [{"Name": "Mary Jane\r\nAna Maria\nEve Adams\r"}],
        preflight,
    )

    assert [unit.values["name"] for unit in units] == ["Mary Jane", "Ana Maria", "Eve Adams"]
    assert [unit.quantity_index for unit in units] == [1, 2, 3]
    assert [unit.quantity for unit in units] == [3, 3, 3]
def test_v2_order_render_expands_x53_purchase_quantity_for_both_customization_modes(tmp_path):
    names = [
        "Samantha", "Kathi", "Sam", "Candice", "Bethany",
        "Becca", "Brooke", "Jess", "Kayleigh", "Linda",
    ]
    headers = ["内部订单号", "订单明细id", "购买数量", "生产部门", "产品中文名称", "字体", "尺寸", "定制信息", "字体颜色"]
    common_bindings = {
        "order_no": "内部订单号",
        "detail_id": "订单明细id",
        "department": "生产部门",
        "product_name": "产品中文名称",
        "font": "字体",
        "style": "尺寸",
        "name": "定制信息",
        "quantity": "购买数量",
        "color": "字体颜色",
    }

    single_bundle = tmp_path / "single.zip"
    _bundle(
        single_bundle,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={"render_mode": "single_customization", "field_bindings": common_bindings},
    )
    single_orders = tmp_path / "x53-single.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    sheet.append(["4104846437", "2574703", 2, "K", "草编包", "F1", "M", "NASA Mom", "蓝色"])
    sheet.append(["4102818044", "2574704", 10, "K", "草编包", "F1", "M", "\n".join(names), "黑色"])
    workbook.save(single_orders)
    single_renderer = CapturingRenderer()
    single_client = LocalDrawFlowClient(V2PublishedCentral(single_bundle), tmp_path / "single", v2_renderer=single_renderer, font_dirs=[])

    single_record = single_client.render({"template_id": "V2ORDER001", "order_file": str(single_orders)})

    assert single_record["status"] == "completed"
    assert single_record["stats"]["items"] == 12
    assert [call["values"]["name"] for call in single_renderer.calls] == ["NASA Mom", "NASA Mom", *names]

    multi_bundle = tmp_path / "multi.zip"
    _bundle(
        multi_bundle,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={"render_mode": "multi_customization", "field_bindings": common_bindings},
    )
    multi_orders = tmp_path / "x53-multi.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    group = "\n".join(names)
    sheet.append(["4102818044", "2574704", 10, "K", "草编包", "F1", "M", group, "黑色"])
    workbook.save(multi_orders)
    multi_renderer = CapturingRenderer()
    multi_client = LocalDrawFlowClient(V2PublishedCentral(multi_bundle), tmp_path / "multi", v2_renderer=multi_renderer, font_dirs=[])

    multi_record = multi_client.render({"template_id": "V2ORDER001", "order_file": str(multi_orders)})

    assert multi_record["status"] == "completed"
    assert multi_record["stats"]["items"] == 10
    assert [call["values"]["name"] for call in multi_renderer.calls] == [group] * 10


def test_v2_order_render_preserves_technical_illustrator_error_in_job_and_gateway_result(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(
        bundle_path,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={
            "field_bindings": {
                "order_no": "订单号",
                "detail_id": "订单明细号",
                "department": "生产部门",
                "product_name": "产品名称",
                "font": "字体",
                "style": "尺寸",
                "name": "定制信息",
                "quantity": "数量",
                "color": "字体颜色",
            }
        },
    )
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["订单号", "订单明细号", "生产部门", "产品名称", "字体", "尺寸", "定制信息", "数量", "字体颜色"])
    sheet.append(["ORDER-K", "LINE-K", "K", "Bracelet", "F1", "M", "Alice", 1, "White"])
    workbook.save(order_path)

    class FailingOrderComposer(CapturingRenderer):
        def compose_order_column(self, **kwargs):
            raise IllustratorBridgeError("Illustrator JSX failed: frames[index].createOutline is not a function")

    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=FailingOrderComposer(),
        font_dirs=[],
    )

    with pytest.raises(LocalClientError) as exc_info:
        client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert exc_info.value.code == "v2_order_render_failed"
    assert "createOutline is not a function" in exc_info.value.technical_message
    job = client.jobs.list_recent(1)[0]
    assert "createOutline is not a function" in job["technical_error"]


def test_v2_newline_names_do_not_split_when_selected_template_has_multiple_name_slots():
    config = {"field_bindings": {"name": "Name"}}
    render_task = {
        "outputs": [
            {
                "key": "Output_main",
                "actions": [
                    {
                        "type": "replace_slot_text",
                        "group": "font",
                        "option_key": "F1",
                        "source_field": "name",
                        "source_part_index": 0,
                        "tail_paths": [],
                        "tails": [],
                    },
                    {
                        "type": "replace_slot_text",
                        "group": "design",
                        "option_key": "Design01",
                        "source_field": "name",
                        "source_part_index": 0,
                        "tail_paths": [],
                        "tails": [],
                    },
                ],
            }
        ]
    }
    preflight = {
        "preflight_rows": [
            {
                "row": 1,
                "outputs": [{"output": "Output_main", "font": "F1", "design": "Design01"}],
            }
        ]
    }

    units = build_v2_order_units(config, render_task, [{"Name": "Alice\nBob"}], preflight)

    assert len(units) == 1
    assert units[0].values["name"] == "Alice\nBob"


def test_v2_department_single_order_combines_duplicate_order_with_independent_styles(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(bundle_path, "V2ORDER001", b"template-ai", with_styles=True)
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["订单号", "订单明细号", "生产部门", "产品名称", "字体颜色", "字体", "尺寸", "定制信息"])
    sheet.append(["ORDER1", "LINE1", "T", "Bracelet", "Gold", "F1", "S", "Alice"])
    sheet.append(["ORDER1", "LINE2", "T", "Bracelet", "Gold", "F1", "M", "Bob"])
    workbook.save(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    record = client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert record["status"] == "completed"
    assert record["outputs"]["single_order_files"][0]["name"] == "ORDER1.ai"
    assert record["outputs"]["single_order_files"][0]["item_count"] == 2
    assert len(renderer.compose_calls) == 2
    assert len(renderer.color_frame_calls) == 1
    assert renderer.compose_calls[0]["input_order_nos"] == ["ORDER1", "ORDER1"]
    assert renderer.compose_calls[0]["label_lines"] == ["ORDER1", "\u91d1\u8272"]
    assert renderer.color_frame_calls[0]["inputs"][0]["order_nos"] == ["ORDER1"]
    rendered_styles = [call["selections"]["Output_main"]["style"] for call in renderer.calls]
    # The two reusable components feed both the single-order file and color summary.
    assert rendered_styles == ["style1", "style2"]
    with zipfile.ZipFile(record["outputs"]["primary_output"]) as archive:
        names = archive.namelist()
        assert "single-orders/ORDER1.ai" in names
        assert any(name.startswith("summary/") and name.endswith(".ai") for name in names)


def test_v2_h_department_outputs_png_single_graphics_and_master_pages(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(
        bundle_path,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={
            "field_bindings": {
                "order_no": "order_no",
                "detail_id": "detail_id",
                "department": "department",
                "product_name": "product_name",
                "font": "font",
                "style": "style",
                "name": "name",
                "color": "color",
            }
        },
    )
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_no", "detail_id", "department", "product_name", "font", "style", "name", "color"])
    sheet.append(["ORDER-H", "LINE-H", "H", "Charm", "F1", "S", "Alice", "White"])
    workbook.save(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    record = client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert record["status"] == "completed"
    assert record["outputs"]["graphic_files"][0]["name"] == "ORDER-H.png"
    assert record["outputs"]["summary_files"][0]["name"].endswith("580mm-master.ai")
    assert renderer.calls[0]["preview_dpi"] == 300
    assert len(renderer.png_master_calls) == 1
    compose_call = renderer.png_master_calls[0]
    assert compose_call["frame_width_mm"] == 580.0
    assert compose_call["frame_height_mm"] == 2000.0
    assert compose_call["items"][0]["graphic_width_mm"] == 50.0
    assert compose_call["items"][0]["graphic_height_mm"] == 30.0
    assert compose_call["items"][0]["width_mm"] == 50.0
    assert compose_call["items"][0]["height_mm"] == 30.0
    assert compose_call["items"][0]["label_embedded"] is False
    with zipfile.ZipFile(record["outputs"]["primary_output"]) as archive:
        names = archive.namelist()
        assert "single-graphics/ORDER-H.png" in names
        assert any(name.startswith("summary/") and name.endswith("580mm-master.ai") for name in names)


def test_v2_h_department_uses_selected_design_dimensions_without_style(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(
        bundle_path,
        "V2ORDER001",
        b"template-ai",
        with_design_dimensions=True,
        config_updates={
            "field_bindings": {
                "order_no": "order_no",
                "detail_id": "detail_id",
                "department": "department",
                "product_name": "product_name",
                "font": "font",
                "design": "design",
                "name": "name",
                "color": "color",
            }
        },
    )
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_no", "detail_id", "department", "product_name", "font", "design", "name", "color"])
    sheet.append(["ORDER-H-DESIGN", "LINE-H-DESIGN", "H", "Charm", "F1", "D3", "Alice", "White"])
    workbook.save(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    record = client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert record["status"] == "completed"
    assert len(renderer.png_master_calls) == 1
    item = renderer.png_master_calls[0]["items"][0]
    assert item["graphic_width_mm"] == 155.035
    assert item["graphic_height_mm"] == 56.652
    assert item["width_mm"] == 155.035
    assert item["height_mm"] == 56.652


def test_v2_w120_department_keeps_png_single_graphics_without_master(tmp_path):
    bundle_path = tmp_path / "published.zip"
    _bundle(
        bundle_path,
        "V2ORDER001",
        b"template-ai",
        with_styles=True,
        config_updates={
            "field_bindings": {
                "order_no": "order_no",
                "detail_id": "detail_id",
                "department": "department",
                "manufacturer": "manufacturer",
                "product_name": "product_name",
                "font": "font",
                "style": "style",
                "name": "name",
                "color": "color",
            }
        },
    )
    order_path = tmp_path / "order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_no", "detail_id", "department", "manufacturer", "product_name", "font", "style", "name", "color"])
    sheet.append(["ORDER-W", "LINE-W", "W", "MY-W120", "Charm", "F1", "M", "Bob", "White"])
    workbook.save(order_path)
    renderer = CapturingRenderer()
    client = LocalDrawFlowClient(
        V2PublishedCentral(bundle_path),
        tmp_path / "local",
        v2_renderer=renderer,
        font_dirs=[],
    )

    record = client.render({"template_id": "V2ORDER001", "order_file": str(order_path)})

    assert record["status"] == "completed"
    assert record["outputs"]["graphic_files"][0]["name"] == "ORDER-W.png"
    assert record["outputs"]["summary_files"] == []
    assert renderer.png_master_calls == []


def test_compose_v2_order_column_javascript_parses_in_node(tmp_path):
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js is unavailable")
    source_file = tmp_path / "compose-v2.js"
    source_file.write_text(
        Path("scripts/illustrator/compose_v2_order_column.jsx").read_text(encoding="utf-8").replace("#target illustrator", "", 1),
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            node,
            "-e",
            "const fs = require('fs'); new Function(fs.readFileSync(process.argv[1], 'utf8'));",
            str(source_file),
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr


def test_compose_v2_order_column_runs_without_native_json_parser(tmp_path):
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js is unavailable")
    task = {
        "type": "compose_v2_order_column",
        "component_contract_file": "out.warnings.json",
        "inputs": [{"path": "input.ai", "component_contract_file": "input.warnings.json"}],
        "output_ai": "out.ai",
        "gap_mm": 8,
        "compatibility": "Illustrator 8",
    }
    source = Path("scripts/illustrator/compose_v2_order_column.jsx").read_text(encoding="utf-8").replace(
        "#target illustrator",
        "",
        1,
    )
    script = f"""
const NativeJSON = JSON;
const source = {json.dumps(source)};
const taskText = NativeJSON.stringify({json.dumps(task)});
const componentContractText = NativeJSON.stringify({{component_contract_version: 1, component_frames: [{{frame_bounds: [10, 40, 60, 0], artwork_bounds_after: [10, 40, 60, 0], tracked_slots: []}}]}});
const folder = {{ exists: true, parent: null, create: () => true }};
let savedAs = '';
global.$ = {{ getenv: () => 'task.json' }};
global.File = function(path) {{
  return {{
    fsName: path,
    exists: path === 'task.json' || path === 'input.ai' || path === 'input.warnings.json',
    parent: folder,
    encoding: '',
    open: () => true,
    read: () => path === 'input.warnings.json' ? componentContractText : taskText,
    close: () => undefined,
    write: () => undefined,
    remove: () => undefined
  }};
}};
global.UserInteractionLevel = {{ DONTDISPLAYALERTS: 0 }};
global.DocumentColorSpace = {{ RGB: 1 }};
global.ElementPlacement = {{ PLACEATEND: 1 }};
global.SaveOptions = {{ DONOTSAVECHANGES: 0 }};
global.Compatibility = {{ ILLUSTRATOR8: 8, ILLUSTRATOR15: 15 }};
global.IllustratorSaveOptions = function() {{}};
  function attach(parent, childNode) {{
    childNode.parent = parent;
    parent.pageItems.push(childNode);
  }}
  function detach(childNode) {{
    if (!childNode.parent || !childNode.parent.pageItems) return;
    const index = childNode.parent.pageItems.indexOf(childNode);
    if (index >= 0) childNode.parent.pageItems.splice(index, 1);
  }}
  function union(boundsList) {{
    return boundsList.reduce((result, bounds) => [
      Math.min(result[0], bounds[0]),
      Math.max(result[1], bounds[1]),
      Math.max(result[2], bounds[2]),
      Math.min(result[3], bounds[3])
    ]);
  }}
  function item(name, bounds) {{
    let box = bounds.slice();
    return {{
      typename: 'GroupItem',
    name,
    hidden: false,
      pageItems: [],
      duplicate: function(targetLayer) {{
        const copy = item(this.name, this.visibleBounds);
        attach(targetLayer, copy);
        return copy;
      }},
      move: function(target) {{
        detach(this);
        attach(target, this);
      }},
      translate: function(dx, dy) {{
        if (this.pageItems.length) {{
          this.pageItems.forEach(child => child.translate(dx, dy));
          return;
        }}
        box = [box[0] + dx, box[1] + dy, box[2] + dx, box[3] + dy];
      }},
      get visibleBounds() {{
        return this.pageItems.length ? union(this.pageItems.map(child => child.visibleBounds)) : box.slice();
      }},
      get geometricBounds() {{ return box.slice(); }}
    }};
  }}
  const sourceLayer = {{ visible: true, pageItems: [item('component', [10, 40, 60, 0])] }};
  const sourceDoc = {{ layers: [sourceLayer], close: () => undefined }};
  const outputLayer = {{ name: 'Layer 1', pageItems: [] }};
  outputLayer.groupItems = {{ add: () => {{
    const group = item('', [0, 0, 0, 0]);
    attach(outputLayer, group);
    return group;
  }} }};
  const outputDoc = {{
    layers: [outputLayer],
    artboards: [{{ artboardRect: [] }}],
    textFrames: [],
    saveAs: file => {{ savedAs = file.fsName; }},
    close: () => undefined
  }};
global.app = {{
  userInteractionLevel: 0,
  documents: {{ add: () => outputDoc }},
  open: () => sourceDoc,
  executeMenuCommand: () => undefined
}};
global.JSON = undefined;
new Function(source)();
global.JSON = NativeJSON;
  console.log(NativeJSON.stringify({{
    savedAs,
    copied: outputLayer.pageItems.map(item => item.name),
    childNames: outputLayer.pageItems[0].pageItems.map(item => item.name),
    artboard: outputDoc.artboards[0].artboardRect
  }}));
"""

    runner = tmp_path / "compose-v2-order-column-runner.js"
    runner.write_text(script, encoding="utf-8")
    result = subprocess.run([node, str(runner)], capture_output=True, text=True, check=False)

    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["savedAs"] == "out.ai"
    assert payload["copied"] == ["ORDER_PACK_BLOCK_0"]
    assert payload["childNames"] == ["ORDER_PACK_ITEM_0_0"]
    assert payload["artboard"] == [0, 0, 50, -40]


def test_template_list_includes_active_v2_publication_without_legacy_registry(tmp_path):
    template_bytes = b"template-ai"
    digest = _sha256(template_bytes)
    payload = _v2_payload("V2ORDER001", digest)
    store = V2TemplateStore(tmp_path / "v2")
    api = V2TemplateApi(store)
    store.save_draft(
        "V2ORDER001",
        metadata={"template_id": "V2ORDER001", "name": "V2 Order Demo"},
        config=payload["config"],
        scan=payload["scan"],
        assets=[
            {
                "filename": "template.ai",
                "role": "template",
                "source_path": _write_template_asset(tmp_path, template_bytes),
                "extension": ".ai",
                "sha256": digest,
            }
        ],
    )
    store.publish_draft("V2ORDER001")
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json")
    handler.v2_template_api = api

    templates = handler._template_list_payloads()

    assert [item["template_id"] for item in handler.registry.list_templates()] == []
    v2 = next(item for item in templates if item["template_id"] == "V2ORDER001")
    assert v2["template_type"] == V2_TEMPLATE_TYPE
    assert v2["pipeline"] == V2_RENDER_PIPELINE
    assert v2["current_version"] == "v0001"
    assert v2["rule_check"]["renderable"] is True


def _write_template_asset(tmp_path: Path, payload: bytes) -> Path:
    path = tmp_path / "template.ai"
    path.write_bytes(payload)
    return path
