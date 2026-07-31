import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

import src.service.http_server as http_server
from src.jjmb_config_grouped_main import build_grouped_task
from src.render_task import RenderTaskError
from src.service.job_store import JobStore
from src.service.http_server import ExclusiveThreadingHTTPServer, RenderRequestHandler
from src.service.llm_rule_parser import LlmRuleParser
from src.service.template_rule_compiler import compile_local_rule_ast
from src.service.render_service import (
    RenderService,
    RenderServiceError,
    _configured_font_options,
    _department_output_settings_for_rows,
    _design_font_options,
    _merge_202508_template_config,
    _missing_202508_font_configs,
    _plan_png_master_pages,
)
from src.service.template_registry import TemplateRegistry
from src.service.template_inspector import TemplateInspector
from src.service.template_onboarding import TemplateOnboardingStore
from src.service.template_publication import TemplatePublicationService
from src.service.department_output import resolve_department_output


def test_http_server_uses_exclusive_windows_port(monkeypatch):
    calls = []

    class FakeSocket:
        def setsockopt(self, *args):
            calls.append(args)

    server = object.__new__(ExclusiveThreadingHTTPServer)
    server.socket = FakeSocket()
    monkeypatch.setattr(http_server.os, "name", "nt")
    monkeypatch.setattr(http_server.socket, "SO_EXCLUSIVEADDRUSE", 12345, raising=False)
    monkeypatch.setattr(http_server.ThreadingHTTPServer, "server_bind", lambda self: calls.append("bind"))

    server.server_bind()

    assert server.allow_reuse_address is False
    assert calls == [(http_server.socket.SOL_SOCKET, 12345, 1), "bind"]


def test_http_render_rejects_concurrent_request():
    errors = []

    class BusyHandler(RenderRequestHandler):
        path = "/api/render"
        render_lock = http_server.threading.Lock()

        def _send_error(self, status, message):
            errors.append((status, message))

    handler = object.__new__(BusyHandler)
    BusyHandler.render_lock.acquire()
    try:
        handler.do_POST()
    finally:
        BusyHandler.render_lock.release()

    assert errors == [(http_server.HTTPStatus.CONFLICT, "DrawFlow 正在处理另一项出图任务，请稍后再试")]


def test_http_render_releases_lock_after_failure():
    errors = []

    class FailingService:
        def submit(self, _payload):
            raise RuntimeError("render failed")

    class FailingHandler(RenderRequestHandler):
        path = "/api/render"
        render_lock = http_server.threading.Lock()
        service = FailingService()

        def _read_render_payload(self):
            return {"template_id": "T1"}

        def _send_error(self, status, message):
            errors.append((status, message))

    object.__new__(FailingHandler).do_POST()

    assert errors == [(http_server.HTTPStatus.BAD_REQUEST, "render failed")]
    assert FailingHandler.render_lock.acquire(blocking=False)
    FailingHandler.render_lock.release()


def write_order_xlsx(
    path: Path,
    *,
    template_id: str = "JJMB202508261001394920",
    department: str = "",
    manufacturer: str = "",
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "订单备注",
            "店名",
            "内部订单号",
            "SKU",
            "产品中文名称",
            "购买数量",
            "生产部门",
            "模板",
            "订单明细id",
            "SPU",
            "字体",
            "定制信息",
            "字体颜色",
            "设计",
            "厂家",
        ]
    )
    sheet.append(
        [
            "",
            "Q-206",
            "ORDER1",
            "SKU1",
            "平纹方形皮质首饰盒",
            "1",
            department,
            template_id,
            "DETAIL1",
            "SPU1",
            "F7",
            "Meg",
            "Gold",
            "Design 3",
            manufacturer,
        ]
    )
    workbook.save(path)


def write_templates_config(path: Path) -> None:
    fake_ai = path.parent / "fake-template.ai"
    fake_ai.write_text("fake ai", encoding="utf-8")
    path.write_text(
        json.dumps(
            {
                "version": 1,
                "templates": [
                    {
                        "template_id": "JJMB202508261001394920",
                        "name": "测试模板",
                        "template_type": "pure_text_color_design",
                        "pipeline": "jjmb_202508",
                        "status": "active",
                        "template_ai": str(fake_ai),
                        "default_columns": 3,
                        "default_hide_boxes": True,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def write_curved_order_xlsx(path: Path, rows: list[tuple[str, str, str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["内部订单号", "订单明细id", "生产部门", "产品中文名称", "字体颜色", "模板", "定制信息"])
    for order_no, detail_id, department, color in rows:
        sheet.append(
            [
                order_no,
                detail_id,
                department,
                "圣诞曲线标题挂件",
                color,
                "JJMB202509231236046265",
                "Font Options:F1\nTitle:Family\nName:1. Kai",
            ]
        )
    workbook.save(path)


def write_curved_templates_config(path: Path) -> None:
    fake_ai = path.parent / "curved-template.ai"
    font_report = path.parent / "curved-font-report.json"
    rules_path = path.parent / "curved.rules.json"
    fake_ai.write_text("fake ai", encoding="utf-8")
    font_report.write_text(
        json.dumps(
            {
                "entries": [
                    {
                        "status": "ok",
                        "font_option": "F1",
                        "font_name": "Test Font",
                        "baseline_ratio": {},
                        "bounds_shape_ratio": {},
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    rules_path.write_text(
        json.dumps(
            {
                "mode": "annotated_ai",
                "capabilities": ["text_on_curve"],
                "slots": [{"name": "Title", "type": "text_on_curve"}],
            }
        ),
        encoding="utf-8",
    )
    path.write_text(
        json.dumps(
            {
                "version": 1,
                "templates": [
                    {
                        "template_id": "JJMB202509231236046265",
                        "name": "曲线标题测试模板",
                        "template_type": "curved_title_text",
                        "pipeline": "jjmb_202509_curved",
                        "status": "active",
                        "template_ai": str(fake_ai),
                        "template_config": str(font_report),
                        "template_rules_config": str(rules_path),
                        "default_columns": 3,
                        "default_hide_boxes": True,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def test_service_dry_run_creates_job_and_render_task(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path)

    service = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    )

    record = service.submit(
        {
            "template_id": "JJMB202508261001394920",
            "order_file": str(order_path),
            "dry_run": True,
            "output_name": "result.ai",
        }
    )

    assert record["status"] == "completed"
    assert record["stats"]["groups"] == 1
    assert record["stats"]["items"] == 1
    assert record["stats"]["dry_run"] is True
    assert "primary_output" not in record["outputs"]
    assert "delivery_plan" in record["outputs"]
    task_path = Path(record["outputs"]["render_task"])
    assert task_path.exists()
    assert json.loads(task_path.read_text(encoding="utf-8"))["type"] == "render_batch"
    task_path = Path(record["outputs"]["render_task_files"][0])
    task = json.loads(task_path.read_text(encoding="utf-8"))
    assert task["groups"][0]["items"][0]["production_label_lines"] == ["ORDER1", "金色"]
    assert task["output"]["color_mode"] == "CMYK"
    assert task["output"]["outline_text"] is True
    assert task["output"]["pathfinder_merge"] is True


def test_service_routes_h_to_per_graphic_pngs_and_cropped_master_png(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department="H")
    workbook = load_workbook(order_path)
    sheet = workbook.active
    second = [cell.value for cell in sheet[2]]
    second[8] = "DETAIL2"
    second[11] = "Amy"
    second[12] = "Black"
    sheet.append(second)
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert "primary_output" not in record["outputs"]
    assert [item["name"] for item in record["outputs"]["graphic_files"]] == ["ORDER1-1.png", "ORDER1-2.png"]
    assert record["outputs"]["delivery_plan"][-1]["path"].endswith("-H-580mm-master.png")
    assert [member["arcname"] for member in record["outputs"]["bundle_plan"]] == [
        "single-graphics/ORDER1-1.png",
        "single-graphics/ORDER1-2.png",
        f"summary/{Path(record['outputs']['delivery_plan'][-1]['path']).name}",
    ]
    assert "output_manifest" in record["outputs"]
    graphic_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "single-graphic-tasks" in str(path)
    )
    master_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "single-graphic-tasks" not in str(path) and path.endswith("render-task-001.json")
    )
    graphic_task = json.loads(graphic_task_path.read_text(encoding="utf-8"))
    master_task = json.loads(master_task_path.read_text(encoding="utf-8"))
    assert graphic_task["output"]["format"] == "png"
    assert graphic_task["output"]["color_mode"] == "CMYK"
    assert graphic_task["groups"][0]["items"][0]["apply_color_to_artwork"] is True
    assert graphic_task["groups"][0]["items"][0]["production_label_lines"] == ["ORDER1"]
    assert master_task["output"]["format"] == "png"
    assert master_task["output"]["fixed_canvas_mm"] == {"width_mm": 580.0, "height_mm": 2000.0}
    assert master_task["output"]["crop_master_height"] is True
    assert master_task["output"]["dpi"] == 300


def test_service_dry_run_ignores_template_text_output_flags(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0]["outline_text"] = False
    config["templates"][0]["pathfinder_merge"] = False
    config_path.write_text(json.dumps(config), encoding="utf-8")

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task_files"][0]).read_text(encoding="utf-8"))
    assert task["output"]["outline_text"] is True
    assert task["output"]["pathfinder_merge"] is True


def test_service_dry_run_uses_department_output_policy_for_task(monkeypatch, tmp_path):
    from dataclasses import replace

    import src.service.render_service as render_service_module

    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department="PW")

    def fake_resolve_department_output(department, manufacturer="", **kwargs):
        rule = resolve_department_output(department, manufacturer)
        if str(department).strip() == "PW":
            return replace(rule, pathfinder_merge=False)
        return rule

    monkeypatch.setattr(render_service_module, "resolve_department_output", fake_resolve_department_output)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    tasks = [
        json.loads(Path(path).read_text(encoding="utf-8"))
        for path in record["outputs"]["render_task_files"]
        if str(path).endswith(".json")
    ]

    assert any(task.get("output", {}).get("outline_text") is True for task in tasks)
    assert any(task.get("output", {}).get("pathfinder_merge") is False for task in tasks)


def test_service_routes_t_to_one_ai_with_color_frame_artboards(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department="T")
    workbook = load_workbook(order_path)
    sheet = workbook.active
    second = [cell.value for cell in sheet[2]]
    second[2] = "ORDER2"
    second[8] = "DETAIL2"
    second[11] = "Beth"
    second[12] = "Silver"
    sheet.append(second)
    third = [cell.value for cell in sheet[2]]
    third[8] = "DETAIL3"
    third[11] = "Mia"
    third[12] = "Gold"
    sheet.append(third)
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert record["outputs"]["delivery_plan"][0]["path"].endswith("-T.ai")
    master_task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    component_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "-color-" in Path(task_path).name
    )
    compose_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "compose-color-frames" in Path(task_path).name
    )
    component_task = json.loads(component_task_path.read_text(encoding="utf-8"))
    compose_task = json.loads(compose_task_path.read_text(encoding="utf-8"))
    assert master_task["type"] == "render_batch"
    assert len(master_task["tasks"]) == len(record["outputs"]["render_task_files"])
    assert [item["name"] for item in record["outputs"]["single_order_files"]] == ["ORDER1.ai", "ORDER2.ai"]
    assert record["outputs"]["single_order_files"][0]["item_count"] == 2
    assert record["outputs"]["single_order_files"][0]["detail_ids"] == ["DETAIL1", "DETAIL3"]
    assert record["outputs"]["single_order_files"][1]["item_count"] == 1
    assert [item["arcname"] for item in record["outputs"]["bundle_plan"][:2]] == [
        "single-orders/ORDER1.ai",
        "single-orders/ORDER2.ai",
    ]
    single_order_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "single-order-tasks" in str(task_path)
        and json.loads(Path(task_path).read_text(encoding="utf-8"))["groups"][0]["order_no"] == "ORDER1"
    )
    single_order_task = json.loads(single_order_task_path.read_text(encoding="utf-8"))
    assert len(single_order_task["groups"]) == 1
    assert len(single_order_task["groups"][0]["items"]) == 2
    assert component_task["output"]["compatibility"] == "CS5"
    assert component_task["output"]["intermediate_component"] is True
    assert component_task["output"]["fixed_canvas_mm"] == {}
    assert component_task["layout"]["pack_order_blocks"] is True
    assert component_task["layout"]["suppress_labels"] is True
    assert component_task["layout"]["master_packing"]["target_width_mm"] == 580.0
    assert component_task["layout"]["master_packing"]["algorithm"] == "adaptive_column_grid"
    assert component_task["groups"][0]["items"][0]["production_label_lines"] == ["ORDER1"]
    assert compose_task["type"] == "compose_color_frames"
    assert compose_task["master_packing"]["target_width_mm"] == 580.0
    assert compose_task["master_packing"]["component_suppress_labels"] is True
    assert compose_task["show_color_header"] is True
    assert compose_task["show_color_frame_boundary"] is True
    assert [frame["color_option"] for frame in compose_task["inputs"]] == ["金色", "银色"]
    assert compose_task["inputs"][0]["order_nos"] == ["ORDER1"]


def test_curved_template_reuses_shared_department_output_pipeline(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "curved-orders.xlsx"
    write_curved_templates_config(config_path)
    write_curved_order_xlsx(
        order_path,
        [
            ("CURVED-T", "DETAIL-T", "T", "Red"),
            ("CURVED-D", "DETAIL-D", "Dept_D", "Black"),
        ],
    )

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202509231236046265", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert [item["name"] for item in record["outputs"]["single_order_files"]] == ["CURVED-T.ai", "CURVED-D.ai"]
    assert [item["department"] for item in record["outputs"]["delivery_plan"]] == ["T"]
    assert all(member["arcname"] != "manifest.json" for member in record["outputs"]["bundle_plan"])
    assert "output_manifest" in record["outputs"]
    batch_task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert batch_task["type"] == "render_batch"

    color_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "-color-" in Path(path).name
    )
    d_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "single-order-tasks" in str(path)
        and json.loads(Path(path).read_text(encoding="utf-8"))["groups"][0]["order_no"] == "CURVED-D"
    )
    color_task = json.loads(color_task_path.read_text(encoding="utf-8"))
    d_task = json.loads(d_task_path.read_text(encoding="utf-8"))
    assert color_task["output"]["fixed_canvas_mm"] == {}
    assert color_task["layout"]["pack_order_blocks"] is True
    assert color_task["layout"]["suppress_labels"] is True
    assert color_task["layout"]["master_packing"]["target_width_mm"] == 580.0
    assert color_task["layout"]["master_packing"]["algorithm"] == "adaptive_column_grid"
    assert color_task["groups"][0]["production_label_lines"] == ["CURVED-T"]
    assert d_task["groups"][0]["production_label_lines"] == ["CURVED-D", "圣诞曲线标题挂件"]


def test_curved_zw_keeps_its_existing_single_ai_delivery(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "curved-zw-orders.xlsx"
    write_curved_templates_config(config_path)
    write_curved_order_xlsx(order_path, [("CURVED-ZW", "DETAIL-ZW", "ZW", "Gold")])

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202509231236046265", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert record["outputs"]["output_ai"].endswith("JJMB202509231236046265-3col.ai")
    assert "output_bundle" not in record["outputs"]
    task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert task["type"] == "jjmb_202509_curved"
    assert task["groups"][0]["production_label_lines"] == ["CURVED-ZW"]


@pytest.mark.parametrize(
    ("department", "width_mm"),
    [
        ("K", 480.0),
        ("ZK", 450.0),
        ("FK", 450.0),
    ],
)
def test_service_routes_color_master_widths_by_department(tmp_path, department, width_mm):
    config_path = tmp_path / f"templates-{department}.json"
    order_path = tmp_path / f"orders-{department}.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department=department)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / f"jobs-{department}"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    component_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "-color-" in Path(task_path).name
    )
    component_task = json.loads(component_task_path.read_text(encoding="utf-8"))
    compose_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "compose-color-frames" in Path(task_path).name
    )
    compose_task = json.loads(compose_task_path.read_text(encoding="utf-8"))
    assert component_task["output"]["fixed_canvas_mm"] == {}
    assert component_task["layout"]["pack_order_blocks"] is True
    assert component_task["layout"]["suppress_labels"] is True
    assert compose_task["master_packing"]["target_width_mm"] == width_mm
    assert compose_task["master_packing"]["algorithm"] == "adaptive_column_grid"


@pytest.mark.parametrize("department", ["PW", "EW"])
def test_service_routes_pw_ew_to_single_orders_and_one_master(tmp_path, department):
    config_path = tmp_path / f"templates-{department}.json"
    order_path = tmp_path / f"orders-{department}.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department=department)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / f"jobs-{department}"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert [item["name"] for item in record["outputs"]["single_order_files"]] == ["ORDER1.ai"]
    assert record["outputs"]["delivery_plan"][0]["path"].endswith(f"-{department}.ai")
    arcnames = [member["arcname"] for member in record["outputs"]["bundle_plan"]]
    assert arcnames[0] == "single-orders/ORDER1.ai"
    assert arcnames[1].startswith("summary/")
    assert arcnames[1].endswith(f"-{department}.ai")
    master_task_path = next(
        Path(task_path)
        for task_path in record["outputs"]["render_task_files"]
        if "single-order-tasks" not in str(task_path)
    )
    master_task = json.loads(master_task_path.read_text(encoding="utf-8"))
    item = master_task["groups"][0]["items"][0]
    assert item["production_label_lines"] == ["ORDER1", "平纹方形皮质首饰盒"]
    assert master_task["output"]["fixed_canvas_mm"] == {}


def test_service_routes_d_department_to_single_orders_without_master(tmp_path):
    config_path = tmp_path / "templates-d.json"
    order_path = tmp_path / "orders-d.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department="Dept_D")

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-d"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert record["stats"]["deliveries"] == 0
    assert record["outputs"]["delivery_plan"] == []
    assert [item["name"] for item in record["outputs"]["single_order_files"]] == ["ORDER1.ai"]
    assert [member["arcname"] for member in record["outputs"]["bundle_plan"]] == [
        "single-orders/ORDER1.ai",
    ]
    assert "output_manifest" in record["outputs"]
    single_task_path = Path(record["outputs"]["render_task_files"][0])
    single_task = json.loads(single_task_path.read_text(encoding="utf-8"))
    item = single_task["groups"][0]["items"][0]
    assert item["production_label_lines"] == ["ORDER1", "平纹方形皮质首饰盒"]


def test_202603_pipeline_rejects_non_graphic_department_output(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    config = json.loads(config_path.read_text(encoding="utf-8"))
    structure_path = tmp_path / "template.config.json"
    structure_path.write_text(
        json.dumps(
            {
                "font_options": {"F7": {"type": "text"}},
                "style_options": {"Style1": {"width_mm": 80, "height_mm": 50}},
                "slots": [{"name": "text_fit_box"}],
            }
        ),
        encoding="utf-8",
    )
    config["templates"][0].update(
        {
            "pipeline": "jjmb_202603_grouped",
            "template_type": "pure_text",
            "template_id": "JJMB202603281027102517",
            "template_config": str(structure_path),
        }
    )
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    write_order_xlsx(order_path, template_id="JJMB202603281027102517", department="T")
    workbook = load_workbook(order_path)
    workbook.active["N1"] = "Style Option"
    workbook.active["N2"] = "Style 1"
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202603281027102517", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "failed"
    assert record["error_code"] == "department_output_pipeline_unsupported"


def test_service_routes_202603_h_to_exact_pngs_and_paginated_master_ai(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    structure_path = tmp_path / "template.config.json"
    rules_path = tmp_path / "template.rules.json"
    write_templates_config(config_path)
    structure_path.write_text(
        json.dumps(
            {
                "font_options": {
                    "F7": {
                        "type": "text",
                        "font_name": "TestFont",
                        "font_size_pt": 48,
                    }
                },
                "style_options": {
                    "Style1": {
                        "width_pt": 226.7716535433,
                        "height_pt": 141.7322834646,
                        "width_mm": 80,
                        "height_mm": 50,
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    rules_path.write_text(json.dumps({"font_options": ["F7"], "style_options": ["Style1"]}), encoding="utf-8")
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0].update(
        {
            "template_id": "JJMB202603281027102517",
            "template_type": "pure_text_style",
            "pipeline": "jjmb_202603_grouped",
            "template_config": str(structure_path),
            "template_rules_config": str(rules_path),
        }
    )
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    write_order_xlsx(order_path, template_id="JJMB202603281027102517", department="H")
    workbook = load_workbook(order_path)
    workbook.active["N1"] = "Style Option"
    workbook.active["N2"] = "Style 1"
    second_row = [cell.value for cell in workbook.active[2]]
    second_row[2] = "ORDER2"
    second_row[8] = "DETAIL2"
    second_row[11] = "Amy"
    workbook.active.append(second_row)
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-h-202603"),
    ).submit(
        {"template_id": "JJMB202603281027102517", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    assert [item["name"] for item in record["outputs"]["graphic_files"]] == ["ORDER1.png", "ORDER2.png"]
    assert record["outputs"]["summary_files"][0]["name"].endswith("-H-580mm-master.ai")
    assert [member["arcname"] for member in record["outputs"]["bundle_plan"]] == [
        "single-graphics/ORDER1.png",
        "single-graphics/ORDER2.png",
        f"summary/{record['outputs']['summary_files'][0]['name']}",
    ]
    assert "output_manifest" in record["outputs"]
    single_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "single-graphic-tasks" in str(path)
    )
    compose_task_path = next(
        Path(path)
        for path in record["outputs"]["render_task_files"]
        if "compose-png-master-pages" in str(path)
    )
    single_task = json.loads(single_task_path.read_text(encoding="utf-8"))
    compose_task = json.loads(compose_task_path.read_text(encoding="utf-8"))
    assert single_task["layout"]["single_graphic_exact"] is True
    assert single_task["layout"]["embed_order_label"] is True
    assert single_task["layout"]["single_graphic_label_height_mm"] == 6.0
    assert single_task["layout"]["single_graphic_label_gap_mm"] == 0.8
    assert single_task["layout"]["single_graphic_bleed_mm"] == 2.0
    assert single_task["export"]["format"] == "png"
    assert single_task["export"]["png_path"].endswith("ORDER1.png")
    assert "fixed_canvas_mm" not in single_task["export"]
    assert compose_task["type"] == "compose_png_master_pages"
    assert compose_task["compatibility"] == "CS5"
    assert compose_task["frame_width_mm"] == 580.0
    assert compose_task["frame_height_mm"] == 2000.0
    assert compose_task["items"][0]["width_mm"] == 84.0
    assert compose_task["items"][0]["height_mm"] == 60.8
    assert compose_task["items"][0]["graphic_width_mm"] == 80.0
    assert compose_task["items"][0]["graphic_height_mm"] == 50.0
    assert compose_task["items"][0]["label_embedded"] is True
    assert compose_task["items"][0]["bleed_mm"] == 2.0
    render_batches = [str(path) for path in record["outputs"]["render_batch_files"]]
    assert any("single-render-batches" in path for path in render_batches)
    assert any("compose-render-batches" in path for path in render_batches)
    single_batches = [
        json.loads(Path(path).read_text(encoding="utf-8"))
        for path in render_batches
        if "single-render-batches" in path
    ]
    compose_batches = [
        json.loads(Path(path).read_text(encoding="utf-8"))
        for path in render_batches
        if "compose-render-batches" in path
    ]
    assert [len(batch["tasks"]) for batch in single_batches] == [1, 1]
    assert [len(batch["tasks"]) for batch in compose_batches] == [1]


def test_202603_h_master_planner_paginates_without_scaling():
    rule = resolve_department_output("H")
    items = [
        {"order_no": f"ORDER-{index:03d}", "width_mm": 80, "height_mm": 50}
        for index in range(300)
    ]

    plan = _plan_png_master_pages(items, rule)

    assert plan["scale"] == 1
    assert len(plan["pages"]) > 1
    assert all(page["artboard_height_mm"] <= 2000 for page in plan["pages"])
    assert sum(page["items"] for page in plan["pages"]) == 300


def test_service_routes_w_manufacturers_to_cs5_master_ai_pngs_and_standard_ai(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path, department="W", manufacturer="MY-W196")
    workbook = load_workbook(order_path)
    sheet = workbook.active
    second = [cell.value for cell in sheet[2]]
    second[8] = "DETAIL2"
    second[11] = "Amy"
    second[12] = "Black"
    sheet.append(second)
    workbook.save(order_path)

    w196_record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-w196-cs5-master"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )
    component_path = next(
        Path(path)
        for path in w196_record["outputs"]["render_task_files"]
        if "master-component" in str(path)
    )
    compose_path = next(
        Path(path)
        for path in w196_record["outputs"]["render_task_files"]
        if "compose-color-frames" in str(path)
    )
    w196_task = json.loads(component_path.read_text(encoding="utf-8"))
    compose_task = json.loads(compose_path.read_text(encoding="utf-8"))
    assert w196_task["output"]["format"] == "ai"
    assert w196_task["output"]["compatibility"] == "CS5"
    assert w196_task["output"]["color_mode"] == "CMYK"
    assert w196_task["layout"]["pack_order_blocks"] is True
    assert w196_task["layout"]["suppress_labels"] is True
    assert w196_task["layout"]["master_packing"]["force_subitem_order_labels"] is False
    assert w196_task["layout"]["master_packing"]["component_suppress_labels"] is True
    assert len(w196_task["groups"]) == 1
    assert len(w196_task["groups"][0]["items"]) == 2
    assert [item["color_option"] for item in w196_task["groups"][0]["items"]] == ["Gold", "Black"]
    assert all(item["apply_color_to_artwork"] is True for item in w196_task["groups"][0]["items"])
    assert all(item["production_label_lines"] == ["ORDER1"] for item in w196_task["groups"][0]["items"])
    assert compose_task["compatibility"] == "CS5"
    assert compose_task["show_color_header"] is False
    assert compose_task["master_packing"]["keep_order_items_together"] is True
    assert [item["name"] for item in w196_record["outputs"]["summary_files"]] == [
        "JJMB202508261001394920-3col-W-MY-W196.ai"
    ]
    assert w196_record["outputs"]["delivery_plan"][0]["path"].endswith("-W-MY-W196.ai")
    assert w196_record["outputs"]["graphic_files"] == []
    assert "output_bundle" not in w196_record["outputs"]

    write_order_xlsx(order_path, department="ZW")
    workbook = load_workbook(order_path)
    sheet = workbook.active
    sheet["O1"] = "外协厂家代码"
    sheet["O2"] = "MY-W196"
    second = [cell.value for cell in sheet[2]]
    second[8] = "DETAIL2"
    second[11] = "Amy"
    sheet.append(second)
    workbook.save(order_path)
    zw_w196_record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-zw-w196-cs5-master"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )
    zw_component_path = next(
        Path(path)
        for path in zw_w196_record["outputs"]["render_task_files"]
        if "master-component" in str(path)
    )
    zw_task = json.loads(zw_component_path.read_text(encoding="utf-8"))
    assert zw_task["output"]["compatibility"] == "CS5"
    assert zw_task["groups"][0]["items"][0]["department"] == "ZW"
    assert all(item["apply_color_to_artwork"] is True for item in zw_task["groups"][0]["items"])
    assert zw_w196_record["outputs"]["delivery_plan"][0]["path"].endswith("-ZW-MY-W196.ai")

    write_order_xlsx(order_path, department="W", manufacturer="OTHER-W")
    standard_record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-standard"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )
    standard_task = json.loads(Path(standard_record["outputs"]["render_task_files"][0]).read_text(encoding="utf-8"))
    assert standard_task["output"]["format"] == "ai"
    assert standard_task["output"]["compatibility"] == "AI_STANDARD"
    assert standard_record["outputs"]["delivery_plan"][0]["path"].endswith("-W-ORDER1.ai")

    write_order_xlsx(order_path, department="W", manufacturer="MY-W120")
    workbook = load_workbook(order_path)
    sheet = workbook.active
    sheet["C2"] = "W-120-001"
    second = [cell.value for cell in sheet[2]]
    second[2] = "W-120-002"
    second[8] = "DETAIL2"
    second[11] = "Amy"
    sheet.append(second)
    workbook.save(order_path)

    png_record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs-png"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )
    assert "output_bundle" not in png_record["outputs"]
    assert [item["name"] for item in png_record["outputs"]["graphic_files"]] == ["W-120-001.png", "W-120-002.png"]
    assert [member["arcname"] for member in png_record["outputs"]["bundle_plan"]] == [
        "single-graphics/W-120-001.png",
        "single-graphics/W-120-002.png",
    ]
    assert "output_manifest" in png_record["outputs"]
    for task_path in png_record["outputs"]["render_task_files"]:
        task = json.loads(Path(task_path).read_text(encoding="utf-8"))
        assert task["output"]["format"] == "png"
        assert task["groups"][0]["items"][0]["production_label_lines"] == [task["groups"][0]["order_no"]]
        assert task["output"]["color_mode"] == "CMYK"
        assert task["layout"]["suppress_labels"] is False
        assert task["groups"][0]["items"][0]["apply_color_to_artwork"] is True


def test_generic_w196_applies_cs5_and_quantity_split(tmp_path):
    config_path = tmp_path / "templates.json"
    rules_path = tmp_path / "template.rules.json"
    order_path = tmp_path / "orders.xlsx"
    fake_ai = tmp_path / "template.ai"
    fake_ai.write_text("fake ai", encoding="utf-8")
    rules_path.write_text(
        json.dumps(
            {
                "status": "confirmed",
                "order_bindings": {
                    "order_no": "Order",
                    "text": "Name",
                    "quantity": "Quantity",
                },
                "slot_mappings": [{"field": "text", "slot": "Name"}],
                "multi_name_customization": {"enabled": True},
                "render_layout": {
                    "type": "name_columns",
                    "output_mode": "single_file",
                    "default": {"group_by": ["order_no"], "header_fields": ["order_no"]},
                    "manufacturer_overrides": {
                        "MY-W196": {"group_by": ["row"], "header_fields": ["order_no"]}
                    },
                },
                "output": {"color_mode": "CMYK"},
            }
        ),
        encoding="utf-8",
    )
    config_path.write_text(
        json.dumps(
            {
                "version": 1,
                "templates": [
                    {
                        "template_id": "GENERIC-W196",
                        "name": "Generic W196",
                        "template_type": "pure_text_color_design",
                        "pipeline": "generic_rules_only",
                        "status": "active",
                        "template_ai": str(fake_ai),
                        "template_rules_config": str(rules_path),
                        "default_columns": 4,
                        "default_hide_boxes": True,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Name", "Quantity", "Department", "Manufacturer"])
    sheet.append(["ORDER-W196", "Alice|Bob", 3, "ZW", "MY-W196"])
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "GENERIC-W196", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert record["status"] == "completed"
    assert task["output"]["compatibility"] == "CS5"
    assert len(task["orders"]) == 3
    assert [order["quantity_index"] for order in task["orders"]] == [1, 2, 3]
    assert [len(order["layout_members"]) for order in task["orders"]] == [1, 1, 1]


def test_generic_w120_outputs_per_graphic_png_bundle_plan(tmp_path):
    config_path = tmp_path / "templates.json"
    rules_path = tmp_path / "template.rules.json"
    order_path = tmp_path / "orders.xlsx"
    fake_ai = tmp_path / "template.ai"
    fake_ai.write_text("fake ai", encoding="utf-8")
    rules_path.write_text(
        json.dumps(
            {
                "status": "confirmed",
                "order_bindings": {
                    "order_no": "Order",
                    "text": "Name",
                    "quantity": "Quantity",
                },
                "slot_mappings": [{"field": "text", "slot": "Name"}],
                "multi_name_customization": {"enabled": True},
                "render_layout": {
                    "type": "name_columns",
                    "output_mode": "single_file",
                    "default": {"group_by": ["row"], "header_fields": ["order_no"]},
                },
                "output": {"color_mode": "CMYK"},
            }
        ),
        encoding="utf-8",
    )
    config_path.write_text(
        json.dumps(
            {
                "version": 1,
                "templates": [
                    {
                        "template_id": "GENERIC-W120",
                        "name": "Generic W120",
                        "template_type": "pure_text_color_design",
                        "pipeline": "generic_rules_only",
                        "status": "active",
                        "template_ai": str(fake_ai),
                        "template_rules_config": str(rules_path),
                        "default_columns": 4,
                        "default_hide_boxes": True,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Name", "Quantity", "Department", "Manufacturer"])
    sheet.append(["ORDER-W120", "Alice|Bob", 2, "ZW", "MY-W120"])
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "GENERIC-W120", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert record["status"] == "completed"
    assert task["output"]["format"] == "png"
    assert task["output"]["transparent_background"] is True
    assert task["render_layout"]["output_mode"] == "per_graphic"
    assert [Path(item["path"]).name for item in record["outputs"]["graphic_files"]] == [
        "ORDER-W120-1.png",
        "ORDER-W120-2.png",
    ]
    assert [member["arcname"] for member in record["outputs"]["bundle_plan"]] == [
        "single-graphics/ORDER-W120-1.png",
        "single-graphics/ORDER-W120-2.png",
    ]
    assert "output_manifest" in record["outputs"]


def test_202508_task_receives_every_configured_font_boldness_mapping(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    rules_path = tmp_path / "template.rules.json"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    rules_path.write_text(
        json.dumps(
            {
                "font_style_rules": [
                    {"font_options": ["F2", "F3", "F10", "F11", "F12"], "boldness": 0.4},
                    {"font_options": ["F5", "F6", "F7", "F8", "F9"], "boldness": 0.5},
                ]
            }
        ),
        encoding="utf-8",
    )
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0]["template_rules_config"] = str(rules_path)
    config_path.write_text(json.dumps(config), encoding="utf-8")

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task_files"][0]).read_text(encoding="utf-8"))
    assert task["font_styles"] == {
        "F2": {"boldness": 0.4},
        "F3": {"boldness": 0.4},
        "F10": {"boldness": 0.4},
        "F11": {"boldness": 0.4},
        "F12": {"boldness": 0.4},
        "F5": {"boldness": 0.5},
        "F6": {"boldness": 0.5},
        "F7": {"boldness": 0.5},
        "F8": {"boldness": 0.5},
        "F9": {"boldness": 0.5},
    }


def test_202508_task_receives_compiled_segment_color_actions(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    rules_path = tmp_path / "template.rules.json"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    rules_path.write_text(
        json.dumps(
            {
                "rule_ast": compile_local_rule_ast(
                    "Name列按 | 分隔，奇数位渲染为红色，偶数位渲染为白色"
                )
            }
        ),
        encoding="utf-8",
    )
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0]["template_rules_config"] = str(rules_path)
    config_path.write_text(json.dumps(config), encoding="utf-8")
    workbook = load_workbook(order_path)
    workbook.active["L2"] = "Alice|Bob|Cara"
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task_files"][0]).read_text(encoding="utf-8"))
    assert len(task["groups"]) == 1
    assert len(task["groups"][0]["items"]) == 1
    assert task["groups"][0]["items"][0]["text"] == "Alice|Bob|Cara"
    assert task["groups"][0]["items"][0]["text_actions"] == [
        {
            "type": "fill_color",
            "strategy": "cycle",
            "values": ["#FF0000", "#FFFFFF"],
            "selector": {"type": "segments", "delimiter": "|"},
        }
    ]


def test_202508_segment_color_rules_still_split_comma_name_lists(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    rules_path = tmp_path / "template.rules.json"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    rules_path.write_text(
        json.dumps(
            {
                "rule_ast": {
                    "$schema": "custom-renderer/template-rule-ast",
                    "version": 1,
                    "source_hash": "0" * 64,
                    "rules": [
                        {
                            "target": {"type": "text", "name": "Name"},
                            "conditions": [],
                            "selector": {"type": "segments", "delimiter": "|"},
                            "operations": [
                                {
                                    "type": "fill_color",
                                    "strategy": "cycle",
                                    "values": ["#FF0000", "#FFFFFF"],
                                }
                            ],
                        }
                    ],
                    "unresolved": [],
                }
            }
        ),
        encoding="utf-8",
    )
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0]["template_rules_config"] = str(rules_path)
    config_path.write_text(json.dumps(config), encoding="utf-8")
    workbook = load_workbook(order_path)
    workbook.active["L2"] = "Alice|Bob,Cara|Dana"
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202508261001394920", "order_file": str(order_path), "dry_run": True}
    )

    task = json.loads(Path(record["outputs"]["render_task_files"][0]).read_text(encoding="utf-8"))
    assert len(task["groups"]) == 1
    assert [item["text"] for item in task["groups"][0]["items"]] == ["Alice|Bob", "Cara|Dana"]
    assert all(
        item["text_actions"]
        == [
            {
                "type": "fill_color",
                "strategy": "cycle",
                "values": ["#FF0000", "#FFFFFF"],
                "selector": {"type": "segments", "delimiter": "|"},
            }
        ]
        for item in task["groups"][0]["items"]
    )


def test_202603_grouped_task_receives_every_configured_font_boldness_mapping(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    structure_path = tmp_path / "template.config.json"
    rules_path = tmp_path / "template.rules.json"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    structure_path.write_text("{}", encoding="utf-8")
    rules_path.write_text(
        json.dumps(
            {
                "font_style_rules": [
                    {"font_options": ["F2", "F3", "F10", "F11", "F12"], "boldness": 0.4},
                    {"font_options": ["F5", "F6", "F7", "F8", "F9"], "boldness": 0.5},
                ]
            }
        ),
        encoding="utf-8",
    )
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["templates"][0].update(
        {
            "template_id": "JJMB202603281027102517",
            "template_type": "pure_text_style",
            "pipeline": "jjmb_202603_grouped",
            "template_config": str(structure_path),
            "template_rules_config": str(rules_path),
        }
    )
    config_path.write_text(json.dumps(config), encoding="utf-8")
    workbook = load_workbook(order_path)
    workbook.active["H2"] = "JJMB202603281027102517"
    workbook.active["N1"] = "Style Option"
    workbook.active["N2"] = "Style 1"
    workbook.save(order_path)

    record = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    ).submit(
        {"template_id": "JJMB202603281027102517", "order_file": str(order_path), "dry_run": True}
    )

    assert record["status"] == "completed", record.get("error")
    task = json.loads(Path(record["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert task["font_styles"] == {
        "F2": {"boldness": 0.4},
        "F3": {"boldness": 0.4},
        "F10": {"boldness": 0.4},
        "F11": {"boldness": 0.4},
        "F12": {"boldness": 0.4},
        "F5": {"boldness": 0.5},
        "F6": {"boldness": 0.5},
        "F7": {"boldness": 0.5},
        "F8": {"boldness": 0.5},
        "F9": {"boldness": 0.5},
    }


def test_202603_grouped_design_asset_task_keeps_its_font_style_mapping(tmp_path):
    order_path = tmp_path / "orders.xlsx"
    design_asset_path = tmp_path / "F10.ai"
    write_order_xlsx(order_path)
    design_asset_path.write_text("placeholder", encoding="utf-8")
    workbook = load_workbook(order_path)
    workbook.active["H2"] = "JJMB202603281027102517"
    workbook.active["K2"] = "F10"
    workbook.active["N1"] = "Style Option"
    workbook.active["N2"] = "Style 1"
    workbook.save(order_path)

    task = build_grouped_task(
        xlsx_path=order_path,
        template_config=tmp_path / "template.config.json",
        output_ai=tmp_path / "out.ai",
        columns=4,
        allowed_font_options=["F10"],
        design_font_options=["F10"],
        design_asset_mappings={"F10": {"path": str(design_asset_path), "group": "F10"}},
        font_styles={"F10": {"boldness": 0.4}},
    )

    task_json = task.to_json_dict()
    assert task_json["groups"][0]["items"][0]["render_kind"] == "design_asset"
    assert task_json["font_styles"] == {"F10": {"boldness": 0.4}}


def test_service_rejects_draft_template_before_creating_job(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path)
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    payload["templates"][0]["status"] = "draft"
    config_path.write_text(json.dumps(payload), encoding="utf-8")
    service = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    )

    with pytest.raises(RenderServiceError, match="模板未启用") as exc_info:
        service.submit({"template_id": "JJMB202508261001394920", "order_file": str(order_path)})

    assert exc_info.value.code == "template_not_active"
    assert service.jobs.list_recent() == []


def test_scan_confirm_publish_then_render_dry_run(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = registry.save_uploaded_ai("DEMO001", "template.ai", b"fake ai")
    template = registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text_color_design",
            "status": "draft",
            "template_ai": registry.to_config_path(ai_path),
        }
    )

    class Bridge:
        def __init__(self, **kwargs):
            pass

        def render(self, script, task_path):
            task = json.loads(task_path.read_text(encoding="utf-8"))
            Path(task["output_json"]).write_text(
                json.dumps(
                    {
                        "layers": [{"name": "Template"}],
                        "items": [
                            {"type": "GroupItem", "name": "F1", "path": "Template/F1"},
                            {"type": "GroupItem", "name": "Design1", "path": "Template/Design1"},
                            {"type": "TextFrame", "name": "Name1", "path": "Template/Name1"},
                        ],
                    }
                ),
                encoding="utf-8",
            )

    store = TemplateOnboardingStore(registry.storage_dir)
    state = TemplateInspector(Bridge).scan(template, store)
    pack = state["draft"]
    pack["validation"]["unresolved_items"] = []
    pack["validation"]["sample"] = {
        "input": {"定制信息": "Meg"},
        "expected": {"Name1": "Meg"},
    }
    pack["rules"]["order_bindings"] = {"text": "定制信息"}
    pack["rules"]["text_policies"] = {"fit": "scale_to_box"}
    record = TemplatePublicationService(registry, store).confirm(
        "DEMO001", pack, change_summary="E2E"
    )
    active = record["template"]
    order_path = tmp_path / "orders.xlsx"
    write_order_xlsx(order_path, template_id=active.template_id)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": active.template_id, "order_file": str(order_path), "dry_run": True}
    )

    assert result["status"] == "completed", result
    assert active.template_rules_config and active.template_rules_config.exists()


def test_template_registry_resolves_defaults(tmp_path):
    config_path = tmp_path / "templates.json"
    write_templates_config(config_path)

    template = TemplateRegistry(config_path).get_template("JJMB202508261001394920")

    assert template.default_columns == 3
    assert template.default_hide_boxes is True
    assert template.pipeline == "jjmb_202508"


def test_template_registry_upserts_uploaded_template_and_rules(tmp_path):
    config_path = tmp_path / "templates.json"
    storage_dir = tmp_path / "templates"
    registry = TemplateRegistry(config_path, storage_dir)

    ai_path = registry.save_uploaded_ai("JJMB202607030001", "source.ai", b"fake ai")
    rule_path = registry.save_template_config(
        "JJMB202607030001",
        '{"template_id":"JJMB202607030001","defaults":{"color":"Gold"}}',
    )
    template = registry.upsert_template(
        {
            "template_id": "JJMB202607030001",
            "name": "新增模板",
            "template_type": "pure_text_color_design",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
            "template_config": registry.to_config_path(rule_path),
            "default_columns": "6",
            "default_hide_boxes": "false",
        }
    )

    assert template.template_id == "JJMB202607030001"
    assert template.template_ai.exists()
    assert template.template_config and template.template_config.exists()
    assert template.default_columns == 6
    assert template.default_hide_boxes is False
    assert template.pipeline == "jjmb_202508"


def test_register_template_uses_reference_ai_when_size_template_is_empty(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    template = handler._register_template(
        {
            "template_id": "JJMB202607030005",
            "name": "只传原始参考模板",
            "template_type": "pure_text_color_design",
            "status": "active",
        },
        {"reference_ai": [{"filename": "source.ai", "content": b"source ai"}]},
    )

    assert template.template_ai and template.template_ai.exists()
    assert template.template_ai_role == "原始参考模板"
    assert template.assets == []
    assert template.status == "draft"


def test_register_template_keeps_existing_reference_when_size_template_is_added(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    first = handler._register_template(
        {
            "template_id": "JJMB202607030015",
            "name": "先传原始参考模板",
            "template_type": "pure_text_color_design",
            "status": "active",
        },
        {"reference_ai": [{"filename": "source.ai", "content": b"source ai"}]},
    )

    updated = handler._register_template(
        {
            "template_id": first.template_id,
            "name": first.name,
            "template_type": first.template_type,
            "status": "active",
        },
        {"template_ai": [{"filename": "layout.ai", "content": b"layout ai"}]},
    )

    assert updated.template_ai and updated.template_ai.read_bytes() == b"layout ai"
    assert updated.template_ai_role == "尺寸/作图区模板"
    assert len(updated.assets) == 1
    assert updated.assets[0]["role"] == "原始参考模板"
    reference_copy = Path(updated.assets[0]["stored_path"])
    assert reference_copy.read_bytes() == b"source ai"


def test_register_template_keeps_reference_ai_as_asset_when_size_template_exists(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    template = handler._register_template(
        {
            "template_id": "JJMB202607030006",
            "name": "参考和尺寸模板",
            "template_type": "pure_text_color_design",
            "status": "active",
        },
        {
            "template_ai": [{"filename": "layout.ai", "content": b"layout ai"}],
            "reference_ai": [{"filename": "source.ai", "content": b"source ai"}],
        },
    )

    assert template.template_ai and template.template_ai.exists()
    assert template.template_ai_role == "尺寸/作图区模板"
    assert len(template.assets) == 1
    assert template.assets[0]["role"] == "原始参考模板"


def test_register_template_ignores_direct_rules_and_preserves_structure_config(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    structure_config = tmp_path / "curved-title-mark-report.json"
    structure_config.write_text('{"entries":[{"font_option":"F14","status":"ok"}]}', encoding="utf-8")
    ai_path = handler.registry.save_uploaded_ai("JJMB202607030007", "template.ai", b"ai")
    handler.registry.upsert_template(
        {
            "template_id": "JJMB202607030007",
            "name": "曲线标题模板",
            "template_type": "curved_title_text",
            "pipeline": "jjmb_202509_curved",
            "status": "draft",
            "template_ai": handler.registry.to_config_path(ai_path),
            "template_config": str(structure_config),
        }
    )

    template = handler._register_template(
        {
            "template_id": "JJMB202607030007",
            "name": "曲线标题模板",
            "template_type": "curved_title_text",
            "status": "active",
            "template_rules_json": '{"raw_text":"业务规则","font_options":["F14"],"parser":{"source":"llm"}}',
            "template_rules_text": "业务规则",
        },
        {},
    )

    assert template.template_config == structure_config
    assert template.template_rules_config is None
    assert not (handler.registry.storage_dir / template.template_id / "template.rules.json").exists()


def test_template_registry_saves_multiple_ai_assets(tmp_path):
    config_path = tmp_path / "templates.json"
    storage_dir = tmp_path / "templates"
    registry = TemplateRegistry(config_path, storage_dir)

    ai_path = registry.save_uploaded_ai("JJMB202607030002", "main.ai", b"main ai")
    assets = registry.save_uploaded_assets(
        "JJMB202607030002",
        [
            {"filename": "design-a.ai", "content": b"asset a"},
            {"filename": "design-b.ai", "content": b"asset b"},
        ],
        role="独立设计模板",
    )
    template = registry.upsert_template(
        {
            "template_id": "JJMB202607030002",
            "name": "多资产模板",
            "template_type": "pure_text_color_design",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
            "assets": assets,
        }
    )

    assert len(template.assets) == 2
    assert template.assets[0]["file_name"] == "design-a.ai"
    assert template.assets[0]["role"] == "独立设计模板"
    assert Path(template.assets[0]["stored_path"]).name == "design-a.ai"
    assert (storage_dir / "JJMB202607030002" / "assets" / "design-a.ai").exists()


def test_render_service_reads_design_font_options_from_llm_rule_shape():
    assert _design_font_options(
        {
            "design_options": {
                "design_font_options": ["F10", "F11", "F12"],
            }
        }
    ) == ["F10", "F11", "F12"]


def test_render_service_normalizes_object_option_rules():
    config = {
        "font_options": [{"id": "F1"}, {"name": "F2"}],
        "design_font_options": [{"font_option": "F10"}],
        "design_options": {
            "design_font_options": [{"id": "F11"}],
        },
    }

    assert _configured_font_options(config) == ["F1", "F2"]
    assert _design_font_options(config) == ["F10", "F11"]


def test_202508_template_config_merge_fills_missing_font_options():
    base = {
        "font_options": {
            "F1": {"font_name": "Base Font"},
        },
        "design_options": {"Design1": {"box": "base"}},
    }
    reference = {
        "font_options": {
            "F1": {"font_name": "Reference Font"},
            "F7": {"font_name": "Reference F7"},
        }
    }

    merged = _merge_202508_template_config(base, reference)

    assert merged["font_options"]["F1"]["font_name"] == "Base Font"
    assert merged["font_options"]["F7"]["font_name"] == "Reference F7"
    assert merged["design_options"] == {"Design1": {"box": "base"}}


def test_202508_missing_font_configs_reports_used_options_only_once():
    config = {"font_options": {"F1": {"font_name": "A"}}}

    assert _missing_202508_font_configs(config, ["F1", "F7", "F7", ""]) == ["F7"]


def test_template_registry_deletes_asset_registration_without_removing_files(tmp_path):
    config_path = tmp_path / "templates.json"
    storage_dir = tmp_path / "templates"
    registry = TemplateRegistry(config_path, storage_dir)

    ai_path = registry.save_uploaded_ai("JJMB202607030003", "main.ai", b"main ai")
    assets = registry.save_uploaded_assets(
        "JJMB202607030003",
        [{"filename": "design-a.ai", "content": b"asset a"}],
        role="独立设计模板",
    )
    asset_path = storage_dir / "JJMB202607030003" / "assets" / "design-a.ai"
    registry.upsert_template(
        {
            "template_id": "JJMB202607030003",
            "name": "待删除模板",
            "template_type": "pure_text_color_design",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
            "assets": assets,
        }
    )

    removed = registry.delete_template_asset("JJMB202607030003", 0)

    assert removed["file_name"] == "design-a.ai"
    assert ai_path.exists()
    assert asset_path.exists()
    assert registry.get_template("JJMB202607030003").assets == []
    assert registry.get_template("JJMB202607030003").status == "draft"


def test_template_registry_deletes_primary_ai_registration_without_removing_file(tmp_path):
    config_path = tmp_path / "templates.json"
    storage_dir = tmp_path / "templates"
    registry = TemplateRegistry(config_path, storage_dir)

    ai_path = registry.save_uploaded_ai("JJMB202607030004", "main.ai", b"main ai")
    registry.upsert_template(
        {
            "template_id": "JJMB202607030004",
            "name": "待删除主模板",
            "template_type": "pure_text_color_design",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
        }
    )

    removed = registry.delete_template_ai("JJMB202607030004")
    template = registry.get_template("JJMB202607030004")

    assert removed["role"] == "尺寸/作图区模板"
    assert ai_path.exists()
    assert template.template_ai is None
    assert template.status == "draft"


def test_template_registry_disables_and_removes_record_without_removing_files(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = registry.save_uploaded_ai("DEMO001", "main.ai", b"main ai")
    registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text",
            "status": "draft",
            "template_ai": registry.to_config_path(ai_path),
        }
    )

    assert registry.set_template_status("DEMO001", "disabled").status == "disabled"
    removed = registry.remove_template_record("DEMO001")

    assert removed["template_id"] == "DEMO001"
    assert ai_path.exists()
    with pytest.raises(KeyError):
        registry.get_template("DEMO001")


def test_confirmed_pack_is_published_to_runtime_config_and_activates(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    registry.upsert_template(
        {"template_id": "DEMO001", "name": "Demo", "template_type": "pure_text", "status": "draft"}
    )
    pack = {"template": {"template_id": "DEMO001"}, "rules": {"font_options": ["F1"]}}

    template = registry.apply_confirmed_rule_pack("DEMO001", pack, activate=True)

    assert template.status == "active"
    assert template.template_rules_config
    assert json.loads(template.template_rules_config.read_text(encoding="utf-8")) == pack


def test_template_registry_does_not_persist_template_text_output_flags(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    template = registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text",
            "outline_text": False,
            "pathfinder_merge": False,
        }
    )

    reloaded = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates").get_template("DEMO001")
    raw = json.loads((tmp_path / "templates.json").read_text(encoding="utf-8"))
    assert not hasattr(template, "outline_text")
    assert not hasattr(template, "pathfinder_merge")
    assert "outline_text" not in reloaded.to_json_dict()
    assert "pathfinder_merge" not in reloaded.to_json_dict()
    assert "outline_text" not in raw["templates"][0]
    assert "pathfinder_merge" not in raw["templates"][0]


def test_confirmed_pack_does_not_persist_template_text_output_flags(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text",
        }
    )
    pack = {
        "template": {"template_id": "DEMO001"},
        "rules": {"output": {"outline_text": True, "pathfinder_merge": False, "color_mode": "RGB"}},
    }

    template = registry.apply_confirmed_rule_pack("DEMO001", pack, activate=False)

    saved_pack = json.loads(template.template_rules_config.read_text(encoding="utf-8"))
    raw = json.loads((tmp_path / "templates.json").read_text(encoding="utf-8"))
    assert saved_pack["rules"]["output"] == {"color_mode": "RGB"}
    assert "outline_text" not in template.to_json_dict()
    assert "pathfinder_merge" not in template.to_json_dict()
    assert "outline_text" not in raw["templates"][0]
    assert "pathfinder_merge" not in raw["templates"][0]


def test_direct_template_rules_save_strips_text_output_flags(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    rules_path = registry.save_template_rules_config(
        "DEMO001",
        json.dumps(
            {
                "output": {
                    "outline_text": True,
                    "pathfinder_merge": False,
                    "color_mode": "CMYK",
                }
            }
        ),
    )

    saved = json.loads(rules_path.read_text(encoding="utf-8"))
    assert saved["output"] == {"color_mode": "CMYK"}


def test_confirmed_pack_strips_top_level_template_text_output_flags(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text",
        }
    )
    pack = {
        "template": {"template_id": "DEMO001"},
        "output": {"outline_text": True, "pathfinder_merge": False, "color_mode": "RGB"},
    }

    template = registry.apply_confirmed_rule_pack("DEMO001", pack, activate=False)

    saved_pack = json.loads(template.template_rules_config.read_text(encoding="utf-8"))
    assert saved_pack["output"] == {"color_mode": "RGB"}


def test_render_task_json_is_ascii_escaped_for_illustrator(tmp_path):
    service = object.__new__(RenderService)
    task_file = tmp_path / "render-task.json"

    service._write_render_task_json(task_file, {"product_name": "单件定制-红色挂件"})

    text = task_file.read_text(encoding="utf-8")
    assert "单件定制" not in text
    assert "\\u5355\\u4ef6\\u5b9a\\u5236" in text
    assert json.loads(text)["product_name"] == "单件定制-红色挂件"


def test_non_render_json_keeps_readable_unicode(tmp_path):
    service = object.__new__(RenderService)
    job_file = tmp_path / "job.json"

    service._write_json(job_file, {"stage": "生成 AI 文件"})

    assert "生成 AI 文件" in job_file.read_text(encoding="utf-8")


def test_shared_output_settings_reject_mixed_department_policies(monkeypatch):
    from dataclasses import replace

    import src.service.render_service as render_service_module

    def fake_resolve_department_output(department, manufacturer="", **kwargs):
        rule = resolve_department_output(department, manufacturer)
        if str(department).strip() == "PW":
            return replace(rule, pathfinder_merge=False)
        return rule

    monkeypatch.setattr(render_service_module, "resolve_department_output", fake_resolve_department_output)

    with pytest.raises(RenderServiceError, match="不同生产部门转曲/去重策略"):
        _department_output_settings_for_rows(
            [{"department": "T"}, {"department": "PW"}],
            {"output": {"color_mode": "CMYK"}},
        )


def test_destructive_template_action_requires_exact_id_confirmation():
    with pytest.raises(ValueError, match="template ID"):
        RenderRequestHandler._require_template_confirmation("DEMO001", {"confirmation": "wrong"})

    RenderRequestHandler._require_template_confirmation("DEMO001", {"confirmation": "DEMO001"})


def test_activate_template_requires_a_renderable_template(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    handler.registry.upsert_template(
        {"template_id": "BROKEN001", "name": "Broken", "template_type": "pure_text", "status": "draft"}
    )

    with pytest.raises(ValueError, match="还不能启用出图"):
        handler._activate_template("BROKEN001")


def test_template_special_rule_compile_endpoint_builds_a_reviewable_ast(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    handler.llm_parser = LlmRuleParser(api_key="", base_url="")
    handler._onboarding_store = lambda: type(
        "Onboarding",
        (),
        {"get_state": lambda _self, _template_id: {"draft": {"rules": {"text_targets": [{"name": "Name"}]}}}},
    )()
    handler.registry.upsert_template(
        {
            "template_id": "RULE001",
            "name": "Rule",
            "template_type": "pure_text",
            "pipeline": "generic_rules_only",
            "status": "draft",
        }
    )

    result = handler._compile_template_special_rules(
        {
            "template_id": "RULE001",
            "natural_text": "Name列的数据，奇数位渲染为红色，偶数位渲染为白色",
        }
    )

    assert result["errors"] == ["自然语言规则模型未配置，当前结果仅供预览，不能确认保存。"]
    assert result["summary"] == ["Name 按 | 分段，循环填充 #FF0000 / #FFFFFF"]
    assert result["compiler"] == {"source": "local", "llm_configured": False}


def test_activate_template_restores_a_renderable_draft(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = handler.registry.save_uploaded_ai("READY001", "template.ai", b"ai")
    handler.registry.upsert_template(
        {
            "template_id": "READY001",
            "name": "Ready",
            "template_type": "pure_text_color_design",
            "pipeline": "jjmb_202508",
            "status": "draft",
            "template_ai": handler.registry.to_config_path(ai_path),
        }
    )

    template = handler._activate_template("READY001")

    assert template.status == "active"


def test_activate_template_restores_a_complete_generic_draft(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = handler.registry.save_uploaded_ai("GENERICREADY001", "template.ai", b"ai")
    rule_path = handler.registry.save_template_rules_config(
        "GENERICREADY001",
        json.dumps(
            {
                "mode": "pure_text",
                "status": "confirmed",
                "font_options": ["F1"],
                "order_bindings": {"text": "Name"},
                "slot_mappings": [{"name": "Name", "field": "text", "type": "replace_text"}],
            }
        ),
    )
    handler.registry.upsert_template(
        {
            "template_id": "GENERICREADY001",
            "name": "Generic ready",
            "template_type": "pure_text",
            "status": "draft",
            "template_ai": handler.registry.to_config_path(ai_path),
            "template_rules_config": handler.registry.to_config_path(rule_path),
        }
    )

    template = handler._activate_template("GENERICREADY001")

    assert template.status == "active"


def test_active_template_upload_failure_leaves_template_in_draft(tmp_path):
    handler = object.__new__(RenderRequestHandler)
    handler.registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = handler.registry.save_uploaded_ai("SAFEUPDATE001", "template.ai", b"old scanned ai")
    handler.registry.upsert_template(
        {
            "template_id": "SAFEUPDATE001",
            "name": "Safe update",
            "template_type": "pure_text_color_design",
            "pipeline": "jjmb_202508",
            "status": "active",
            "template_ai": handler.registry.to_config_path(ai_path),
        }
    )

    with pytest.raises(ValueError, match="必须是 .ai 格式"):
        handler._register_template(
            {
                "template_id": "SAFEUPDATE001",
                "name": "Safe update",
                "template_type": "pure_text_color_design",
                "pipeline": "jjmb_202508",
            },
            {
                "template_ai": [{"filename": "replacement.ai", "content": b"new unscanned ai"}],
                "design_font_assets": [{"filename": "invalid.txt", "content": b"not an ai asset"}],
            },
        )

    template = handler.registry.get_template("SAFEUPDATE001")

    assert template.status == "draft"
    assert template.template_ai and template.template_ai.read_bytes() == b"new unscanned ai"


def test_template_registry_rejects_directory_colliding_ids(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    with pytest.raises(ValueError, match="template_id may only contain"):
        registry.upsert_template(
            {"template_id": "DEMO/001", "name": "Bad", "template_type": "pure_text"}
        )


def test_template_registry_rejects_case_only_collision_before_upload(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    registry.upsert_template({"template_id": "Demo", "name": "One", "template_type": "pure_text"})

    with pytest.raises(ValueError, match="differs only by case"):
        registry.validate_template_id("demo")


def test_template_registry_serializes_concurrent_updates(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")

    def add(template_id):
        return registry.upsert_template(
            {"template_id": template_id, "name": template_id, "template_type": "pure_text"}
        ).template_id

    with ThreadPoolExecutor(max_workers=2) as executor:
        assert set(executor.map(add, ["DEMO001", "DEMO002"])) == {"DEMO001", "DEMO002"}

    assert {item.template_id for item in registry.list_templates()} == {"DEMO001", "DEMO002"}


def test_activation_rejects_missing_ai_for_generic_pipeline(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    template = registry.upsert_template(
        {"template_id": "DEMO001", "name": "Demo", "template_type": "pure_text", "status": "draft"}
    )
    with pytest.raises(ValueError, match="AI file does not exist"):
        TemplatePublicationService._validate_activation(template, {"rules": {}})


def test_activation_rejects_registered_asset_when_file_is_missing(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = registry.save_uploaded_ai("DEMO001", "template.ai", b"ai")
    template = registry.upsert_template(
        {
            "template_id": "DEMO001",
            "name": "Demo",
            "template_type": "pure_text_color_design",
            "status": "draft",
            "template_ai": registry.to_config_path(ai_path),
            "assets": [{"file_name": "missing.ai", "stored_path": str(tmp_path / "missing.ai")}],
        }
    )

    with pytest.raises(ValueError, match="file is missing"):
        TemplatePublicationService._validate_activation(
            template,
            {"rules": {"asset_mappings": [{"option": "Design1", "asset": "missing.ai"}]}},
        )


def test_service_never_enables_diagnostic_boxes_from_request_flags(tmp_path):
    config_path = tmp_path / "templates.json"
    order_path = tmp_path / "orders.xlsx"
    write_templates_config(config_path)
    write_order_xlsx(order_path)

    service = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    )

    record = service.submit(
        {
            "template_id": "JJMB202508261001394920",
            "order_file": str(order_path),
            "dry_run": "true",
            "hide_boxes": "false",
        }
    )

    assert record["status"] == "completed"
    task_path = Path(record["outputs"]["render_task_files"][0])
    task = json.loads(task_path.read_text(encoding="utf-8"))
    assert task["layout"]["show_style_boxes"] is False


def test_grouped_pipeline_reports_template_mismatch_when_no_groups(tmp_path):
    order_path = tmp_path / "orders.xlsx"
    write_order_xlsx(order_path)

    try:
        build_grouped_task(
            xlsx_path=order_path,
            template_config=tmp_path / "template.config.json",
            output_ai=tmp_path / "out.ai",
            columns=4,
        )
    except RenderTaskError as exc:
        message = str(exc)
        assert "订单表没有解析到可渲染内容" in message
        assert "JJMB202508261001394920" in message
        assert "JJMB202603281027102517" in message
    else:
        raise AssertionError("template mismatch should fail with a readable error")


def test_service_rejects_missing_order_file(tmp_path):
    config_path = tmp_path / "templates.json"
    write_templates_config(config_path)
    service = RenderService(
        registry=TemplateRegistry(config_path),
        jobs=JobStore(tmp_path / "jobs"),
    )

    try:
        service.submit({"template_id": "JJMB202508261001394920", "order_file": ""})
    except Exception as exc:
        assert "缺少订单表格" in str(exc)
        assert getattr(exc, "code", "") == "missing_order_file"
    else:
        raise AssertionError("missing order_file should fail")


def test_job_store_lists_recent_jobs(tmp_path):
    jobs = JobStore(tmp_path / "jobs")
    first = jobs.create({"template_id": "A"})
    second = jobs.create({"template_id": "B"})

    recent = jobs.list_recent(2)

    assert [item["job_id"] for item in recent] == [second["job_id"], first["job_id"]]
