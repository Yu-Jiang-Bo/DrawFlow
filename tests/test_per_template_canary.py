from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.renderer.illustrator_bridge import IllustratorBridgeError
from src.service.multi_template_order import MultiTemplateOrderBatch, MultiTemplateOrderRow, TemplateOrderGroup
from src.service.multi_template_preflight import MultiTemplateGroupPreflight, MultiTemplatePreflightResult
from src.service.multi_template_snapshot import TemplateSnapshot
from src.service.per_template_canary import PerTemplateCanaryRenderer
from src.service.render_service import render_failure_scope


def _row(template_id: str, excel_row: int, order_no: str) -> MultiTemplateOrderRow:
    return MultiTemplateOrderRow(
        "订单",
        excel_row,
        order_no,
        template_id,
        {"订单号": order_no, "模板": template_id, "定制信息": f"Name-{order_no}"},
        (order_no, template_id, f"Name-{order_no}"),
    )


def _preflight(tmp_path: Path, template_ids: tuple[str, ...], *, status: str = "ready") -> MultiTemplatePreflightResult:
    groups = tuple(TemplateOrderGroup(template_id, (_row(template_id, index + 2, f"ORDER-{template_id}"),)) for index, template_id in enumerate(template_ids))
    headers = ("订单号", "模板", "定制信息")
    batch = MultiTemplateOrderBatch("订单", headers, tuple(row for group in groups for row in group.rows), groups, ())
    summaries = []
    snapshots = []
    for group in groups:
        workbook_path = tmp_path / "groups" / group.template_id / "orders.xlsx"
        workbook_path.parent.mkdir(parents=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "订单"
        sheet.append(headers)
        sheet.append(group.rows[0].raw_values)
        workbook.save(workbook_path)
        workbook.close()
        digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
        row = group.rows[0]
        summaries.append(MultiTemplateGroupPreflight(
            group.template_id,
            1,
            (row.excel_row,),
            (row.order_no,),
            str(workbook_path),
            digest,
            status == "ready",
            {},
            {"row_metrics": {str(row.excel_row): {"planned_output_units": 1, "variable_text_length": len(str(row.values["定制信息"]))}}},
        ))
        snapshots.append(TemplateSnapshot("legacy", group.template_id, "generic_rules_only", "v1", "a" * 64, "", "", "", "", ()))
    return MultiTemplatePreflightResult(status, "source-sha", "订单", batch, tuple(snapshots), tuple(summaries), ())


class FakeAdapter:
    def __init__(self, outcomes: dict[str, object] | None = None) -> None:
        self.outcomes = dict(outcomes or {})
        self.calls: list[tuple[str, Path, Path, bool]] = []

    def render_canary(self, group, snapshot, *, group_workbook, work_dir):
        workbook = Path(group_workbook)
        self.calls.append((group.template_id, workbook, Path(work_dir), False))
        outcome = self.outcomes.get(group.template_id, {"status": "completed"})
        if isinstance(outcome, Exception):
            raise outcome
        return {
            "job_id": f"canary-{group.template_id}",
            "job_dir": str(Path(work_dir) / "jobs" / group.template_id),
            "outputs": {"primary_output": str(Path(work_dir) / "diagnostic.zip")},
            **outcome,
        }


class WorkbookMutatingAdapter(FakeAdapter):
    def render_canary(self, group, snapshot, *, group_workbook, work_dir):
        workbook = load_workbook(group_workbook)
        try:
            workbook["订单"].cell(row=2, column=2).value = "B"
            workbook.save(group_workbook)
        finally:
            workbook.close()
        return super().render_canary(group, snapshot, group_workbook=group_workbook, work_dir=work_dir)


class RecoveryGate:
    def __init__(self, available: bool) -> None:
        self.available = available
        self.calls = 0

    def check(self) -> bool:
        self.calls += 1
        return self.available


def test_template_canary_failure_does_not_prevent_later_templates(tmp_path):
    adapter = FakeAdapter({
        "A": {"status": "failed", "failure_scope": "template", "error_code": "render_failed"},
    })

    result = PerTemplateCanaryRenderer(adapter=adapter).run(_preflight(tmp_path, ("A", "B", "C")), work_dir=tmp_path / "parent")

    assert result.status == "completed_with_errors"
    assert [(item.template_id, item.status, item.failure_scope) for item in result.groups] == [
        ("A", "canary_failed", "template"),
        ("B", "ready", ""),
        ("C", "ready", ""),
    ]
    assert [call[0] for call in adapter.calls] == ["A", "B", "C"]
    workbook = load_workbook(adapter.calls[1][1], read_only=True, data_only=True)
    try:
        rows = list(workbook["订单"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows == [("订单号", "模板", "定制信息"), ("ORDER-B", "B", "Name-ORDER-B")]
    assert result.groups[1].canary_workbook.endswith("canary\\B\\orders.xlsx")
    assert result.groups[1].child_job_id == "canary-B"


def test_canary_marks_each_template_started_before_its_adapter_call(tmp_path):
    started: list[str] = []

    class StartAwareAdapter(FakeAdapter):
        def render_canary(self, group, snapshot, *, group_workbook, work_dir):
            assert started[-1] == group.template_id
            return super().render_canary(group, snapshot, group_workbook=group_workbook, work_dir=work_dir)

    adapter = StartAwareAdapter()

    result = PerTemplateCanaryRenderer(adapter=adapter).run(
        _preflight(tmp_path, ("A", "B")),
        work_dir=tmp_path / "parent",
        on_group_started=started.append,
    )

    assert result.status == "completed"
    assert started == ["A", "B"]
    assert [call[0] for call in adapter.calls] == ["A", "B"]


def test_legacy_and_v2_formal_failure_codes_continue_without_a_handwritten_scope(tmp_path):
    for template_id, error_code in (("A", "illustrator_render_failed"), ("B", "v2_order_render_failed")):
        adapter = FakeAdapter({template_id: {"status": "failed", "error_code": error_code}})

        result = PerTemplateCanaryRenderer(adapter=adapter).run(
            _preflight(tmp_path / template_id, (template_id, "C")),
            work_dir=tmp_path / template_id / "parent",
        )

        assert result.status == "completed_with_errors"
        assert [(item.template_id, item.status) for item in result.groups] == [
            (template_id, "canary_failed"),
            ("C", "ready"),
        ]


def test_system_failure_stops_later_canaries(tmp_path):
    adapter = FakeAdapter({"A": OSError("disk unavailable")})

    result = PerTemplateCanaryRenderer(adapter=adapter).run(_preflight(tmp_path, ("A", "B", "C")), work_dir=tmp_path / "parent")

    assert result.status == "interrupted"
    assert [(item.template_id, item.status, item.failure_scope) for item in result.groups] == [("A", "interrupted", "system")]
    assert [call[0] for call in adapter.calls] == ["A"]
    assert result.error_code == "canary_runtime_unavailable"


def test_recovered_com_canary_failure_keeps_later_templates_running(tmp_path):
    recovery = RecoveryGate(True)
    adapter = FakeAdapter({
        "A": {
            "status": "failed",
            "error_code": "illustrator_render_failed",
            "error": "HRESULT -2147417851",
            "failure_scope": "system",
        },
    })

    result = PerTemplateCanaryRenderer(adapter=adapter, illustrator_recovery=recovery).run(
        _preflight(tmp_path, ("A", "B")), work_dir=tmp_path / "parent",
    )

    assert result.status == "completed_with_errors"
    assert [(item.template_id, item.status, item.failure_scope) for item in result.groups] == [
        ("A", "canary_failed", "template"),
        ("B", "ready", ""),
    ]
    assert [call[0] for call in adapter.calls] == ["A", "B"]
    assert recovery.calls == 1


def test_unrecovered_com_canary_failure_stops_later_templates(tmp_path):
    recovery = RecoveryGate(False)
    adapter = FakeAdapter({
        "A": {
            "status": "failed",
            "error_code": "illustrator_render_failed",
            "error": "HRESULT -2147417851",
            "failure_scope": "system",
        },
    })

    result = PerTemplateCanaryRenderer(adapter=adapter, illustrator_recovery=recovery).run(
        _preflight(tmp_path, ("A", "B")), work_dir=tmp_path / "parent",
    )

    assert result.status == "interrupted"
    assert [(item.template_id, item.status, item.failure_scope) for item in result.groups] == [("A", "interrupted", "system")]
    assert [call[0] for call in adapter.calls] == ["A"]
    assert recovery.calls == 1


def test_direct_recoverable_com_canary_exception_uses_fresh_session_gate(tmp_path):
    recovery = RecoveryGate(True)
    adapter = FakeAdapter({
        "A": IllustratorBridgeError("COM HRESULT -2147417851", failure_scope="system"),
    })

    result = PerTemplateCanaryRenderer(adapter=adapter, illustrator_recovery=recovery).run(
        _preflight(tmp_path, ("A", "B")), work_dir=tmp_path / "parent",
    )

    assert result.status == "completed_with_errors"
    assert [(item.template_id, item.status, item.failure_scope) for item in result.groups] == [
        ("A", "canary_failed", "template"),
        ("B", "ready", ""),
    ]
    assert [call[0] for call in adapter.calls] == ["A", "B"]
    assert recovery.calls == 1


def test_parent_session_retries_only_the_failed_canary_group(tmp_path):
    recovery_calls = []

    class RetryOnceAdapter(FakeAdapter):
        def render_canary(self, group, snapshot, *, group_workbook, work_dir):
            if group.template_id == "A" and sum(call[0] == "A" for call in self.calls) == 0:
                return super().render_canary(
                    group,
                    snapshot,
                    group_workbook=group_workbook,
                    work_dir=work_dir,
                ) | {
                    "status": "failed",
                    "error_code": "v2_order_render_failed",
                    "error": "Illustrator unavailable",
                    "failure_scope": "system",
                    "_technical_failure": "HRESULT -2147023174",
                }
            return super().render_canary(group, snapshot, group_workbook=group_workbook, work_dir=work_dir)

    adapter = RetryOnceAdapter()
    result = PerTemplateCanaryRenderer(adapter=adapter).run(
        _preflight(tmp_path, ("A", "B")),
        work_dir=tmp_path / "parent",
        recover_illustrator=lambda: recovery_calls.append("reset") or True,
    )

    assert result.status == "completed"
    assert [call[0] for call in adapter.calls] == ["A", "A", "B"]
    assert recovery_calls == ["reset"]


def test_changed_group_workbook_stops_before_starting_a_canary(tmp_path):
    preflight = _preflight(tmp_path, ("A", "B"))
    Path(preflight.groups[0].group_workbook).write_bytes(b"changed")
    adapter = FakeAdapter()

    result = PerTemplateCanaryRenderer(adapter=adapter).run(preflight, work_dir=tmp_path / "parent")

    assert result.status == "interrupted"
    assert adapter.calls == []
    assert result.groups[0].failure_scope == "system"


def test_changed_canary_workbook_is_not_marked_ready(tmp_path):
    adapter = WorkbookMutatingAdapter()

    result = PerTemplateCanaryRenderer(adapter=adapter).run(_preflight(tmp_path, ("A", "B")), work_dir=tmp_path / "parent")

    assert result.status == "interrupted"
    assert [(group.template_id, group.status, group.error_code) for group in result.groups] == [
        ("A", "interrupted", "canary_workbook_changed"),
    ]
    assert [call[0] for call in adapter.calls] == ["A"]


def test_canary_refuses_to_run_from_a_failed_static_preflight(tmp_path):
    adapter = FakeAdapter()

    result = PerTemplateCanaryRenderer(adapter=adapter).run(_preflight(tmp_path, ("A",), status="preflight_failed"), work_dir=tmp_path / "parent")

    assert result.status == "preflight_failed"
    assert result.groups == ()
    assert adapter.calls == []


def test_legacy_illustrator_failure_scope_distinguishes_template_and_unrecoverable_com():
    assert render_failure_scope(IllustratorBridgeError("JSX slot is missing", failure_scope="template")) == "template"
    assert render_failure_scope(IllustratorBridgeError("HRESULT -2146959355", failure_scope="system")) == "system"
