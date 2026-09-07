from __future__ import annotations

import hashlib
from datetime import datetime
from decimal import Decimal
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from src.jjmb_202508_main import parse_items as parse_202508_items
from src.jjmb_202509_curved_main import TEMPLATE_ID as CURVED_TEMPLATE_ID
from src.jjmb_202509_curved_main import parse_items as parse_202509_items
from src.jjmb_202509_curved_main import read_xlsx_rows as read_202509_rows
from src.jjmb_config_grouped_main import build_grouped_task
from src.jjmb_order_parser import read_xlsx_rows as read_202603_rows
from src.jjmb_template_main import TEMPLATE_ID as CONFIG_TEMPLATE_ID
from src.service.generic_rule_renderer import _read_rows as read_generic_rows
from src.service.multi_template_group_workbooks import GroupWorkbookWriter
from src.service import multi_template_excel_reader
from src.service import multi_template_order
from src.service.job_store import JobStore
from src.service.multi_template_order import (
    PARENT_JOB_STATUSES,
    TEMPLATE_GROUP_STATUSES,
    MultiTemplateIssue,
    MultiTemplateOrderBatch,
    MultiTemplateOrderParser,
    MultiTemplateOrderRow,
)
from src.service.v2_order_render_support import read_order_rows


def write_workbook(path, rows, *, sheet_name="订单"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def cache_numeric_formula_result(path, formula, result):
    temporary = path.with_name(f"{path.stem}.cached{path.suffix}")
    with ZipFile(path) as source:
        contents = {member.filename: source.read(member.filename) for member in source.infolist()}
        sheet_name = "xl/worksheets/sheet1.xml"
        formula_text = formula.removeprefix("=").encode("utf-8")
        expected = b"<f>" + formula_text + b"</f><v />"
        replacement = b"<f>" + formula_text + b"</f><v>" + str(result).encode("utf-8") + b"</v>"
        assert expected in contents[sheet_name]
        contents[sheet_name] = contents[sheet_name].replace(expected, replacement, 1)
    with ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for name, content in contents.items():
            target.writestr(name, content)
    temporary.replace(path)


def test_parser_groups_once_in_first_template_order_and_keeps_excel_rows(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["内部订单号", "模板", "数量"], ["A-1", "TEMPLATE-A", 1], ["B-1", "TEMPLATE-B", 2], ["A-2", "TEMPLATE-A", 3], [None, None, None]])

    batch = MultiTemplateOrderParser().parse(source)

    assert batch.sheet_name == "订单"
    assert batch.headers == ("内部订单号", "模板", "数量")
    assert batch.issues == ()
    assert [row.excel_row for row in batch.rows] == [2, 3, 4]
    assert [row.order_no for row in batch.rows] == ["A-1", "B-1", "A-2"]
    assert [group.template_id for group in batch.groups] == ["TEMPLATE-A", "TEMPLATE-B"]
    assert [[row.excel_row for row in group.rows] for group in batch.groups] == [[2, 4], [3]]
    assert MultiTemplateOrderBatch.from_dict(batch.to_dict()) == batch


def test_parser_opens_one_workbook_and_scans_the_selected_sheet_once(tmp_path, monkeypatch):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["订单号", "模板"], ["A-1", "TEMPLATE-A"]])
    open_calls = []
    scan_calls = 0
    original_open = multi_template_order.load_workbook
    original_parse = multi_template_excel_reader._FormulaAwareWorksheetParser.parse

    def counted_open(*args, **kwargs):
        open_calls.append((kwargs.get("read_only"), kwargs.get("data_only")))
        return original_open(*args, **kwargs)

    def counted_parse(parser):
        nonlocal scan_calls
        scan_calls += 1
        yield from original_parse(parser)

    monkeypatch.setattr(multi_template_order, "load_workbook", counted_open)
    monkeypatch.setattr(multi_template_excel_reader._FormulaAwareWorksheetParser, "parse", counted_parse)

    batch = MultiTemplateOrderParser().parse(source)

    assert batch.issues == ()
    assert open_calls == [(True, True)]
    assert scan_calls == 1


def test_parser_requires_one_exact_template_header_and_reports_header_errors(tmp_path):
    missing = tmp_path / "missing.xlsx"
    duplicate = tmp_path / "duplicate.xlsx"
    padded = tmp_path / "padded.xlsx"
    write_workbook(missing, [["内部订单号", "模板ID"], ["A-1", "TEMPLATE-A"]])
    write_workbook(duplicate, [["模板", "模板"], ["TEMPLATE-A", "TEMPLATE-B"]])
    write_workbook(padded, [["订单号", " 模板 "], ["A-1", "TEMPLATE-A"]])

    missing_batch = MultiTemplateOrderParser().parse(missing)
    duplicate_batch = MultiTemplateOrderParser().parse(duplicate)
    padded_batch = MultiTemplateOrderParser().parse(padded)

    assert [issue.code for issue in missing_batch.issues] == ["template_column_missing"]
    assert [issue.excel_row for issue in missing_batch.issues] == [1]
    assert [issue.code for issue in duplicate_batch.issues] == ["template_column_duplicate"]
    assert duplicate_batch.rows == ()
    assert [issue.code for issue in padded_batch.issues] == ["template_column_missing"]


def test_parser_collects_empty_and_invalid_template_rows_without_losing_valid_groups(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["订单号", "模板", "说明"], ["OK-1", "TEMPLATE-A", "normal"], ["EMPTY-1", "  ", "must report"], ["INVALID-1", "../outside", "must report"], ["OK-2", "TEMPLATE-B", "normal"]])

    batch = MultiTemplateOrderParser().parse(source)

    assert [group.template_id for group in batch.groups] == ["TEMPLATE-A", "TEMPLATE-B"]
    assert [(issue.code, issue.excel_row, issue.order_no) for issue in batch.issues] == [("template_id_missing", 3, "EMPTY-1"), ("template_id_invalid", 4, "INVALID-1")]
    assert all("C:" not in issue.message for issue in batch.issues)


def test_parser_trims_template_values_and_writes_cached_formula_values_for_legacy_consumers(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["订单编号", "模板", "数量", "公式列"], ["A", " TEMPLATE-A ", 1, "=1+1"], ["B", "template-a", 2, "literal"], ["C", 1001, 3, 6]])
    cache_numeric_formula_result(source, "=1+1", 2)

    batch = MultiTemplateOrderParser().parse(source)

    assert batch.issues == ()
    assert [group.template_id for group in batch.groups] == ["TEMPLATE-A", "template-a", "1001"]
    assert batch.rows[2].values["数量"] == 3
    assert batch.rows[0].values["公式列"] == 2
    grouped = GroupWorkbookWriter().write(batch, tmp_path / "parent-job")
    assert grouped["TEMPLATE-A"] != grouped["template-a"]
    formula_value = load_workbook(grouped["TEMPLATE-A"], read_only=True, data_only=True).active["D2"].value
    assert formula_value == 2


def test_parser_rejects_formula_cells_without_a_cached_value(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["订单号", "模板", "数量"], ["A-1", "TEMPLATE-A", "=1+1"]])

    batch = MultiTemplateOrderParser().parse(source)

    assert batch.rows == ()
    assert [(issue.code, issue.excel_row, issue.order_no) for issue in batch.issues] == [("formula_value_missing", 2, "A-1")]
    with pytest.raises(ValueError, match="解析问题"):
        GroupWorkbookWriter().write(batch, tmp_path / "parent-job")


def test_parser_groups_one_thousand_rows_across_thirty_templates_in_one_pass(tmp_path):
    source = tmp_path / "orders.xlsx"
    rows = [["内部订单号", "模板", "数量"]]
    rows.extend([f"ORDER-{index}", f"T{index % 30:02d}", index] for index in range(1000))
    write_workbook(source, rows)

    batch = MultiTemplateOrderParser().parse(source)

    assert batch.issues == ()
    assert len(batch.rows) == 1000
    assert len(batch.groups) == 30
    assert [group.template_id for group in batch.groups] == [f"T{index:02d}" for index in range(30)]
    assert sum(len(group.rows) for group in batch.groups) == 1000


def test_models_serialize_without_mutable_value_leaks_and_round_trip_through_job_store(tmp_path):
    shipped_at = datetime(2026, 8, 26, 9, 30)
    row = MultiTemplateOrderRow(
        "订单",
        2,
        "A-1",
        "TEMPLATE-A",
        {"模板": "TEMPLATE-A", "交付时间": shipped_at, "金额": Decimal("12.50")},
        ("TEMPLATE-A", shipped_at, Decimal("12.50")),
    )
    issue = MultiTemplateIssue("template_not_found", "模板不存在。", "检查模板 ID。", "TEMPLATE-A", 2, "A-1")

    with pytest.raises(TypeError):
        row.values["模板"] = "OTHER"

    assert MultiTemplateOrderRow.from_dict(row.to_dict()) == row
    assert MultiTemplateIssue.from_dict(issue.to_dict()) == issue
    batch = MultiTemplateOrderBatch("订单", ("模板", "  customer field  ", "交付时间", "金额"), (row,), (), (issue,))
    record = JobStore(tmp_path / "jobs").create({"multi_template_batch": batch.to_dict()})
    restored = MultiTemplateOrderBatch.from_dict(JobStore(tmp_path / "jobs").load(record["job_id"])["request"]["multi_template_batch"])
    assert restored.rows[0].values["交付时间"] == shipped_at
    assert restored.rows[0].raw_values[2] == Decimal("12.50")
    assert restored.headers[1] == "  customer field  "
    assert {"ready", "interrupted", "completed_with_errors"} <= PARENT_JOB_STATUSES
    assert {"pending", "canary_failed", "succeeded"} <= TEMPLATE_GROUP_STATUSES


def test_group_workbook_writer_isolates_rows_and_preserves_source_file(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["内部订单号", "模板", "数量", "重复", "重复"], ["A-1", "TEMPLATE-A", 1, "first", "last"], ["B-1", "TEMPLATE-B", 2, "first-b", "last-b"], ["A-2", "TEMPLATE-A", 3, "first-2", "last-2"]], sheet_name="业务订单")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    batch = MultiTemplateOrderParser().parse(source, sheet_name="业务订单")

    paths = GroupWorkbookWriter().write(batch, tmp_path / "parent-job")

    assert list(paths) == ["TEMPLATE-A", "TEMPLATE-B"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    for template_id, path in paths.items():
        workbook = load_workbook(path, read_only=True, data_only=True)
        values = list(workbook["业务订单"].iter_rows(values_only=True))
        assert values[0] == ("内部订单号", "模板", "数量", "重复", "重复")
        assert {row[1] for row in values[1:]} == {template_id}
    values = list(load_workbook(paths["TEMPLATE-A"], read_only=True).active.iter_rows(values_only=True))
    assert values[1:] == [("A-1", "TEMPLATE-A", 1, "first", "last"), ("A-2", "TEMPLATE-A", 3, "first-2", "last-2")]


def test_group_workbook_writer_refuses_a_batch_with_parse_issues(tmp_path):
    source = tmp_path / "orders.xlsx"
    write_workbook(source, [["订单号", "模板"], ["A-1", ""]])

    with pytest.raises(ValueError, match="解析问题"):
        GroupWorkbookWriter().write(MultiTemplateOrderParser().parse(source), tmp_path / "parent-job")


def test_parser_requires_sheet_name_when_a_workbook_has_multiple_sheets(tmp_path):
    source = tmp_path / "orders.xlsx"
    workbook = Workbook()
    workbook.active.title = "summary"
    workbook.active.append(["说明"])
    orders = workbook.create_sheet("orders")
    orders.append(["订单号", "模板"])
    orders.append(["A-1", "TEMPLATE-A"])
    workbook.save(source)

    unspecified = MultiTemplateOrderParser().parse(source)
    selected = MultiTemplateOrderParser().parse(source, sheet_name="orders")

    assert [issue.code for issue in unspecified.issues] == ["sheet_name_required"]
    assert [group.template_id for group in selected.groups] == ["TEMPLATE-A"]


def test_existing_single_template_parsers_consume_group_workbooks(tmp_path):
    def grouped_path(name, template_id, headers, values, *, cached_formula=None):
        source = tmp_path / f"{name}.xlsx"
        write_workbook(source, [headers, values])
        if cached_formula:
            cache_numeric_formula_result(source, *cached_formula)
        batch = MultiTemplateOrderParser().parse(source)
        return GroupWorkbookWriter().write(batch, tmp_path / f"{name}-parent")[template_id]

    legacy_path = grouped_path(
        "legacy",
        "LEGACY-A",
        ["内部订单号", "订单明细id", "生产部门", "产品中文名称", "字体颜色", "模板", "定制信息", "字体", "设计"],
        ["ORDER-1", "DETAIL-1", "K", "产品", "Gold", "LEGACY-A", "Alice", "F1", "Design1"],
    )
    legacy_rows = read_202603_rows(legacy_path)
    assert len(parse_202508_items(legacy_rows, template_id="LEGACY-A")) == 1

    curved_path = grouped_path(
        "curved",
        CURVED_TEMPLATE_ID,
        ["内部订单号", "订单明细id", "生产部门", "产品中文名称", "字体颜色", "模板", "定制信息"],
        ["ORDER-2", "DETAIL-2", "K", "产品", "Gold", CURVED_TEMPLATE_ID, "Font Option: F1\nTitle: A\nNames: Alice"],
    )
    curved_items = parse_202509_items(read_202509_rows(curved_path))
    assert curved_items and {item.order_no for item in curved_items} == {"ORDER-2"}

    grouped_path_202603 = grouped_path(
        "config-grouped",
        CONFIG_TEMPLATE_ID,
        ["内部订单号", "订单明细id", "购买数量", "模板", "Style Option", "Font Option", "Personalization", "生产部门"],
        ["ORDER-3", "DETAIL-3", 1, CONFIG_TEMPLATE_ID, "Style1", "F1", "Alice", "K"],
    )
    task = build_grouped_task(grouped_path_202603, tmp_path / "template.config.json", tmp_path / "out.ai", 4, allowed_font_options={"F1"})
    assert len(task.groups) == 1

    generic_path = grouped_path(
        "generic",
        "GENERIC-A",
        ["订单号", "模板", "Name", "Quantity"],
        ["ORDER-4", "GENERIC-A", "Alice", "=1+1"],
        cached_formula=("=1+1", 2),
    )
    assert read_generic_rows(generic_path, sheet_name="订单")[0]["模板"] == "GENERIC-A"
    assert read_generic_rows(generic_path, sheet_name="订单")[0]["Quantity"] == 2

    v2_path = grouped_path("v2", "V2-A", ["订单号", "模板", "Name"], ["ORDER-5", "V2-A", "Alice"])
    assert read_order_rows({"order_file": str(v2_path), "sheet_name": "订单"})[0]["模板"] == "V2-A"
