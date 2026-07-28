from pathlib import Path
import json

import pytest
from openpyxl import Workbook

from src.service import render_service as render_service_module
from src.service.generic_rule_renderer import GenericRuleRenderError, build_generic_render_task
from src.service.job_store import JobStore
from src.service.render_service import RenderService
from src.service.rule_center import check_template_definition
from src.service.template_inspector import TemplateInspector
from src.service.template_onboarding import TemplateOnboardingStore
from src.service.template_publication import TemplatePublicationService
from src.service.template_registry import TemplateRegistry
from src.service.template_rule_ast import migrate_legacy_rule_ast


def make_template(tmp_path, assets=None):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = registry.save_uploaded_ai("GENERIC001", "template.ai", b"ai")
    return registry, registry.upsert_template(
        {
            "template_id": "GENERIC001",
            "name": "Generic",
            "template_type": "pure_text",
            "pipeline": "generic_rules_only",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
            "assets": assets or [],
        }
    )


def make_orders(path: Path, custom="Alice | Bob", quantity=None, quantity_header="Quantity"):
    workbook = Workbook()
    sheet = workbook.active
    headers = ["Order", "Custom", "Font", "Design", "Style", "Color"]
    row = ["A-1", custom, "F2", "Design1", "Style3", "Gold"]
    if quantity is not None:
        headers.append(quantity_header)
        row.append(quantity)
    sheet.append(headers)
    sheet.append(row)
    workbook.save(path)


def base_rules():
    return {
        "order_bindings": {
            "order_no": "Order",
            "text": "Custom",
            "font": "Font",
            "design": "Design",
            "style": "Style",
            "color": "Color",
        },
        "slot_mappings": [
            {"field": "text", "slot": "Name1", "delimiter": "|", "sequence_index": 1},
            {"field": "text", "slot": "Name2", "delimiter": "|", "sequence_index": 2},
        ],
        "option_groups": [
            {"role": "font_options", "values": ["F1", "F2"]},
            {"role": "design_options", "values": ["Design1", "Design2"]},
            {"role": "style_options", "values": ["Style1", "Style3"]},
        ],
        "dimensions": {"Name1": {"width_mm": 20, "height_mm": 5}},
        "text_policies": {"fit": "scale_to_box"},
        "output": {"color_mode": "CMYK"},
        "transforms": {"outline_text": True},
    }


def exact_box_rules(*, single_file=True):
    rules = base_rules()
    rules["order_bindings"] = {**rules["order_bindings"], "year": "Year"}
    rules["slot_mappings"] = [
        {"field": "text", "slot": "Name"},
        {"field": "year", "slot": "Year", "optional": True},
    ]
    rules["dimensions"] = {
        "Name": {"width_mm": 17, "height_mm": 10},
        "Year": {"width_mm": 11, "height_mm": 5},
    }
    layout = {
        "type": "name_columns",
        "name": {"delimiter": "|", "segment_box_target": "Name", "fill_box_exactly": True},
        "footer": {"box_target": "Year", "optional": True, "fill_box_exactly": True},
        "default": {"group_by": ["row"], "footer_field": "year"},
    }
    if single_file:
        layout["output_mode"] = "single_file"
    rules["render_layout"] = layout
    return rules


def make_orders_with_year(path: Path, count=1):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Font", "Design", "Style", "Color", "Year"])
    for index in range(count):
        sheet.append([f"A-{index + 1}", "Alice | Bob", "F2", "Design1", "Style3", "Gold", 2026])
    workbook.save(path)


def test_builds_generic_task_from_bindings_and_split_variables(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)

    task = build_generic_render_task(
        template, base_rules(), order_path, tmp_path / "output.ai", columns=3
    )

    assert task["type"] == "generic_template_rules"
    assert task["layout"]["columns"] == 3
    assert task["orders"][0]["order_no"] == "A-1"
    assert task["orders"][0]["selections"] == {
        "font": "F2",
        "design": "Design1",
        "style": "Style3",
        "color": "Gold",
    }
    assert task["orders"][0]["variables"] == [
        {"target": "Name1", "field": "text", "value": "Alice"},
        {"target": "Name2", "field": "text", "value": "Bob"},
    ]
    assert task["dimensions"]["Name1"]["width_mm"] == 20
    assert task["output"]["color_mode"] == "CMYK"


def test_multi_name_customization_copies_the_full_text_by_quantity(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob", quantity=3)
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 3
    assert [order["quantity_index"] for order in task["orders"]] == [1, 2, 3]
    assert [order["quantity"] for order in task["orders"]] == [3, 3, 3]
    assert [order["variables"] for order in task["orders"]] == [
        [{"target": "Name", "field": "text", "value": "Alice | Bob"}],
        [{"target": "Name", "field": "text", "value": "Alice | Bob"}],
        [{"target": "Name", "field": "text", "value": "Alice | Bob"}],
    ]


def test_multi_name_quantity_creates_independent_name_columns_cards(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob", quantity=3)
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}
    rules["render_layout"] = {
        "type": "name_columns",
        "output_mode": "single_file",
        "packing": "masonry",
        "default": {"group_by": ["order_no"], "header_fields": ["order_no"]},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 3
    assert [len(order["layout_members"]) for order in task["orders"]] == [1, 1, 1]
    assert [order["layout_members"][0]["quantity_index"] for order in task["orders"]] == [1, 2, 3]
    assert [order["layout_members"][0]["order_no"] for order in task["orders"]] == ["A-1", "A-1", "A-1"]


def test_multi_name_cards_keep_bound_year_in_every_copy(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Font", "Design", "Style", "Color", "Quantity", "Year"])
    sheet.append(["A-1", "Alice | Bob", "F2", "Design1", "Style3", "Gold", 3, 2025])
    workbook.save(order_path)
    rules = base_rules()
    rules["order_bindings"].update({"quantity": "Quantity", "year": "Year"})
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}
    rules["render_layout"] = {
        "type": "name_columns",
        "output_mode": "single_file",
        "default": {"group_by": ["row"], "footer_field": "year"},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert [order["layout_members"][0]["values"]["year"] for order in task["orders"]] == [2025, 2025, 2025]


def test_multi_name_cards_allow_missing_optional_year_column(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob", quantity=2)
    rules = base_rules()
    rules["order_bindings"].update({"quantity": "Quantity", "year": "Year"})
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}
    rules["render_layout"] = {
        "type": "name_columns",
        "output_mode": "single_file",
        "default": {"group_by": ["row"], "footer_field": "year"},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 2
    assert [order["layout_members"][0]["values"]["year"] for order in task["orders"]] == [None, None]


def test_year_binding_is_required_when_not_used_as_an_optional_footer(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["order_bindings"]["year"] = "Year"

    with pytest.raises(GenericRuleRenderError, match="Year"):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_multi_name_customization_keeps_one_copy_for_quantity_one(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob", quantity=1)
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 1
    assert task["orders"][0]["quantity_index"] == 1
    assert task["orders"][0]["variables"] == [
        {"target": "Name", "field": "text", "value": "Alice | Bob"}
    ]


def test_quantity_does_not_copy_single_content_template_when_switch_is_off(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob", quantity=3)

    rules = base_rules()
    rules["text_policies"]["split"] = {
        "delimiter": "|",
        "max_parts": 2,
        "overflow": "reject",
        "trim": True,
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 1
    assert task["orders"][0]["variables"] == [
        {"target": "Name1", "field": "text", "value": "Alice"},
        {"target": "Name2", "field": "text", "value": "Bob"},
    ]


def test_multi_name_customization_honors_a_configured_quantity_column(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, quantity=2, quantity_header="Pieces")
    rules = base_rules()
    rules["order_bindings"]["quantity"] = "Pieces"
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 2
    assert [order["quantity_index"] for order in task["orders"]] == [1, 2]


@pytest.mark.parametrize("quantity", ["", 0, -1, "1.5", "three"])
def test_multi_name_customization_rejects_invalid_quantity(tmp_path, quantity):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, quantity=quantity)
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}

    with pytest.raises(GenericRuleRenderError, match="订单数量"):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_multi_name_customization_requires_a_quantity_column(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["multi_name_customization"] = {"enabled": True}

    with pytest.raises(GenericRuleRenderError, match="数量列"):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_passes_declared_font_option_styles_to_generic_task(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["font_option_styles"] = {
        "F2": {"font_name": "Milkshake", "font_family": "Milkshake"}
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["font_option_styles"] == {
        "F2": {"font_name": "Milkshake", "font_family": "Milkshake"}
    }


def test_generic_task_resolves_case_and_common_chinese_order_columns(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["\u5185\u90e8\u8ba2\u5355\u53f7", "\u6a21\u677f", "\u5b57\u4f53", "Name"])
    sheet.append(["A-1", template.template_id, "F5", "Alice|Bob"])
    sheet.append(["B-1", "OTHER_TEMPLATE", "F7", "Skip Me"])
    workbook.save(order_path)
    rules = base_rules()
    rules["order_bindings"] = {"text": "name"}
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["font_style_rules"] = [{"font_options": ["F5"], "boldness": 0.5}]

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert len(task["orders"]) == 1
    assert task["orders"][0]["order_no"] == "A-1"
    assert task["orders"][0]["selections"]["font"] == "F5"
    assert task["orders"][0]["variables"] == [
        {"target": "Name", "field": "text", "value": "Alice|Bob", "font_style": {"boldness": 0.5}}
    ]


def test_resolves_selected_design_asset(tmp_path):
    registry, original = make_template(tmp_path)
    assets = registry.save_uploaded_assets(
        "GENERIC001", [{"filename": "design-1.ai", "content": b"asset"}]
    )
    template = registry.upsert_template(
        {
            "template_id": "GENERIC001",
            "name": "Generic",
            "template_type": "asset_split",
            "pipeline": "generic_rules_only",
            "status": "active",
            "template_ai": str(original.template_ai),
            "assets": assets,
        }
    )
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["asset_mappings"] = [
        {"option": "Design1", "asset": "design-1.ai", "target": "DesignSlot"}
    ]

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["assets"][0]["target"] == "DesignSlot"
    assert task["orders"][0]["assets"][0]["asset"].endswith("design-1.ai")


def test_asset_mappings_support_font_and_style_options(tmp_path):
    registry, original = make_template(tmp_path)
    assets = registry.save_uploaded_assets(
        "GENERIC001",
        [
            {"filename": "font-f2.ai", "content": b"font"},
            {"filename": "style-3.ai", "content": b"style"},
        ],
    )
    template = registry.upsert_template({**original.to_json_dict(), "assets": assets})
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["asset_mappings"] = [
        {"option": "F2", "asset": "font-f2.ai", "target": "FontSlot"},
        {"option": "Style3", "asset": "style-3.ai", "target": "StyleSlot"},
    ]

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert {item["target"] for item in task["orders"][0]["assets"]} == {"FontSlot", "StyleSlot"}


def test_runtime_split_uses_confirmed_trim_max_parts_and_overflow(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice | Bob | Carol")
    rules = base_rules()
    rules["text_policies"] = {
        "fit": "scale_to_box",
        "split": {"delimiter": "|", "trim": False, "max_parts": 2, "overflow": "truncate"},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert [item["value"] for item in task["orders"][0]["variables"]] == ["Alice ", " Bob "]

    rules["text_policies"]["split"]["overflow"] = "reject"
    with pytest.raises(GenericRuleRenderError, match="cannot satisfy split policy"):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_numeric_zero_is_preserved_as_template_text(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom=0)

    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name1"}]
    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["variables"][0]["value"] == "0"


def test_passes_configured_name_color_cycle_to_exact_name_target(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice|Bob|Carol")
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["name_color_cycle"] = {
        "delimiter": "|",
        "colors": ["#D71920", "#000000", "#0000FF"],
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["variables"] == [
        {
            "target": "Name",
            "field": "text",
            "value": "Alice|Bob|Carol",
            "name_color_cycle": {
                "delimiter": "|",
                "colors": ["#D71920", "#000000", "#0000FF"],
            },
        }
    ]


def test_name_columns_layout_groups_t_department_by_order_and_color(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Font", "Color", "Department"])
    sheet.append(["A-1", "Alice|Bob", "F2", "Black", "T"])
    sheet.append(["A-1", "Carol|Dan", "F3", "Black", "T"])
    sheet.append(["A-1", "Eve", "F4", "Red", "T"])
    sheet.append(["B-1", "Frank", "F5", "Black", "T"])
    workbook.save(order_path)
    rules = {
        **base_rules(),
        "order_bindings": {"order_no": "Order", "text": "Custom", "font": "Font", "color": "Color", "department": "Department"},
        "slot_mappings": [{"field": "text", "slot": "Name"}],
        "render_layout": {
            "type": "name_columns",
            "packing": "masonry",
            "page_height_mm": 1320,
            "card_gap_mm": 8,
            "default": {"group_by": ["row"]},
            "department_overrides": {"T": {"group_by": ["order_no", "color"], "header_fields": ["order_no", "color"]}},
        },
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["render_layout"]["type"] == "name_columns"
    assert task["render_layout"]["packing"] == "masonry"
    assert task["render_layout"]["page_height_mm"] == 1320
    assert len(task["orders"]) == 3
    assert [member["order_no"] for member in task["orders"][0]["layout_members"]] == ["A-1", "A-1"]
    assert task["orders"][0]["layout_mode"]["header_fields"] == ["order_no", "color"]


def test_boxed_name_columns_validate_one_to_seven_names_and_keep_year_optional(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders-with-year.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Year"])
    sheet.append(["A-1", "One|Two|Three|Four|Five|Six|Seven", "2026"])
    workbook.save(order_path)
    rules = {
        "order_bindings": {"order_no": "Order", "text": "Custom", "year": "Year"},
        "slot_mappings": [
            {"field": "text", "slot": "Name"},
            {"field": "year", "slot": "Year", "optional": True},
        ],
        "dimensions": {
            "Name": {"width_mm": 17, "height_mm": 10},
            "Year": {"width_mm": 11, "height_mm": 5},
        },
        "text_policies": {"fit": "scale_to_box"},
        "render_layout": {
            "type": "name_columns",
            "name": {"delimiter": "|", "segment_box_target": "Name", "min_parts": 1, "max_parts": 7},
            "footer": {"box_target": "Year", "optional": True},
            "default": {"group_by": ["row"], "footer_field": "year"},
        },
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["variables"] == [
        {"target": "Name", "field": "text", "value": "One|Two|Three|Four|Five|Six|Seven"},
        {"target": "Year", "field": "year", "value": "2026"},
    ]

    no_year_path = tmp_path / "orders-without-year.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom"])
    sheet.append(["B-1", "OnlyName"])
    workbook.save(no_year_path)
    no_year_task = build_generic_render_task(template, rules, no_year_path, tmp_path / "without-year.ai")

    assert no_year_task["orders"][0]["variables"] == [
        {"target": "Name", "field": "text", "value": "OnlyName"}
    ]


@pytest.mark.parametrize(
    ("custom", "message"),
    [
        ("One||Three", "空姓名段"),
        ("One|Two|Three|Four|Five|Six|Seven|Eight", "1 至 7"),
    ],
)
def test_boxed_name_columns_reject_invalid_name_segments(tmp_path, custom, message):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Custom"])
    sheet.append([custom])
    workbook.save(order_path)
    rules = {
        "order_bindings": {"text": "Custom"},
        "slot_mappings": [{"field": "text", "slot": "Name"}],
        "text_policies": {"fit": "scale_to_box"},
        "render_layout": {
            "type": "name_columns",
            "name": {"delimiter": "|", "segment_box_target": "Name", "min_parts": 1, "max_parts": 7},
        },
    }

    with pytest.raises(GenericRuleRenderError, match=message):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_unconfigured_name_columns_keep_existing_segment_validation_behavior(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Custom"])
    sheet.append(["One||Three|Four|Five|Six|Seven|Eight"])
    workbook.save(order_path)
    rules = {
        "order_bindings": {"text": "Custom"},
        "slot_mappings": [{"field": "text", "slot": "Name"}],
        "text_policies": {"fit": "scale_to_box"},
        "render_layout": {"type": "name_columns", "name": {"delimiter": "|"}},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["variables"][0]["value"] == "One||Three|Four|Five|Six|Seven|Eight"


def test_boxed_name_columns_keep_optional_year_for_t_department_override(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Year", "Department", "Color"])
    sheet.append(["T-1", "One|Two", "2026", "T", "Black"])
    workbook.save(order_path)
    rules = {
        "order_bindings": {
            "order_no": "Order",
            "text": "Custom",
            "year": "Year",
            "department": "Department",
            "color": "Color",
        },
        "slot_mappings": [
            {"field": "text", "slot": "Name"},
            {"field": "year", "slot": "Year", "optional": True},
        ],
        "text_policies": {"fit": "scale_to_box"},
        "render_layout": {
            "type": "name_columns",
            "name": {"delimiter": "|", "segment_box_target": "Name", "min_parts": 1, "max_parts": 7},
            "footer": {"box_target": "Year", "optional": True},
            "default": {"group_by": ["row"], "footer_field": "year"},
            "department_overrides": {
                "T": {"group_by": ["order_no", "color"], "header_fields": ["order_no", "color"]}
            },
        },
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert task["orders"][0]["layout_mode"]["footer_field"] == "year"
    assert {variable["target"] for variable in task["orders"][0]["variables"]} == {"Name", "Year"}


def test_confirmed_jjmb_202510_template_enables_exact_box_fill_only_for_its_layout():
    rule_path = Path("templates/JJMB202510241154389614/template.rules.json")
    rules = json.loads(rule_path.read_text(encoding="utf-8"))["rules"]

    assert rules["render_layout"]["name"]["segment_box_target"] == "Name"
    assert rules["render_layout"]["name"]["fill_box_exactly"] is True
    assert rules["render_layout"]["footer"] == {
        "box_target": "Year",
        "optional": True,
        "fill_box_exactly": True,
    }
    assert rules["output"] == {"color_mode": "CMYK"}


def test_generic_service_attaches_layout_audit_for_single_file_exact_boxes(tmp_path):
    registry, template = make_template(tmp_path)
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": template.template_id, "profile": "composite"},
        "rules": exact_box_rules(single_file=True),
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    template = registry.apply_confirmed_rule_pack(template.template_id, pack, activate=True)
    order_path = tmp_path / "orders.xlsx"
    make_orders_with_year(order_path)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": True}
    )

    assert result["status"] == "completed", result.get("error")
    task = json.loads(Path(result["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert Path(task["layout_audit_file"]).name == "layout-audit.tsv"
    assert task["layout_audit_files"] == [task["layout_audit_file"]]
    assert result["outputs"]["layout_audit_file"] == task["layout_audit_file"]
    assert result["outputs"]["layout_audit_files"] == task["layout_audit_files"]
    assert result["stats"]["exact_text_box_audit"] is True


def test_generic_service_uses_separate_layout_audits_for_exact_box_chunks(tmp_path, monkeypatch):
    registry, template = make_template(tmp_path)
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": template.template_id, "profile": "composite"},
        "rules": exact_box_rules(single_file=False),
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    template = registry.apply_confirmed_rule_pack(template.template_id, pack, activate=True)
    order_path = tmp_path / "many-orders.xlsx"
    make_orders_with_year(order_path, count=17)
    calls = []

    class Bridge:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def render(self, script, task_path):
            task = json.loads(Path(task_path).read_text(encoding="utf-8"))
            calls.append((Path(task_path).name, len(task["orders"]), Path(task["layout_audit_file"]).name))
            for output in task["output_ai_files"]:
                Path(output).write_text("ai", encoding="utf-8")

        def close(self):
            return None

    monkeypatch.setattr(render_service_module, "IllustratorBridge", Bridge)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": False}
    )

    assert result["status"] == "completed", result.get("error")
    assert calls == [
        ("render-task-001.json", 8, "layout-audit-001.tsv"),
        ("render-task-002.json", 8, "layout-audit-002.tsv"),
        ("render-task-003.json", 1, "layout-audit-003.tsv"),
    ]
    task = json.loads(Path(result["outputs"]["render_task"]).read_text(encoding="utf-8"))
    assert [Path(path).name for path in task["layout_audit_files"]] == [
        "layout-audit-001.tsv",
        "layout-audit-002.tsv",
        "layout-audit-003.tsv",
    ]
    assert result["outputs"]["layout_audit_files"] == task["layout_audit_files"]


def test_passes_selected_font_boldness_to_every_text_variable(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["font_style_rules"] = [
        {"font_options": ["F2", "F3", "F10", "F11", "F12"], "boldness": 0.4},
        {"font_options": ["F5", "F6", "F7", "F8", "F9"], "boldness": 0.5},
    ]

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert [variable["font_style"] for variable in task["orders"][0]["variables"]] == [
        {"boldness": 0.4},
        {"boldness": 0.4},
    ]


def test_compiles_confirmed_rule_ast_into_generic_runtime_actions(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice|Bob|Carol")
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name"}]
    rules["rule_ast"] = migrate_legacy_rule_ast(
        name_color_cycle={"delimiter": "|", "colors": ["Red", "Black", "Blue"]},
        font_style_rules=[{"font_options": ["F2"], "boldness": 0.4}],
    )

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    variable = task["orders"][0]["variables"][0]
    assert [action["type"] for action in variable["actions"]] == ["fill_color", "stroke_width"]
    assert "name_color_cycle" not in variable
    assert "font_style" not in variable


def test_does_not_pass_name_color_cycle_to_numbered_name_targets(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path, custom="Alice|Bob")
    rules = base_rules()
    rules["slot_mappings"] = [{"field": "text", "slot": "Name1"}]
    rules["name_color_cycle"] = {"delimiter": "|", "colors": ["#D71920", "#000000"]}

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    assert "name_color_cycle" not in task["orders"][0]["variables"][0]


def test_builds_effective_transform_for_each_selected_option(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["transforms"] = {
        "outline_text": True,
        "Design1": {"rotation_deg": 90},
        "Style3": {"scale_percent": 110, "offset_x_mm": 2},
    }

    task = build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")

    transforms = task["orders"][0]["transforms"]
    assert transforms["outline_text"] is True
    assert transforms["targets"] == [
        {"target": "Design1", "rotation_deg": 90},
        {"target": "Style3", "scale_percent": 110, "offset_x_mm": 2},
    ]


def test_rejects_missing_bound_order_column(tmp_path):
    _, template = make_template(tmp_path)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)
    rules = base_rules()
    rules["order_bindings"]["title"] = "Missing Header"

    with pytest.raises(GenericRuleRenderError, match="missing bound columns"):
        build_generic_render_task(template, rules, order_path, tmp_path / "output.ai")


def test_generic_pipeline_dry_run_consumes_confirmed_rule_pack(tmp_path):
    registry, template = make_template(tmp_path)
    rules = base_rules()
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": template.template_id, "profile": "composite"},
        "rules": rules,
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    template = registry.apply_confirmed_rule_pack(template.template_id, pack, activate=True)
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": True}
    )

    assert check_template_definition(template)["renderable"] is True
    assert result["status"] == "completed"
    assert result["stats"] == {"orders": 1, "variables": 2, "assets": 0, "dry_run": True}
    task = Path(result["outputs"]["render_task"]).read_text(encoding="utf-8")
    assert '"target": "Name1"' in task
    assert '"value": "Alice"' in task


def test_generic_pipeline_splits_large_illustrator_runs_into_chunks(tmp_path, monkeypatch):
    registry, template = make_template(tmp_path)
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": template.template_id, "profile": "composite"},
        "rules": base_rules(),
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    template = registry.apply_confirmed_rule_pack(template.template_id, pack, activate=True)
    order_path = tmp_path / "many-orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Order", "Custom", "Font", "Design", "Style", "Color"])
    for index in range(17):
        sheet.append([f"A-{index}", f"Alice {index} | Bob {index}", "F2", "Design1", "Style3", "Gold"])
    workbook.save(order_path)
    calls = []

    class Bridge:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def render(self, script, task_path):
            task = json.loads(Path(task_path).read_text(encoding="utf-8"))
            calls.append((Path(task_path).name, len(task["orders"]), self.kwargs))
            for output in task["output_ai_files"]:
                Path(output).write_text("ai", encoding="utf-8")

        def close(self):
            return None

    monkeypatch.setattr(render_service_module, "IllustratorBridge", Bridge)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": False}
    )

    assert result["status"] == "completed", result.get("error")
    assert calls == [
        ("render-task-001.json", 8, {"visible": False, "fresh_instance": True, "reuse_instance": True}),
        ("render-task-002.json", 8, {"visible": False, "fresh_instance": True, "reuse_instance": True}),
        ("render-task-003.json", 1, {"visible": False, "fresh_instance": True, "reuse_instance": True}),
    ]
    assert len(result["outputs"]["output_ai_files"]) == 17


def test_generic_pipeline_retries_only_transient_illustrator_com_errors(tmp_path, monkeypatch):
    from src.renderer.illustrator_bridge import IllustratorBridgeError

    registry, template = make_template(tmp_path)
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": template.template_id, "profile": "composite"},
        "rules": base_rules(),
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    template = registry.apply_confirmed_rule_pack(template.template_id, pack, activate=True)
    order_path = tmp_path / "one-order.xlsx"
    make_orders(order_path)
    calls = []

    class Bridge:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def render(self, script, task_path):
            calls.append(Path(task_path).name)
            if len(calls) < 3:
                raise IllustratorBridgeError("执行 Illustrator JSX 失败: (-2147417851, 'server unavailable')")
            task = json.loads(Path(task_path).read_text(encoding="utf-8"))
            for output in task["output_ai_files"]:
                Path(output).write_text("ai", encoding="utf-8")

        def reset(self):
            return None

        def close(self):
            return None

    monkeypatch.setattr(render_service_module, "IllustratorBridge", Bridge)
    monkeypatch.setattr(render_service_module.time, "sleep", lambda seconds: None)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": False}
    )

    assert result["status"] == "completed", result.get("error")
    assert calls == ["render-task.json", "render-task.json", "render-task.json"]


def test_render_service_uses_generic_rules_when_legacy_pipeline_has_executable_pack(tmp_path):
    registry = TemplateRegistry(tmp_path / "templates.json", tmp_path / "templates")
    ai_path = registry.save_uploaded_ai("LEGACYGENERIC001", "template.ai", b"ai")
    pack = {
        "$schema": "custom-renderer/template-rule-pack",
        "template": {"template_id": "LEGACYGENERIC001", "profile": "unclassified"},
        "rules": base_rules(),
        "assets": {"items": [], "policy": {"mode": "inline"}},
        "capabilities": ["replace_text", "scale_to_box"],
        "validation": {"status": "confirmed", "unresolved_items": []},
    }
    rules_path = registry.save_template_rules_config("LEGACYGENERIC001", json.dumps(pack))
    template = registry.upsert_template(
        {
            "template_id": "LEGACYGENERIC001",
            "name": "Legacy generic",
            "template_type": "pure_text_color_design",
            "pipeline": "jjmb_202508",
            "status": "active",
            "template_ai": registry.to_config_path(ai_path),
            "template_rules_config": registry.to_config_path(rules_path),
        }
    )
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs-legacy")).submit(
        {"template_id": template.template_id, "order_file": str(order_path), "dry_run": True}
    )

    assert result["status"] == "completed"
    assert result["stats"]["orders"] == 1
    assert result["stats"]["variables"] == 2


def test_new_generic_template_scan_confirm_and_dry_run(tmp_path):
    registry, template = make_template(tmp_path)

    class ScanBridge:
        def __init__(self, **kwargs):
            pass

        def render(self, script, task_path):
            task = json.loads(task_path.read_text(encoding="utf-8"))
            Path(task["output_json"]).write_text(
                json.dumps(
                    {
                        "layers": [{"name": "Template"}],
                        "items": [
                            {"type": "TextFrame", "name": "F1", "path": "Template/FONT_STYLES/F1"},
                            {"type": "TextFrame", "name": "Name1", "path": "Template/Name1"},
                        ],
                    }
                ),
                encoding="utf-8",
            )

    store = TemplateOnboardingStore(registry.storage_dir)
    state = TemplateInspector(ScanBridge).scan(template, store)
    pack = state["draft"]
    pack["validation"]["unresolved_items"] = []
    pack["validation"]["sample"] = {
        "input": {"Order": "A-1", "Custom": "Alice", "Font": "F1"},
        "expected": {"Name1": "Alice"},
    }
    pack["rules"]["order_bindings"] = {
        "order_no": "Order",
        "text": "Custom",
        "font": "Font",
    }
    pack["rules"]["slot_mappings"] = [{"field": "text", "slot": "Name1"}]
    pack["rules"]["text_policies"] = {"fit": "scale_to_box"}
    published = TemplatePublicationService(registry, store).confirm(
        template.template_id, pack, change_summary="generic e2e"
    )
    order_path = tmp_path / "orders.xlsx"
    make_orders(order_path)

    result = RenderService(registry=registry, jobs=JobStore(tmp_path / "jobs-e2e")).submit(
        {"template_id": published["template"].template_id, "order_file": str(order_path), "dry_run": True}
    )

    assert published["template"].status == "active"
    assert result["status"] == "completed"
    assert result["stats"]["variables"] == 1
