from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.service import multi_template_preflight as preflight_module
from src.service.multi_template_preflight import MultiTemplatePreflight
from src.service.multi_template_snapshot import TemplateResolutionBatch, TemplateResolutionIssue, TemplateSnapshot
from src.service.single_template_render_adapter import SingleTemplatePreflightResult


def _write_orders(path: Path, rows: list[tuple[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.append(["订单号", "模板", "生产部门", "厂家"])
    for index, (order_no, template_id) in enumerate(rows, start=1):
        sheet.append([order_no, template_id, "K" if index % 2 else "H", "Factory-A"])
    workbook.save(path)


def _snapshot(template_id: str) -> TemplateSnapshot:
    return TemplateSnapshot("legacy", template_id, "generic_rules_only", "v1", "a" * 64, "", "", "", "", ())


class FakeResolver:
    def __init__(self, *, snapshots=(), issues=()) -> None:
        self.snapshots = tuple(snapshots)
        self.issues = tuple(issues)
        self.calls: list[tuple[list[str], Path]] = []

    def resolve_many(self, template_ids, *, snapshot_root):
        self.calls.append((list(template_ids), Path(snapshot_root)))
        return TemplateResolutionBatch(self.snapshots, self.issues)


class FakeAdapter:
    def __init__(self, failures=None) -> None:
        self.failures = dict(failures or {})
        self.calls: list[tuple[str, Path, Path]] = []
        self.render_calls = 0

    def preflight(self, group, snapshot, *, group_workbook, work_dir):
        self.calls.append((group.template_id, Path(group_workbook), Path(work_dir)))
        failure = self.failures.get(group.template_id)
        if failure:
            return SingleTemplatePreflightResult(group.template_id, False, {}, {}, *failure)
        return SingleTemplatePreflightResult(
            group.template_id,
            True,
            {"template_id": group.template_id, "dry_run": True},
            {
                "stats": {"orders": len(group.rows)},
                "department_plan": {"template_id": group.template_id},
                "row_metrics": {
                    str(row.excel_row): {"planned_output_units": 1, "variable_text_length": 0}
                    for row in group.rows
                },
            },
        )


class FailingGroupWriter:
    def write(self, batch, parent_job_dir):
        raise OSError("disk unavailable")


def test_preflight_resolves_and_dry_runs_each_group_in_first_seen_order(tmp_path):
    source = tmp_path / "orders.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A"), ("B-1", "TEMPLATE-B"), ("A-2", "TEMPLATE-A")])
    resolver = FakeResolver(snapshots=(_snapshot("TEMPLATE-A"), _snapshot("TEMPLATE-B")))
    adapter = FakeAdapter()

    result = MultiTemplatePreflight(resolver=resolver, adapter=adapter).preflight(source, work_dir=tmp_path / "parent")

    assert result.status == "ready"
    assert result.can_render is True
    assert result.issues == ()
    assert resolver.calls[0][0] == ["TEMPLATE-A", "TEMPLATE-B"]
    assert [call[0] for call in adapter.calls] == ["TEMPLATE-A", "TEMPLATE-B"]
    assert [group.order_count for group in result.groups] == [2, 1]
    assert [group.excel_rows for group in result.groups] == [(2, 4), (3,)]
    assert result.source_order_sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    assert result.to_dict()["template_snapshots"][0]["template_id"] == "TEMPLATE-A"
    assert adapter.render_calls == 0
    for group in result.groups:
        assert group.can_render is True
        workbook = load_workbook(group.group_workbook, read_only=True, data_only=True)
        values = list(workbook["订单"].iter_rows(values_only=True))
        workbook.close()
        assert {row[1] for row in values[1:]} == {group.template_id}
        assert hashlib.sha256(Path(group.group_workbook).read_bytes()).hexdigest() == group.group_workbook_sha256


def test_selected_preflight_only_resolves_and_writes_the_requested_recovery_group(tmp_path):
    source = tmp_path / "orders.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A"), ("B-1", "TEMPLATE-B"), ("C-1", "TEMPLATE-C")])
    resolver = FakeResolver(snapshots=(_snapshot("TEMPLATE-B"),))
    adapter = FakeAdapter()

    result = MultiTemplatePreflight(resolver=resolver, adapter=adapter).preflight(
        source,
        work_dir=tmp_path / "recovery",
        template_ids=("TEMPLATE-B",),
    )

    assert result.status == "ready"
    assert [group.template_id for group in result.groups] == ["TEMPLATE-B"]
    assert resolver.calls[0][0] == ["TEMPLATE-B"]
    assert [call[0] for call in adapter.calls] == ["TEMPLATE-B"]
    group_dirs = [path for path in (tmp_path / "recovery" / "groups").iterdir() if path.is_dir()]
    assert len(group_dirs) == 1
    assert Path(result.groups[0].group_workbook).is_relative_to(group_dirs[0])


def test_preflight_collects_parse_resolution_and_group_errors_without_rendering(tmp_path):
    source = tmp_path / "orders.xlsx"
    _write_orders(source, [
        ("EMPTY", ""),
        ("MISSING", "MISSING"),
        ("DISABLED", "DISABLED"),
        ("V2BAD", "V2BAD"),
        ("CONFIG", "CONFIG"),
        ("FONTS", "FONTS"),
        ("ORDERS", "ORDERS"),
        ("GOOD", "GOOD"),
    ])
    resolver = FakeResolver(
        snapshots=(_snapshot("CONFIG"), _snapshot("FONTS"), _snapshot("ORDERS"), _snapshot("GOOD")),
        issues=(
            TemplateResolutionIssue("MISSING", "template_not_found", "模板不存在。", "检查模板 ID。"),
            TemplateResolutionIssue("DISABLED", "template_not_active", "模板已停用。", "启用模板。"),
            TemplateResolutionIssue("V2BAD", "template_not_published", "模板未发布。", "发布模板。"),
        ),
    )
    adapter = FakeAdapter({
        "CONFIG": ("template_config_missing", "模板配置缺失，请重新发布模板后再试。"),
        "FONTS": ("missing_required_fonts", "本机缺少模板字体：F9"),
        "ORDERS": ("v2_order_preflight_failed", "该模板的订单字段或选项未通过预检，请检查订单内容。"),
    })

    result = MultiTemplatePreflight(resolver=resolver, adapter=adapter).preflight(source, work_dir=tmp_path / "parent")

    assert result.status == "preflight_failed"
    assert result.can_render is False
    assert [call[0] for call in adapter.calls] == ["CONFIG", "FONTS", "ORDERS", "GOOD"]
    assert adapter.render_calls == 0
    assert {(issue.template_id, issue.code, issue.excel_row, issue.order_no) for issue in result.issues} == {
        ("", "template_id_missing", 2, "EMPTY"),
        ("MISSING", "template_not_found", 3, "MISSING"),
        ("DISABLED", "template_not_active", 4, "DISABLED"),
        ("V2BAD", "template_not_published", 5, "V2BAD"),
        ("CONFIG", "template_config_missing", 6, "CONFIG"),
        ("FONTS", "missing_required_fonts", 7, "FONTS"),
        ("ORDERS", "v2_order_preflight_failed", 8, "ORDERS"),
    }
    assert result.groups[-1].template_id == "GOOD"
    assert result.groups[-1].can_render is True
    assert all("C:\\" not in issue.message for issue in result.issues)


def test_preflight_rejects_an_order_sheet_without_valid_order_rows(tmp_path):
    source = tmp_path / "orders.xlsx"
    _write_orders(source, [])

    result = MultiTemplatePreflight(resolver=FakeResolver(), adapter=FakeAdapter()).preflight(source, work_dir=tmp_path / "parent")

    assert result.status == "preflight_failed"
    assert [(issue.code, issue.excel_row) for issue in result.issues] == [("orders_empty", 1)]


def test_preflight_maps_group_workbook_write_failure_to_every_affected_order_row(tmp_path):
    source = tmp_path / "orders.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A"), ("B-1", "TEMPLATE-B"), ("A-2", "TEMPLATE-A")])
    adapter = FakeAdapter()

    result = MultiTemplatePreflight(
        resolver=FakeResolver(snapshots=(_snapshot("TEMPLATE-A"), _snapshot("TEMPLATE-B"))),
        adapter=adapter,
        group_writer=FailingGroupWriter(),
    ).preflight(source, work_dir=tmp_path / "parent")

    assert result.status == "preflight_failed"
    assert adapter.calls == []
    assert {(issue.template_id, issue.code, issue.excel_row, issue.order_no) for issue in result.issues} == {
        ("TEMPLATE-A", "preflight_storage_failed", 2, "A-1"),
        ("TEMPLATE-B", "preflight_storage_failed", 3, "B-1"),
        ("TEMPLATE-A", "preflight_storage_failed", 4, "A-2"),
    }


def test_preflight_maps_source_or_group_hash_read_failures_to_business_issues(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A")])
    original_sha256 = preflight_module._sha256_file

    monkeypatch.setattr(preflight_module, "_sha256_file", lambda path: "" if Path(path) == source else original_sha256(path))
    source_result = MultiTemplatePreflight(
        resolver=FakeResolver(snapshots=(_snapshot("TEMPLATE-A"),)), adapter=FakeAdapter(),
    ).preflight(source, work_dir=tmp_path / "source-parent")

    def unreadable_group(path):
        return "" if "groups" in Path(path).parts else original_sha256(path)

    monkeypatch.setattr(preflight_module, "_sha256_file", unreadable_group)
    group_result = MultiTemplatePreflight(
        resolver=FakeResolver(snapshots=(_snapshot("TEMPLATE-A"),)), adapter=FakeAdapter(),
    ).preflight(source, work_dir=tmp_path / "group-parent")

    assert (source_result.status, source_result.issues[-1].code) == ("preflight_failed", "source_hash_unavailable")
    assert {(issue.template_id, issue.code, issue.excel_row) for issue in group_result.issues} == {
        ("TEMPLATE-A", "preflight_storage_failed", 2),
    }


def test_preflight_rejects_a_successful_adapter_result_without_row_metrics(tmp_path):
    class MetricsMissingAdapter(FakeAdapter):
        def preflight(self, group, snapshot, *, group_workbook, work_dir):
            return SingleTemplatePreflightResult(group.template_id, True, {}, {"stats": {"orders": len(group.rows)}})

    source = tmp_path / "orders.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A")])

    result = MultiTemplatePreflight(
        resolver=FakeResolver(snapshots=(_snapshot("TEMPLATE-A"),)), adapter=MetricsMissingAdapter(),
    ).preflight(source, work_dir=tmp_path / "parent")

    assert result.status == "preflight_failed"
    assert [(issue.code, issue.excel_row) for issue in result.issues] == [("preflight_plan_metrics_missing", 2)]


def test_preflight_rejects_fractional_row_metrics(tmp_path):
    class FractionalMetricsAdapter(FakeAdapter):
        def preflight(self, group, snapshot, *, group_workbook, work_dir):
            return SingleTemplatePreflightResult(
                group.template_id,
                True,
                {},
                {"row_metrics": {str(group.rows[0].excel_row): {"planned_output_units": 1.5, "variable_text_length": 0}}},
            )

    source = tmp_path / "orders.xlsx"
    _write_orders(source, [("A-1", "TEMPLATE-A")])

    result = MultiTemplatePreflight(
        resolver=FakeResolver(snapshots=(_snapshot("TEMPLATE-A"),)), adapter=FractionalMetricsAdapter(),
    ).preflight(source, work_dir=tmp_path / "parent")

    assert [(issue.code, issue.excel_row) for issue in result.issues] == [("preflight_plan_metrics_missing", 2)]
