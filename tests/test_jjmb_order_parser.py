from pathlib import Path

import pytest

from src.jjmb_combined_main import combined_personalization_text
from src.jjmb_config_grouped_main import build_grouped_task
from src.jjmb_order_parser import (
    parse_custom_info,
    parse_order_items,
    read_xlsx_rows,
    split_personalization,
)
from src.render_task import RenderTaskError
from src.jjmb_template_main import expand_values


def test_parse_custom_info_with_multiline_personalization():
    parsed = parse_custom_info(
        "Style Option:Style 2\n"
        "Font Option:F3\n"
        "Personalization:Mr Clarke\n"
        "Mrs Clarke"
    )

    assert parsed["style_option"] == "Style 2"
    assert parsed["font_option"] == "F3"
    assert parsed["personalization"] == "Mr Clarke\nMrs Clarke"


def test_split_personalization_by_lines_and_pipe():
    assert split_personalization("A\nB\nC") == ["A", "B", "C"]
    assert split_personalization("A|B|C") == ["A", "B", "C"]
    assert split_personalization("1. Payt\n2. Ave\n3) Ken\n4\u3001Luc") == ["Payt", "Ave", "Ken", "Luc"]
    assert split_personalization("Jiuwan|04.12.2026") == ["Jiuwan", "04.12.2026"]


def test_parse_order_items_normalizes_options():
    rows = [
        {
            "内部订单号": "4013367078",
            "订单明细id": "2313435",
            "购买数量": "2",
            "生产部门": "H",
            "产品中文名称": "测试产品",
            "字体颜色": "Gold",
            "外协厂家代码": "MY-H",
            "模板": "JJMB202603281027102517",
            "定制信息": "Style Option:Style 2\nFont Option:F3\nPersonalization:Mr Clarke\nMrs Clarke",
        }
    ]

    items = parse_order_items(rows, template_id="JJMB202603281027102517")

    assert len(items) == 1
    assert items[0].order_no == "4013367078"
    assert items[0].detail_id == "2313435"
    assert items[0].quantity == 2
    assert items[0].style_option == "Style2"
    assert items[0].font_option == "F3"
    assert items[0].personalization_values == ["Mr Clarke", "Mrs Clarke"]
    assert items[0].department == "H"
    assert items[0].manufacturer == "MY-H"
    assert items[0].product_name == "测试产品"
    assert items[0].color_option == "Gold"


def test_parse_order_items_from_named_sheet_split_columns(tmp_path):
    xlsx = tmp_path / "orders.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1.append(["内部订单号", "订单明细id", "购买数量", "模板", "定制信息"])
    sheet1.append(["OLD", "1", "1", "JJMB202603281027102517", "Style Option:Style 1\nFont Option:F1\nPersonalization:Old"])
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2.append(["内部订单号", "订单明细id", "购买数量", "模板", "Style Option", "Font Option", "Personalization"])
    sheet2.append(["ORDER2", "2", "1", "JJMB202603281027102517", "Style 4", "F10", "Tom|Jery"])
    workbook.save(xlsx)

    rows = read_xlsx_rows(xlsx, sheet_name="Sheet2")
    items = parse_order_items(rows, template_id="JJMB202603281027102517")

    assert len(items) == 1
    assert items[0].order_no == "ORDER2"
    assert items[0].style_option == "Style4"
    assert items[0].font_option == "F10"
    assert items[0].personalization_values == ["Tom", "Jery"]

    task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "sheet2.ai",
        columns=4,
        allowed_font_options=[f"F{i}" for i in range(1, 13)],
        sheet_name="Sheet2",
    )

    assert task.groups[0].order_no == "ORDER2"
    assert [item.text for item in task.groups[0].items] == ["Tom", "Jery"]


def test_expand_values_does_not_duplicate_by_quantity():
    item = parse_order_items(
        [
            {
                "内部订单号": "401",
                "订单明细id": "1",
                "购买数量": "5",
                "模板": "JJMB202603281027102517",
                "定制信息": "Style Option:Style 1\nFont Option:F1\nPersonalization:Only One",
            }
        ],
        template_id="JJMB202603281027102517",
    )[0]

    assert expand_values(item) == ["Only One"]


def test_combined_sheet_keeps_multiline_personalization_together():
    assert combined_personalization_text(["A", "B", "C"]) == "A\nB\nC"


def test_grouped_sheet_groups_items_by_order_number(tmp_path):
    xlsx = tmp_path / "orders.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["内部订单号", "订单明细id", "购买数量", "模板", "定制信息"])
    sheet.append(["ORDER1", "1", "1", "JJMB202603281027102517", "Style Option:Style 1\nFont Option:F1\nPersonalization:A\nB"])
    sheet.append(["ORDER1", "2", "1", "JJMB202603281027102517", "Style Option:Style 2\nFont Option:F2\nPersonalization:B"])
    sheet.append(["ORDER2", "3", "1", "JJMB202603281027102517", "Style Option:Style 1\nFont Option:F10\nPersonalization:C"])
    workbook.save(xlsx)

    task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "out.ai",
        columns=4,
    )

    assert len(task.groups) == 1
    assert task.groups[0].order_no == "ORDER1"
    assert [item.text for item in task.groups[0].items] == ["A", "B", "B"]
    payload = task.to_json_dict()
    assert payload["layout"]["show_style_boxes"] is False
    assert payload["style"]["color_name"] == "white"
    assert payload["export"]["color_mode"] == "CMYK"

    design_task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "out-design.ai",
        columns=4,
        allowed_font_options=[f"F{i}" for i in range(1, 13)],
    )

    assert len(design_task.groups) == 2
    assert design_task.groups[1].order_no == "ORDER2"
    assert design_task.groups[1].items[0].font_option == "F10"
    assert design_task.groups[1].items[0].text == "C"

    rgb_task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "out-rgb.ai",
        columns=4,
        color_mode="RGB",
    )

    assert rgb_task.to_json_dict()["export"]["color_mode"] == "RGB"


def test_grouped_sheet_keeps_design_font_as_single_asset_item(tmp_path):
    xlsx = tmp_path / "orders.xlsx"
    asset = tmp_path / "designs.ai"
    asset.write_text("fake ai", encoding="utf-8")
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["内部订单号", "订单明细id", "购买数量", "模板", "定制信息"])
    sheet.append(["ORDER1", "1", "1", "JJMB202603281027102517", "Style Option:Style 4\nFont Option:F10\nPersonalization:Tom|Jery"])
    workbook.save(xlsx)

    task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "design.ai",
        columns=4,
        allowed_font_options=[f"F{i}" for i in range(1, 13)],
        design_font_options=["F10", "F11", "F12"],
        design_asset_path=asset,
    )

    item = task.groups[0].items[0]
    assert item.render_kind == "design_asset"
    assert item.text == "Tom|Jery"
    assert item.text_parts == ["Tom", "Jery"]
    payload = task.to_json_dict()["groups"][0]["items"][0]
    assert payload["render_kind"] == "design_asset"
    assert payload["design_group"] == "F10"
    assert payload["design_asset"] == str(asset.resolve())
    assert payload["text_parts"] == ["Tom", "Jery"]


def test_grouped_sheet_expands_multiline_design_pairs_to_separate_items(tmp_path, monkeypatch):
    xlsx = tmp_path / "orders.xlsx"
    xlsx.write_bytes(b"placeholder")
    asset = tmp_path / "designs.ai"
    asset.write_text("fake ai", encoding="utf-8")

    from src.jjmb_order_parser import JJMBOrderItem

    monkeypatch.setattr("src.jjmb_config_grouped_main.read_xlsx_rows", lambda path, sheet_name=None: [{}])
    monkeypatch.setattr(
        "src.jjmb_config_grouped_main.parse_order_items",
        lambda rows, template_id=None: [
            JJMBOrderItem(
                order_no="4024990683",
                detail_id="2337497",
                quantity=2,
                template="JJMB202603281027102517",
                style_option="Style1",
                font_option="F10",
                personalization_values=["Bride|Sara", "Groom|Jordan"],
            )
        ],
    )

    task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "design.ai",
        columns=4,
        allowed_font_options=[f"F{i}" for i in range(1, 13)],
        design_font_options=["F10", "F11", "F12"],
        design_asset_path=asset,
    )

    items = task.groups[0].items
    assert len(items) == 2
    assert [item.text for item in items] == ["Bride|Sara", "Groom|Jordan"]
    assert [item.text_parts for item in items] == [["Bride", "Sara"], ["Groom", "Jordan"]]
    assert [item.quantity_index for item in items] == [1, 2]
    payloads = [item.to_json_dict() for item in items]
    assert [payload["text_parts"] for payload in payloads] == [["Bride", "Sara"], ["Groom", "Jordan"]]
    assert all("design_instances" not in payload for payload in payloads)
    assert all(payload["render_kind"] == "design_asset" for payload in payloads)

def test_grouped_sheet_uses_rule_selected_design_asset_and_group(tmp_path):
    xlsx = tmp_path / "orders.xlsx"
    asset = tmp_path / "f10-design.ai"
    asset.write_text("fake ai", encoding="utf-8")
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["内部订单号", "订单明细id", "购买数量", "模板", "定制信息"])
    sheet.append(["ORDER1", "1", "1", "JJMB202603281027102517", "Style Option:Style 4\nFont Option:F10\nPersonalization:Tom|Jery"])
    workbook.save(xlsx)

    task = build_grouped_task(
        xlsx_path=xlsx,
        template_config=Path("template.config.json"),
        output_ai=tmp_path / "design.ai",
        columns=4,
        allowed_font_options=["F10"],
        design_font_options=["F10"],
        design_asset_mappings={"F10": {"path": str(asset), "group": "F10-artwork"}},
    )

    item = task.groups[0].items[0]
    assert item.design_asset == str(asset)
    assert item.design_group == "F10-artwork"

    with pytest.raises(RenderTaskError, match="映射指向不存在的文件"):
        build_grouped_task(
            xlsx_path=xlsx,
            template_config=Path("template.config.json"),
            output_ai=tmp_path / "missing-design.ai",
            columns=4,
            allowed_font_options=["F10"],
            design_font_options=["F10"],
            design_asset_mappings={"F10": {"path": "", "group": "F10-artwork"}},
        )

    with pytest.raises(RenderTaskError, match="缺少独立设计资产映射"):
        build_grouped_task(
            xlsx_path=xlsx,
            template_config=Path("template.config.json"),
            output_ai=tmp_path / "unmapped-design.ai",
            columns=4,
            allowed_font_options=["F10"],
            design_font_options=["F10"],
            design_asset_path=asset,
            design_asset_mappings={},
        )
